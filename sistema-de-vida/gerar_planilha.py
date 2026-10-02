# -*- coding: utf-8 -*-
"""
Gera o "Sistema de Gestão da Vida" em Excel (.xlsx).

Uso:
    python3 gerar_planilha.py                 # gera modelo vazio + exemplo preenchido
    python3 gerar_planilha.py --ano 2027      # outro ano

Saídas (na mesma pasta deste script):
    Sistema_de_Vida_<ano>.xlsx            -> modelo vazio, pronto para usar
    Sistema_de_Vida_<ano>_EXEMPLO.xlsx    -> mesmo sistema com dados fictícios

Tudo que o usuário vê é calculado por fórmulas do Excel: o script só desenha a
estrutura, as fórmulas, as validações, a formatação condicional e os gráficos.
"""
import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, RadarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun)
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# =============================================================================
# Paleta (cores com função, não decoração)
# =============================================================================
INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
LINE, GRIDC, PAGE, WHITE = "D9D7D0", "E1E0D9", "F6F5F2", "FFFFFF"
DARK = "1F2430"
AUTO_BG = "F1F0EC"          # célula calculada
AUTO_HDR = "6B6A65"         # cabeçalho de coluna calculada
EDIT_HDR = "E8F1FC"         # cabeçalho editável (nomes de hábitos)
# Domínios (abas e faixas): ordem categórica validada
BLUE, ORANGE, AQUA, YELLOW = "2A78D6", "EB6834", "1BAF7A", "EDA100"
MAGENTA, VIOLET, GRAY = "E87BA4", "4A3AA7", "898781"
# Status (reservados: sempre com ícone + texto)
GOOD_T, WARN_T, CRIT_T = "006300", "8A5A00", "B42323"
GOOD_BG, WARN_BG, CRIT_BG, NEUTRAL_BG = "E3F4E3", "FEF3D6", "FBE4E4", "F0EFEC"
TODAY_BG = "FFF4CC"
SEQ_LO, SEQ_HI = "CDE2FB", "1C5CAB"
DIV_LO, DIV_MID, DIV_HI = "F4B9B9", "F0EFEC", "9EC5F4"

FONT = "Arial"
EUR = '€ #,##0.00;[Red]-€ #,##0.00;"–"'
EUR0 = '€ #,##0;[Red]-€ #,##0;€ 0'
PCT = '0%'
DATE = 'dd/mm/yyyy'
DATE_S = 'dd/mm'


class S:
    DASH = "Dashboard"
    RODA = "Roda da Vida"
    METAS = "Metas"
    TAR = "Tarefas"
    FIN = "Finanças"
    ORC = "Orçamento"
    SAU = "Saúde"
    HAB = "Hábitos"
    CAR = "Carreira"
    REL = "Relações"
    APR = "Aprendizado"
    LAZ = "Lazer"
    CASA = "Casa"
    GUIA = "Guia"
    CFG = "Config"
    CALC = "Cálculos"


ORDER = [S.DASH, S.RODA, S.METAS, S.TAR, S.FIN, S.ORC, S.SAU, S.HAB, S.CAR,
         S.REL, S.APR, S.LAZ, S.CASA, S.GUIA, S.CFG, S.CALC]
DOMAIN = {S.DASH: DARK, S.RODA: BLUE, S.METAS: BLUE, S.TAR: BLUE, S.FIN: AQUA,
          S.ORC: AQUA, S.SAU: ORANGE, S.HAB: ORANGE, S.CAR: VIOLET, S.APR: VIOLET,
          S.REL: MAGENTA, S.LAZ: MAGENTA, S.CASA: YELLOW, S.GUIA: GRAY,
          S.CFG: GRAY, S.CALC: GRAY}


def q(sheet):
    return f"'{sheet}'!"


def A(sheet, c1, r1, c2=None, r2=None):
    """Referência absoluta com nome de aba entre aspas."""
    if c2 is None:
        return f"{q(sheet)}${c1}${r1}"
    return f"{q(sheet)}${c1}${r1}:${c2}${r2}"


# =============================================================================
# Listas (aba Config)
# =============================================================================
AREAS = ["Saúde física", "Saúde mental", "Finanças", "Carreira", "Aprendizado",
         "Família", "Amor & parceria", "Amizades & social", "Lazer & criatividade",
         "Propósito & espiritualidade", "Casa & organização"]
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
MESES_ABR = [m[:3] for m in MESES]
CAT_DESP = [("Moradia", "Essencial"), ("Contas da casa", "Essencial"),
            ("Mercado", "Essencial"), ("Transporte", "Essencial"),
            ("Saúde & bem-estar", "Essencial"), ("Dívidas & financiamentos", "Essencial"),
            ("Impostos & taxas", "Essencial"), ("Família & ajudas", "Essencial"),
            ("Restaurantes & cafés", "Estilo de vida"), ("Assinaturas", "Estilo de vida"),
            ("Cuidados pessoais", "Estilo de vida"), ("Lazer & passeios", "Estilo de vida"),
            ("Viagens", "Estilo de vida"), ("Roupas", "Estilo de vida"),
            ("Casa & utensílios", "Estilo de vida"), ("Presentes & doações", "Estilo de vida"),
            ("Educação & cursos", "Crescimento"), ("Livros & materiais", "Crescimento"),
            ("Outros", "Estilo de vida")]
CAT_REC = ["Salário", "Renda extra / freelance", "Benefícios", "Rendimentos",
           "Reembolsos", "Presentes recebidos", "Outras receitas"]
CAT_APORTE = ["Reserva de emergência", "Investimentos", "Previdência", "Objetivo específico"]
CONTAS = ["Conta corrente", "Cartão de crédito", "Cartão de débito", "PIX / transferência",
          "Dinheiro", "Vale-refeição"]
PERIOD = [("Mensal", 1), ("Bimestral", 2), ("Trimestral", 3), ("Semestral", 6), ("Anual", 12)]
RELACAO = ["Família", "Parceria", "Amizade", "Trabalho", "Mentoria", "Comunidade"]

# (nome do intervalo, título, itens)  -> uma coluna cada, a partir de F
LISTS = [
    ("lst_Areas", "Áreas da vida", AREAS),
    ("Meses", "Meses", MESES),
    ("lst_MesSel", "Seleção de mês", ["Mês atual"] + MESES),
    ("lst_CatDesp", "Categorias de despesa", [c for c, _ in CAT_DESP]),
    ("lst_GrupoDesp", "Grupo da despesa", [g for _, g in CAT_DESP]),
    ("lst_CatRec", "Categorias de receita", CAT_REC),
    ("lst_CatAporte", "Destinos de aporte", CAT_APORTE),
    ("lst_Contas", "Contas / formas de pagamento", CONTAS),
    ("lst_TipoLanc", "Tipo de lançamento", ["Receita", "Despesa", "Aporte"]),
    ("lst_FixVar", "Fixo / variável", ["Fixo", "Variável"]),
    ("lst_StatusTarefa", "Status de tarefa", ["A fazer", "Em andamento", "Aguardando", "Concluída", "Cancelada"]),
    ("lst_Prioridade", "Prioridade", ["Alta", "Média", "Baixa"]),
    ("lst_StatusProjeto", "Status de projeto", ["Ideia", "Planejado", "Em andamento", "Pausado", "Concluído"]),
    ("lst_StatusMeta", "Status de meta", ["Ativa", "Pausada", "Concluída", "Abandonada"]),
    ("lst_Treino", "Tipos de treino", ["Musculação", "Corrida", "Caminhada", "Bicicleta", "Natação",
                                       "Yoga / alongamento", "Esporte coletivo", "Funcional / HIIT", "Outro"]),
    ("lst_TipoConsulta", "Tipo de consulta", ["Check-up", "Exame", "Dentista", "Especialista", "Vacina", "Terapia", "Outro"]),
    ("lst_StatusConsulta", "Status de consulta", ["A agendar", "Agendado", "Realizado"]),
    ("lst_Relacao", "Tipo de relação", RELACAO),
    ("lst_Circulo", "Círculo", ["Íntimo", "Próximo", "Social"]),
    ("lst_TipoInter", "Tipo de interação", ["Encontro", "Ligação", "Videochamada", "Mensagem", "Evento / grupo", "Ajuda / favor"]),
    ("lst_TipoApr", "Tipo de aprendizado", ["Livro", "Curso", "Idioma", "Certificação", "Podcast / vídeo", "Artigo", "Workshop", "Mentoria"]),
    ("lst_StatusApr", "Status de aprendizado", ["Quero fazer", "Em andamento", "Pausado", "Concluído", "Abandonado"]),
    ("lst_AtivEstudo", "Atividade de estudo", ["Leitura", "Aula", "Prática", "Exercícios", "Revisão", "Conversação"]),
    ("lst_CatLazer", "Categoria de lazer", ["Hobby criativo", "Esporte / ar livre", "Cultura", "Passeio", "Viagem",
                                            "Social", "Descanso", "Jogos", "Música"]),
    ("lst_StatusSonho", "Status de sonho", ["Sonho", "Planejando", "Agendado", "Realizado"]),
    ("lst_Etapa", "Etapa da candidatura", ["Interesse", "Aplicado", "Entrevista", "Teste / case", "Proposta",
                                           "Aceito", "Recusado", "Desisti"]),
    ("lst_TipoMarco", "Tipo de marco", ["Conquista", "Entrega relevante", "Feedback recebido", "Certificação",
                                        "Reconhecimento", "Promoção / aumento", "Networking", "Palestra / conteúdo"]),
    ("lst_CatComp", "Categoria de competência", ["Técnica", "Ferramenta", "Idioma", "Comportamental", "Gestão"]),
    ("lst_CatManut", "Categoria de rotina", ["Limpeza", "Manutenção", "Digital / backup", "Finanças / admin",
                                             "Veículo", "Plantas", "Outro"]),
    ("lst_Periodo", "Periodicidade", [p for p, _ in PERIOD]),
    ("lst_PeriodoMeses", "Meses por período", [m for _, m in PERIOD]),
    ("lst_Uso", "Uso", ["Alto", "Médio", "Baixo"]),
    ("lst_SimNao", "Sim / não", ["Sim", "Não"]),
]
LIST_COL = {}   # nome -> letra da coluna na Config
LIST_FIRST, LIST_LAST = 6, 40

# =============================================================================
# Estilo
# =============================================================================


def F(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def side(c=LINE, style="thin"):
    return Side(style=style, color=c)


BOX = Border(left=side(), right=side(), top=side(), bottom=side())
NOBORDER = Border()


def al(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def merge(ws, rng, value=None, font=None, fill_=None, align=None, border=None, fmt=None):
    c1 = rng.split(":")[0]
    cells = ws[rng] if ":" in rng else ((ws[rng],),)
    for row in cells:
        for c in row:
            if font:
                c.font = font
            if fill_:
                c.fill = fill_
            if align:
                c.alignment = align
            if border:
                c.border = border
            if fmt:
                c.number_format = fmt
    if value is not None:
        ws[c1] = value
    if ":" in rng and rng.split(":")[0] != rng.split(":")[1]:
        ws.merge_cells(rng)
    return ws[c1]


def link(cell, sheet, text=None, color=BLUE, size=9, bold=False):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
    if text:
        cell.value = text
    cell.font = F(size, bold, color, underline="single")


def sheet_title(ws, title, subtitle, last_col, legend=True):
    color = DOMAIN[ws.title]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color
    ws.column_dimensions["A"].width = 2.5
    last = CI(last_col)
    for c in range(1, last + 1):
        ws.cell(1, c).fill = fill(DARK)
    for r in (1, 2, 3):
        ws.cell(r, 1).fill = fill(color)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 8
    c = ws.cell(1, 2, title)
    c.font = F(16, True, WHITE)
    c.alignment = al(v="center")
    c = ws.cell(2, 2, subtitle)
    c.font = F(9, False, INK2, italic=True)
    c = ws.cell(1, last)
    link(c, S.DASH, "◀ Dashboard", WHITE, 9, True)
    c.alignment = al("right")
    if legend:
        c = ws.cell(3, 2, "Legenda:  branco = você preenche  ·  cinza = automático (não editar)  ·  "
                          "✔ ok   ▲ atenção   ✖ crítico   ○ sem dados")
        c.font = F(8, False, MUTED)


def section(ws, row, c1, c2, text, color=None, height=22):
    color = color or DOMAIN[ws.title]
    for c in range(CI(c1), CI(c2) + 1):
        cell = ws.cell(row, c)
        cell.border = Border(bottom=side(color, "medium"))
    cell = ws[f"{c1}{row}"]
    cell.value = text
    cell.font = F(11, True, INK)
    cell.alignment = al(v="bottom")
    if height:
        ws.row_dimensions[row].height = height


def headers(ws, row, specs, height=30):
    """specs: [(col, texto, 'in'|'auto'|'edit', largura_merge_ate)]"""
    for spec in specs:
        col, text, kind = spec[0], spec[1], spec[2]
        end = spec[3] if len(spec) > 3 else col
        bg = {"in": DARK, "auto": AUTO_HDR, "edit": EDIT_HDR}[kind]
        fc = INK if kind == "edit" else WHITE
        merge(ws, f"{col}{row}:{end}{row}", text, F(9, True, fc), fill(bg),
              al("center", "center", True), Border(left=side(WHITE), right=side(WHITE)))
    if height:
        ws.row_dimensions[row].height = height


def body(ws, c1, c2, r1, r2, kind="in", fmt=None, h="left", wrap=False, size=9):
    bg = fill(WHITE) if kind == "in" else fill(AUTO_BG)
    color = INK if kind == "in" else INK2
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=CI(c1), max_col=CI(c2)):
        for c in row:
            c.fill = bg
            c.border = BOX
            c.font = F(size, False, color)
            c.alignment = al(h, "center", wrap, 1 if h == "left" else 0)
            if fmt:
                c.number_format = fmt


def formulas(ws, col, r1, r2, fn):
    for r in range(r1, r2 + 1):
        ws[f"{col}{r}"] = fn(r)


def merge_rows(ws, c1, c2, r1, r2):
    for r in range(r1, r2 + 1):
        ws.merge_cells(f"{c1}{r}:{c2}{r}")


# ---- validação ----------------------------------------------------------------

def dv_list(ws, rng, source, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "stop" if strict else "warning"
    dv.errorTitle = "Valor fora da lista"
    dv.error = "Escolha um item da lista (as listas ficam na aba Config)."
    if prompt:
        dv.showInputMessage = True
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_num(ws, rng, lo, hi, whole=False, prompt=None, title=None):
    dv = DataValidation(type="whole" if whole else "decimal", operator="between",
                        formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Valor inválido"
    dv.error = f"Use um número entre {lo} e {hi}."
    if prompt:
        dv.showInputMessage = True
        dv.promptTitle = title or ""
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_date(ws, rng, prompt="Data no formato dd/mm/aaaa"):
    dv = DataValidation(type="date", operator="between", formula1="36526", formula2="73051",
                        allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Data inválida"
    dv.error = "Digite uma data (ex.: 15/03/2026). Atalho: Ctrl+; insere a data de hoje."
    dv.showInputMessage = True
    dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


# ---- formatação condicional --------------------------------------------------

def cf_status(ws, rng, bold=True):
    tl = rng.split(":")[0].replace("$", "")
    for sym, color in (("✔", GOOD_T), ("▲", WARN_T), ("✖", CRIT_T), ("○", MUTED),
                       ("‖", MUTED), ("–", MUTED)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({tl},1)="{sym}"'],
                                                       font=Font(color=color, bold=bold)))


def cf_zebra(ws, rng, key_col):
    tl_row = int("".join(ch for ch in rng.split(":")[0] if ch.isdigit()))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(${key_col}{tl_row}<>"",MOD(ROW(),2)=0)'], fill=fill("FAFAF8")))


def cf_bar(ws, rng, lo=0, hi=1, color=BLUE):
    ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=lo, end_type="num",
                                                   end_value=hi, color=color, showValue=True))


def cf_diverge(ws, rng, lo, mid, hi, reverse=False):
    a, b = (DIV_HI, DIV_LO) if reverse else (DIV_LO, DIV_HI)
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=lo, start_color=a,
                                                      mid_type="num", mid_value=mid, mid_color=DIV_MID,
                                                      end_type="num", end_value=hi, end_color=b))


# ---- faixa de KPIs no topo de cada aba ----------------------------------------

K = {}   # endereços dos indicadores, usados pelo Cálculos e pelo Dashboard


def kpi_strip(ws, row, items):
    """items: (chave, col1, col2, rótulo, fórmula, formato, fórmula_sub)"""
    color = DOMAIN[ws.title]
    for key, c1, c2, label, fml, fmt, sub in items:
        merge(ws, f"{c1}{row}:{c2}{row}", label.upper(), F(7.5, True, MUTED), fill(WHITE),
              al("left", "bottom", indent=1))
        merge(ws, f"{c1}{row+1}:{c2}{row+1}", fml, F(16, True, INK), fill(WHITE),
              al("left", "center", indent=1), fmt=fmt)
        merge(ws, f"{c1}{row+2}:{c2}{row+2}", sub, F(8, False, INK2), fill(WHITE),
              al("left", "top", wrap=True, indent=1))
        for r in range(row, row + 3):
            for c in range(CI(c1), CI(c2) + 1):
                cell = ws.cell(r, c)
                cell.border = Border(top=side(color, "thick") if r == row else None,
                                     bottom=side(LINE) if r == row + 2 else None,
                                     left=side(LINE) if c == CI(c1) else None,
                                     right=side(LINE) if c == CI(c2) else None)
        K[key] = A(ws.title, c1, row + 1)
        K[key + "_sub"] = A(ws.title, c1, row + 2)
    cf_status(ws, f"B{row+2}:{CL(max(CI(i[2]) for i in items))}{row+2}")
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row + 1].height = 28
    ws.row_dimensions[row + 2].height = 25


def in_month(col_rng, start="RefIni", end="RefFim"):
    return f'{col_rng},">="&{start},{col_rng},"<="&{end}'


def month_bounds(m):
    return f"DATE(Ano,{m},1)", f"DATE(Ano,{m}+1,0)"


def smin(cond, values):
    """Mínimo condicional compatível com Excel 2010+ (sem MINIFS): 99999 = nenhum."""
    return f"SUMPRODUCT(MIN({cond}*{values}+(1-{cond})*99999))"


# =============================================================================
# Config
# =============================================================================
PARAMS = [  # (linha, nome, rótulo, valor padrão, formato, nota)
    (6, "Nome", "Seu nome", "", None, "Aparece no topo do Dashboard."),
    (7, "Ano", "Ano deste arquivo", None, "0", "Um arquivo por ano. Mudar o ano reposiciona as datas das abas Saúde e Hábitos: faça isso só ao iniciar um arquivo novo e vazio."),
    (8, "Moeda", "Moeda", "€ (EUR)", None, "Informativo. Os valores usam o formato €; para trocar, altere o formato numérico das colunas de valor."),
    (9, "MetaPoup", "Meta: taxa de poupança", 0.20, PCT, "Parte da receita que sobra (receitas − despesas) ÷ receitas."),
    (10, "MetaReserva", "Meta: reserva de emergência (meses)", 6, "0", "Quantos meses de despesas a reserva deve cobrir."),
    (11, "MetaSono", "Meta: horas de sono por noite", 7.5, "0.0", ""),
    (12, "MetaTreinosSem", "Meta: treinos por semana", 4, "0", ""),
    (13, "MetaPassos", "Meta: passos por dia", 8000, "#,##0", ""),
    (14, "PesoAlvo", "Peso alvo (kg)", None, "0.0", "Opcional."),
    (15, "MetaEstudoMes", "Meta: horas de estudo por mês", 24, "0", ""),
    (16, "MetaLivrosAno", "Meta: livros no ano", 12, "0", ""),
    (17, "MetaLazerSem", "Meta: horas de lazer por semana", 6, "0", "Lazer de verdade: tempo sem obrigação."),
    (18, "AlertaDocs", "Avisar documentos com quantos dias de antecedência", 90, "0", ""),
    (19, "AlertaAniv", "Avisar aniversários com quantos dias de antecedência", 30, "0", ""),
    (20, "FollowUpDias", "Follow-up de candidatura após (dias)", 10, "0", ""),
]
DERIVED = [  # (linha, nome, rótulo, fórmula, formato)
    (23, "Hoje", "Hoje", "=TODAY()", DATE),
    (24, "RefMes", "Mês de referência (nº)",
     f'=IF({q(S.DASH)}$Z$1="Mês atual",IF(YEAR(Hoje)=Ano,MONTH(Hoje),IF(YEAR(Hoje)>Ano,12,1)),IFERROR(MATCH({q(S.DASH)}$Z$1,Meses,0),1))', "0"),
    (25, "RefIni", "Início do mês de referência", "=DATE(Ano,RefMes,1)", DATE),
    (26, "RefFim", "Fim do mês de referência", "=DATE(Ano,RefMes+1,0)", DATE),
    (27, "Corte", "Data de corte (até quando contar)", "=MIN(Hoje,RefFim)", DATE),
    (28, "DiasRef", "Dias já decorridos no mês de referência", "=IF(RefIni>Hoje,0,IF(RefFim<=Hoje,DAY(RefFim),DAY(Hoje)))", "0"),
    (29, "IniAno", "Início do ano", "=DATE(Ano,1,1)", DATE),
    (30, "FimAno", "Fim do ano", "=DATE(Ano,12,31)", DATE),
    (31, "DiasAno", "Dias decorridos no ano (até o corte)", "=MAX(0,MIN(Corte,FimAno)-IniAno+1)", "0"),
    (32, "IdxHoje", "Posição de hoje nas abas diárias", "=IF(Hoje<IniAno,0,MIN(FimAno,Hoje)-IniAno+1)", "0"),
    (33, "RotuloMes", "Rótulo do mês", '=INDEX(Meses,RefMes)&" "&Ano', None),
    (34, "DiasMes", "Dias no mês de referência", "=DAY(RefFim)", "0"),
]


def build_config(wb, ws, ano, nome):
    sheet_title(ws, "Configurações & listas",
                "Ajuste seus parâmetros e metas. As listas alimentam os menus suspensos de todas as abas.", "AN",
                legend=False)
    setw(ws, {"B": 46, "C": 16, "D": 60, "E": 3})
    section(ws, 5, "B", "D", "PARÂMETROS E METAS PESSOAIS")
    for r, name, label, val, fmt, note in PARAMS:
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = F(9)
        c = ws[f"C{r}"]
        c.value = ano if name == "Ano" else (nome if name == "Nome" else val)
        c.font = F(10, True, INK)
        c.fill = fill(WHITE)
        c.border = Border(left=side(BLUE), right=side(BLUE), top=side(BLUE), bottom=side(BLUE))
        c.alignment = al("center")
        if fmt:
            c.number_format = fmt
        ws[f"D{r}"] = note
        ws[f"D{r}"].font = F(8, False, MUTED, italic=True)
        ws[f"D{r}"].alignment = al(wrap=True)
        wb.defined_names[name] = DefinedName(name, attr_text=A(S.CFG, "C", r))
    ws.row_dimensions[7].height = 26
    dv_num(ws, "C7", 2000, 2100, True)
    dv_num(ws, "C9", 0, 1)
    section(ws, 22, "B", "D", "DATAS DE REFERÊNCIA (automático)")
    for r, name, label, fml, fmt in DERIVED:
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = F(9, False, INK2)
        c = ws[f"C{r}"]
        c.value = fml
        c.font = F(9, False, INK2)
        c.fill = fill(AUTO_BG)
        c.border = BOX
        c.alignment = al("center")
        if fmt:
            c.number_format = fmt
        wb.defined_names[name] = DefinedName(name, attr_text=A(S.CFG, "C", r))
    ws["D24"] = "Escolhido no topo do Dashboard (célula 'Mês de referência')."
    ws["D24"].font = F(8, False, MUTED, italic=True)

    # listas
    col = CI("F")
    for name, title, items in LISTS:
        L = CL(col)
        LIST_COL[name] = L
        ws.column_dimensions[L].width = 24
        merge(ws, f"{L}5", title, F(9, True, WHITE), fill(DARK), al("center", "center", True))
        ws.row_dimensions[5].height = 30
        for i in range(LIST_FIRST, LIST_LAST + 1):
            c = ws[f"{L}{i}"]
            c.border = BOX
            c.font = F(9)
        for i, it in enumerate(items):
            ws[f"{L}{LIST_FIRST + i}"] = it
        fixed = name in ("lst_Areas", "Meses", "lst_MesSel")
        if fixed:
            ref = A(S.CFG, L, LIST_FIRST, L, LIST_FIRST + len(items) - 1)
        else:
            rng = A(S.CFG, L, LIST_FIRST, L, LIST_LAST)
            ref = f"OFFSET({A(S.CFG, L, LIST_FIRST)},0,0,MAX(1,COUNTA({rng})),1)"
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
        col += 1
    ws["F4"] = ("Listas: edite ou acrescente itens embaixo de cada título (sem deixar linhas vazias no meio). "
                "As 11 áreas da vida podem ser renomeadas, mas mantenha a quantidade e a ordem.")
    ws["F4"].font = F(8, False, MUTED, italic=True)
    ws.freeze_panes = "A6"


# =============================================================================
# Finanças (lançamentos)
# =============================================================================
FIN_R1, FIN_R2 = 10, 1509


def fin(col):
    return A(S.FIN, col, FIN_R1, col, FIN_R2)


def build_financas(ws):
    sheet_title(ws, "Finanças · lançamentos",
                "Registre cada receita, despesa e aporte. Orçamento, patrimônio e Dashboard se atualizam sozinhos.", "K")
    setw(ws, {"B": 12, "C": 12, "D": 26, "E": 36, "F": 14, "G": 22, "H": 11, "I": 32, "J": 16, "K": 14})
    r1, r2 = FIN_R1, FIN_R2
    B, C, D, Fv = (f"${c}${r1}:${c}${r2}" for c in "BCDF")
    sumt = lambda t: f'SUMIFS({Fv},{C},"{t}",{in_month(B)})'
    kpi_strip(ws, 5, [
        ("fin_rec", "B", "C", "Receitas do mês", f"={sumt('Receita')}", EUR0, '="Mês: "&RotuloMes'),
        ("fin_desp", "D", "D", "Despesas do mês", f"={sumt('Despesa')}", EUR0,
         f'=IF({A(S.ORC,"D",43)}>0,IF(D6>{A(S.ORC,"D",43)},"✖ ","✔ ")&ROUND(D6/{A(S.ORC,"D",43)}*100,0)&"% do orçamento","○ Orçamento não definido")'),
        ("fin_res", "E", "E", "Resultado (receitas − despesas)", "=B6-D6", EUR0,
         '=IF(B6+D6=0,"○ Sem lançamentos",IF(E6>=0,"✔ Sobrou dinheiro","✖ Gastou mais do que ganhou"))'),
        ("fin_apo", "F", "G", "Aportes (poupança / investimento)", f"={sumt('Aporte')}", EUR0,
         '="Saldo em caixa: € "&FIXED(E6-F6,0)'),
        ("fin_taxa", "H", "I", "Taxa de poupança", '=IF(B6>0,E6/B6,"–")', PCT,
         '=IF(ISNUMBER(H6),IF(H6>=MetaPoup,"✔ ","▲ ")&"meta "&ROUND(MetaPoup*100,0)&"%","○ Sem receitas no mês")'),
        ("fin_n", "J", "K", "Lançamentos no mês", f"=COUNTIFS({in_month(B)})", "0",
         f'=IF(COUNTIF($K${r1}:$K${r2},"▲*")>0,"▲ "&COUNTIF($K${r1}:$K${r2},"▲*")&" linha(s) incompleta(s)","✔ Tudo completo")'),
    ])
    headers(ws, 9, [("B", "Data", "in"), ("C", "Tipo", "in"), ("D", "Categoria", "in"),
                    ("E", "Descrição", "in"), ("F", "Valor (€)", "in"), ("G", "Conta / forma", "in"),
                    ("H", "Fixo / variável", "in"), ("I", "Notas", "in"), ("J", "Grupo", "auto"),
                    ("K", "Verificação", "auto")])
    body(ws, "B", "B", r1, r2, fmt=DATE, h="center")
    body(ws, "C", "E", r1, r2)
    body(ws, "F", "F", r1, r2, fmt='#,##0.00', h="right")
    body(ws, "G", "I", r1, r2)
    body(ws, "J", "K", r1, r2, "auto")
    cat = A(S.CFG, LIST_COL["lst_CatDesp"], 6, LIST_COL["lst_CatDesp"], 40)
    grp = A(S.CFG, LIST_COL["lst_GrupoDesp"], 6, LIST_COL["lst_GrupoDesp"], 40)
    formulas(ws, "J", r1, r2, lambda r: f'=IF(C{r}<>"Despesa","",IFERROR(INDEX({grp},MATCH(D{r},{cat},0)),"—"))')
    formulas(ws, "K", r1, r2, lambda r: f'=IF(COUNTA(B{r}:F{r})=0,"",IF(OR(B{r}="",C{r}="",D{r}="",F{r}=""),"▲ Completar","✔"))')
    dv_date(ws, f"B{r1}:B{r2}")
    dv_list(ws, f"C{r1}:C{r2}", "lst_TipoLanc", "Receita, Despesa ou Aporte (dinheiro guardado/investido)")
    dv_list(ws, f"D{r1}:D{r2}", f'IF($C{r1}="Receita",lst_CatRec,IF($C{r1}="Aporte",lst_CatAporte,lst_CatDesp))',
            "A lista muda conforme o Tipo")
    dv_num(ws, f"F{r1}:F{r2}", 0, 10000000, prompt="Valor sempre positivo; o Tipo define se entra ou sai.")
    dv_list(ws, f"G{r1}:G{r2}", "lst_Contas")
    dv_list(ws, f"H{r1}:H{r2}", "lst_FixVar")
    cf_zebra(ws, f"B{r1}:K{r2}", "B")
    for t, color in (("Receita", GOOD_T), ("Aporte", BLUE)):
        ws.conditional_formatting.add(f"C{r1}:C{r2}", FormulaRule(formula=[f'$C{r1}="{t}"'], font=Font(color=color, bold=True)))
        ws.conditional_formatting.add(f"F{r1}:F{r2}", FormulaRule(formula=[f'$C{r1}="{t}"'], font=Font(color=color, bold=True)))
    cf_status(ws, f"K{r1}:K{r2}")
    ws.auto_filter.ref = f"B9:K{r2}"
    ws.freeze_panes = "C10"


# =============================================================================
# Orçamento (anual, planejado × realizado) + patrimônio
# =============================================================================
ORC_DESP1, ORC_DESP2 = 18, 42
ORC_REC1, ORC_REC2 = 48, 62
ORC_APO1, ORC_APO2 = 67, 74
MCOLS = [CL(5 + i) for i in range(12)]   # E..P


def build_orcamento(ws):
    sheet_title(ws, "Orçamento · planejado × realizado e patrimônio",
                "Defina quanto quer gastar por categoria. O realizado vem de Finanças. No fim de cada mês, atualize o patrimônio.", "U")
    setw(ws, {"B": 30, "C": 14, "D": 13, **{c: 10.5 for c in MCOLS}, "Q": 12, "R": 11, "S": 12, "T": 10, "U": 24, "V": 4})
    ws.column_dimensions["V"].hidden = True
    Bf, Cf, Df, Ff = fin("B"), fin("C"), fin("D"), fin("F")

    def month_sum(tipo, m, cat_cell=None):
        ini, fim = month_bounds(m)
        extra = f",{Df},{cat_cell}" if cat_cell else ""
        return f'SUMIFS({Ff},{Cf},"{tipo}"{extra},{Bf},">="&{ini},{Bf},"<="&{fim})'

    def head(row, first, d_label, group=False):
        specs = [("B", first, "in"), ("C", "Grupo" if group else "", "in"), ("D", d_label, "in")]
        specs += [(c, MESES_ABR[i], "auto") for i, c in enumerate(MCOLS)]
        specs += [("Q", "Total ano", "auto"), ("R", "Média/mês até a ref.", "auto"), ("S", "Mês de referência", "auto"),
                  ("T", "% do plano", "auto"), ("U", "Status no mês", "auto")]
        headers(ws, row, specs, 30)

    avg_ref = lambda r: f"SUMPRODUCT((COLUMN($E{r}:$P{r})-COLUMN($E{r})+1<=RefMes)*$E{r}:$P{r})/RefMes"

    # --- resumo mensal
    section(ws, 5, "B", "U", "RESUMO MENSAL")
    head(6, "Indicador", "Plano / mês")
    rows = [(7, "Receitas"), (8, "Despesas"), (9, "Resultado (receitas − despesas)"), (10, "Aportes (poupança / invest.)"),
            (11, "Saldo em caixa (resultado − aportes)"), (12, "Taxa de poupança"), (13, "Orçamento de despesas"),
            (14, "Desvio (real − orçado)")]
    for r, label in rows:
        ws[f"B{r}"] = label
    body(ws, "B", "C", 7, 14, "in")
    body(ws, "D", "U", 7, 14, "auto", fmt='#,##0;[Red]-#,##0;"–"', h="right")
    for r in range(7, 15):
        ws[f"B{r}"].font = F(9, True, INK)
    ws["D7"] = f"=D{ORC_REC2+1}"
    ws["D8"] = f"=D{ORC_DESP2+1}"
    ws["D9"] = "=D7-D8"
    ws["D12"] = '=IF(D7>0,D9/D7,"")'
    for i, c in enumerate(MCOLS):
        m = i + 1
        ws[f"{c}7"] = f"={month_sum('Receita', m)}"
        ws[f"{c}8"] = f"={month_sum('Despesa', m)}"
        ws[f"{c}9"] = f"={c}7-{c}8"
        ws[f"{c}10"] = f"={month_sum('Aporte', m)}"
        ws[f"{c}11"] = f"={c}9-{c}10"
        ws[f"{c}12"] = f'=IF({c}7>0,{c}9/{c}7,"")'
        ws[f"{c}13"] = f"=$D${ORC_DESP2+1}"
        ws[f"{c}14"] = f'=IF({c}8=0,"",{c}8-{c}13)'
    for r in (7, 8, 9, 10, 11):
        ws[f"Q{r}"] = f"=SUM(E{r}:P{r})"
        ws[f"R{r}"] = f"={avg_ref(r)}"
        ws[f"S{r}"] = f"=INDEX(E{r}:P{r},RefMes)"
    ws["Q12"], ws["R12"], ws["S12"] = '=IF(Q7>0,Q9/Q7,"")', '=IF(R7>0,R9/R7,"")', "=INDEX(E12:P12,RefMes)"
    ws["Q13"], ws["S13"] = "=D8*12", "=D8"
    ws["Q14"], ws["S14"] = "=SUM(E14:P14)", "=INDEX(E14:P14,RefMes)"
    for c in ["D"] + MCOLS + ["Q", "R", "S"]:
        ws[f"{c}12"].number_format = '0%;[Red]-0%;"–"'
    ws["T7"] = '=IF(D7>0,S7/D7,"")'
    ws["T8"] = '=IF(D8>0,S8/D8,"")'
    ws["T7"].number_format = ws["T8"].number_format = PCT
    ws["U7"] = '=IF(T7="","",IF(T7>=1,"✔ Receita prevista atingida","▲ "&ROUND(T7*100,0)&"% do previsto"))'
    ws["U8"] = '=IF(T8="","",IF(T8>1,"✖ Acima do orçamento",IF(T8>DiasRef/DiasMes+0.1,"▲ Gastando rápido","✔ Dentro do orçamento")))'
    ws["U12"] = '=IF(S12="","",IF(S12>=MetaPoup,"✔ Meta de poupança atingida","▲ Abaixo da meta de "&ROUND(MetaPoup*100,0)&"%"))'
    cf_status(ws, "U7:U14")

    def cat_block(sec_row, title, d_label, r1, r2, tipo, list_name, with_group):
        section(ws, sec_row, "B", "U", title)
        head(sec_row + 1, "Categoria", d_label, with_group)
        L = LIST_COL[list_name]
        body(ws, "B", "C", r1, r2, "auto")
        body(ws, "D", "D", r1, r2, "in", fmt='#,##0;;"–"', h="right")
        body(ws, "E", "S", r1, r2, "auto", fmt='#,##0;[Red]-#,##0;"–"', h="right")
        body(ws, "T", "T", r1, r2, "auto", fmt=PCT, h="right")
        body(ws, "U", "U", r1, r2, "auto")
        for i, r in enumerate(range(r1, r2 + 1)):
            src = A(S.CFG, L, LIST_FIRST + i)
            ws[f"B{r}"] = f'=IF({src}="","",{src})'
            if with_group:
                g = A(S.CFG, LIST_COL["lst_GrupoDesp"], LIST_FIRST + i)
                ws[f"C{r}"] = f'=IF(B{r}="","",{g}&"")'
            for j, c in enumerate(MCOLS):
                ws[f"{c}{r}"] = f'=IF($B{r}="","",{month_sum(tipo, j + 1, f"$B{r}")})'
            ws[f"Q{r}"] = f'=IF(B{r}="","",SUM(E{r}:P{r}))'
            ws[f"R{r}"] = f'=IF(B{r}="","",{avg_ref(r)})'
            ws[f"S{r}"] = f'=IF(B{r}="","",INDEX(E{r}:P{r},RefMes))'
            ws[f"T{r}"] = f'=IF(OR(B{r}="",N(D{r})=0),"",S{r}/D{r})'
            if tipo == "Despesa":
                ws[f"U{r}"] = (f'=IF(T{r}="",IF(AND(B{r}<>"",N(S{r})>0),"○ Sem orçamento",""),IF(T{r}>1,"✖ Estourou",'
                               f'IF(T{r}>DiasRef/DiasMes+0.15,"▲ Ritmo alto","✔ OK")))')
            else:
                ws[f"U{r}"] = f'=IF(T{r}="","",IF(T{r}>=1,"✔ Atingido","▲ "&ROUND(T{r}*100,0)&"%"))'
        t = r2 + 1
        ws[f"B{t}"] = "TOTAL"
        for c in ["D"] + MCOLS + ["Q", "R", "S"]:
            ws[f"{c}{t}"] = f"=SUM({c}{r1}:{c}{r2})"
        ws[f"T{t}"] = f'=IF(D{t}>0,S{t}/D{t},"")'
        body(ws, "B", "U", t, t, "auto", fmt='#,##0;[Red]-#,##0;"–"', h="right")
        ws[f"T{t}"].number_format = PCT
        for c in range(2, 22):
            ws.cell(t, c).font = F(9, True, INK)
            ws.cell(t, c).border = Border(top=side(INK, "thin"), bottom=side(INK, "thin"))
        ws[f"B{t}"].alignment = al("left", indent=1)
        dv_num(ws, f"D{r1}:D{r2}", 0, 10000000, prompt="Quanto você planeja por mês nesta categoria.")
        cf_status(ws, f"U{r1}:U{r2}")
        cf_bar(ws, f"T{r1}:T{r2}", 0, 1, BLUE)
        return t

    t = cat_block(16, "DESPESAS POR CATEGORIA", "Orçamento / mês", ORC_DESP1, ORC_DESP2, "Despesa", "lst_CatDesp", True)
    # linha de conferência
    r = t + 1
    ws[f"B{r}"] = "Fora da lista de categorias"
    for c in MCOLS + ["Q", "S"]:
        ws[f"{c}{r}"] = f"={c}8-{c}{t}" if c not in ("Q", "S") else f"=SUM(E{r}:P{r})" if c == "Q" else f"=INDEX(E{r}:P{r},RefMes)"
    body(ws, "B", "S", r, r, "auto", fmt='#,##0;[Red]-#,##0;"–"', h="right")
    ws[f"B{r}"].font = F(8, False, MUTED, italic=True)
    ws.conditional_formatting.add(f"E{r}:S{r}", FormulaRule(formula=[f"ABS(E{r})>0.005"], fill=fill(WARN_BG)))
    ws[f"U{r}"] = "← deveria ser zero"
    ws[f"U{r}"].font = F(8, False, MUTED, italic=True)

    cat_block(46, "RECEITAS POR CATEGORIA", "Previsto / mês", ORC_REC1, ORC_REC2, "Receita", "lst_CatRec", False)
    cat_block(65, "APORTES POR DESTINO", "Planejado / mês", ORC_APO1, ORC_APO2, "Aporte", "lst_CatAporte", False)

    # --- patrimônio
    section(ws, 77, "B", "U", "PATRIMÔNIO · preencha os saldos no fim de cada mês")
    specs = [("B", "Item", "in"), ("C", "", "in"), ("D", "", "in")]
    specs += [(c, MESES_ABR[i], "in") for i, c in enumerate(MCOLS)]
    specs += [("Q", "", "auto"), ("R", "", "auto"), ("S", "Mês de referência", "auto"), ("T", "", "auto"), ("U", "", "auto")]
    headers(ws, 78, specs)
    items = [(79, "Contas & dinheiro em caixa", "in"), (80, "Reserva de emergência", "in"), (81, "Investimentos", "in"),
             (82, "Outros bens (carro, imóvel…)", "in"), (83, "(−) Dívidas", "in"), (84, "PATRIMÔNIO LÍQUIDO", "auto"),
             (85, "Variação no mês", "auto"), (86, "Reserva cobre (meses de despesa)", "auto")]
    for r, label, kind in items:
        ws[f"B{r}"] = label
        body(ws, "B", "D", r, r, "auto")
        body(ws, "E", "P", r, r, kind, fmt='#,##0;[Red]-#,##0;"–"', h="right")
        body(ws, "Q", "U", r, r, "auto", fmt='#,##0;[Red]-#,##0;"–"', h="right")
    ws["B84"].font = F(9, True, INK)
    for i, c in enumerate(MCOLS):
        ws[f"{c}84"] = f'=IF(COUNT({c}79:{c}83)=0,"",N({c}79)+N({c}80)+N({c}81)+N({c}82)-N({c}83))'
        ws[f"{c}84"].font = F(9, True, INK)
        if i > 0:
            p = MCOLS[i - 1]
            ws[f"{c}85"] = f'=IF(OR({c}84="",{p}84=""),"",{c}84-{p}84)'
        else:
            ws[f"{c}85"] = '=""'
        ws[f"{c}86"] = f'=IF(OR({c}80="",SUM($E$8:{c}8)=0),"",{c}80/(SUM($E$8:{c}8)/COUNTIF($E$8:{c}8,">0")))'
        ws[f"{c}86"].number_format = '0.0;;"–"'
    ws["V84"] = "=SUMPRODUCT(MAX(($E$84:$P$84<>\"\")*(COLUMN($E$84:$P$84)-COLUMN($E$84)+1<=RefMes)*(COLUMN($E$84:$P$84)-COLUMN($E$84)+1)))"
    for r in range(79, 87):
        ws[f"S{r}"] = f'=IF($V$84=0,"",INDEX(E{r}:P{r},$V$84))'
    ws["S86"].number_format = '0.0;;"–"'
    ws["T84"] = '=IF($V$84=0,"",IF($V$84=RefMes,"","dado de "&INDEX(Meses,$V$84)))'
    ws["T84"].font = F(8, False, MUTED, italic=True)
    ws["U86"] = '=IF(S86="","",IF(S86>=MetaReserva,"✔ Meta atingida","▲ Meta: "&MetaReserva&" meses"))'
    cf_status(ws, "U79:U86")
    dv_num(ws, "E79:P83", -100000000, 100000000)

    # --- distribuição 50/30/20
    section(ws, 88, "B", "U", "DISTRIBUIÇÃO DA RECEITA NO MÊS · referência 50/30/20")
    headers(ws, 89, [("B", "Destino", "auto"), ("C", "Valor", "auto"), ("D", "% da receita", "auto"),
                     ("E", "Referência", "auto", "F"), ("G", "Leitura", "auto", "L")])
    dist = [(90, "Essencial", 0.5, "<="), (91, "Estilo de vida", 0.3, "<="), (92, "Crescimento", None, None),
            (93, "Poupança (resultado do mês)", 0.2, ">=")]
    D1, D2 = ORC_DESP1, ORC_DESP2
    for r, label, ref, op in dist:
        ws[f"B{r}"] = label
        if r < 93:
            ws[f"C{r}"] = f'=SUMIFS($S${D1}:$S${D2},$C${D1}:$C${D2},"{label}")'
        else:
            ws[f"C{r}"] = "=MAX(0,S9)"
        ws[f"D{r}"] = f'=IF(S7>0,C{r}/S7,"")'
        merge(ws, f"E{r}:F{r}")
        merge(ws, f"G{r}:L{r}")
        ws[f"E{r}"] = ref if ref is not None else "—"
        if ref is None:
            ws[f"G{r}"] = '=IF(D92="","","○ Investimento em você: cursos, livros")'
        elif op == "<=":
            ws[f"G{r}"] = f'=IF(D{r}="","",IF(D{r}<=E{r}+0.02,"✔ Dentro da referência","▲ Acima da referência"))'
        else:
            ws[f"G{r}"] = f'=IF(D{r}="","",IF(D{r}>=E{r},"✔ Dentro da referência","▲ Abaixo da referência"))'
        body(ws, "B", "L", r, r, "auto")
        ws[f"C{r}"].number_format = '#,##0'
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = PCT
    cf_status(ws, "G90:G93")
    cf_bar(ws, "D90:D93", 0, 1, AQUA)
    K["orc_total"] = A(S.ORC, "D", ORC_DESP2 + 1)
    K["pl_ref"] = A(S.ORC, "S", 84)
    K["pl_var"] = A(S.ORC, "S", 85)
    K["reserva_meses"] = A(S.ORC, "S", 86)
    ws.freeze_panes = "E5"


# =============================================================================
# Saúde (registro diário + resumo mensal + consultas)
# =============================================================================
SAU_R1, SAU_R2 = 11, 376


def build_saude(ws):
    sheet_title(ws, "Saúde & bem-estar",
                "Uma linha por dia (as datas já estão prontas). Leva 1 minuto: sono, humor, energia, treino. Peso pode ser semanal.", "Z")
    setw(ws, {"B": 12, "C": 6, "D": 8, "E": 9, "F": 8, "G": 8, "H": 9, "I": 17, "J": 9, "K": 9, "L": 8,
              "M": 7, "N": 10, "O": 42, "P": 3, "Q": 11, "R": 12, "S": 8, "T": 8, "U": 10, "V": 10,
              "W": 10, "X": 10, "Y": 9, "Z": 11})
    r1, r2 = SAU_R1, SAU_R2
    rg = lambda c: f"${c}${r1}:${c}${r2}"
    Bd = rg("B")
    avg = lambda c: f"IFERROR(AVERAGEIFS({rg(c)},{in_month(Bd)}),\"–\")"
    c = ws["O3"]
    c.value = f'=HYPERLINK("#\'{S.SAU}\'!B"&MAX({r1},IdxHoje+{r1 - 1}),"→ Ir para a linha de hoje")'
    c.font = F(9, True, BLUE, underline="single")
    kpi_strip(ws, 5, [
        ("sau_sono", "B", "C", "Sono médio", f"={avg('D')}", '0.0" h"',
         '=IF(ISNUMBER(B6),IF(B6>=MetaSono,"✔ ","▲ ")&"meta "&FIXED(MetaSono,1)&" h","○ Sem registros")'),
        ("sau_qual", "D", "E", "Qualidade do sono", f"={avg('E')}", '0.0"/5"', '="1 = péssima · 5 = ótima"'),
        ("sau_humor", "F", "G", "Humor médio", f"={avg('F')}", '0.0"/5"',
         f'=IF(ISNUMBER(F6),IF(F6>=3.5,"✔ ",IF(F6>=2.5,"▲ ","✖ ")),"")&"Energia "&IFERROR(FIXED(AVERAGEIFS({rg("G")},{in_month(Bd)}),1),"–")&" · Estresse "&IFERROR(FIXED(AVERAGEIFS({rg("H")},{in_month(Bd)}),1),"–")'),
        ("sau_treinos", "H", "I", "Treinos no mês", f'=COUNTIFS({rg("I")},"?*",{in_month(Bd)})', "0",
         f'=IF(DiasRef=0,"○ Mês ainda não começou",IF(H6>=MetaTreinosSem*DiasRef/7,"✔ ","▲ ")&"meta "&ROUND(MetaTreinosSem*DiasRef/7,0)&" · "&FIXED(SUMIFS({rg("J")},{in_month(Bd)}),0)&" min")'),
        ("sau_passos", "J", "K", "Passos por dia", f"={avg('K')}", "#,##0",
         '=IF(ISNUMBER(J6),IF(J6>=MetaPassos,"✔ ","▲ ")&"meta "&FIXED(MetaPassos,0),"○ Sem registros")'),
        ("sau_peso", "L", "N", "Peso atual",
         f'=IF({A(S.CALC,"C",151)}=0,"–",INDEX({rg("L")},{A(S.CALC,"C",151)}))', '0.0" kg"',
         f'=IF({A(S.CALC,"C",152)}=0,"○ Sem pesagens","Início do ano "&FIXED(INDEX({rg("L")},{A(S.CALC,"C",152)}),1)&IF(ISNUMBER(PesoAlvo)," · alvo "&FIXED(PesoAlvo,1),""))'),
        ("sau_dias", "O", "O", "Dias com registro no mês",
         f'=SUMPRODUCT(({Bd}>=RefIni)*({Bd}<=RefFim)*((({rg("D")}<>"")+({rg("F")}<>"")+({rg("I")}<>""))>0))', "0",
         '=IF(DiasRef=0,"","de "&DiasRef&" dias já decorridos · "&IFERROR(ROUND(O6/DiasRef*100,0),0)&"%")'),
    ])
    headers(ws, 10, [("B", "Data", "auto"), ("C", "Dia", "auto"), ("D", "Sono (h)", "in"),
                     ("E", "Qualid. sono (1–5)", "in"), ("F", "Humor (1–5)", "in"), ("G", "Energia (1–5)", "in"),
                     ("H", "Estresse (1–5)", "in"), ("I", "Treino", "in"), ("J", "Treino (min)", "in"),
                     ("K", "Passos", "in"), ("L", "Peso (kg)", "in"), ("M", "Água (L)", "in"),
                     ("N", "Alimentação (1–5)", "in"), ("O", "Notas · como me senti, o que marcou o dia", "in")], 34)
    body(ws, "B", "C", r1, r2, "auto", h="center")
    ws[f"B{r1}"] = "=DATE(Ano,1,1)"
    formulas(ws, "B", r1 + 1, r2, lambda r: f'=IF(B{r-1}="","",IF(YEAR(B{r-1}+1)=Ano,B{r-1}+1,""))')
    formulas(ws, "C", r1, r2, lambda r: f'=IF(B{r}="","",CHOOSE(WEEKDAY(B{r}),"Dom","Seg","Ter","Qua","Qui","Sex","Sáb"))')
    for r in range(r1, r2 + 1):
        ws[f"B{r}"].number_format = DATE
    body(ws, "D", "H", r1, r2, h="center", fmt="0.0")
    for col in "EFGH":
        body(ws, col, col, r1, r2, h="center", fmt="0")
    body(ws, "I", "I", r1, r2)
    body(ws, "J", "J", r1, r2, h="center", fmt="0")
    body(ws, "K", "K", r1, r2, h="center", fmt="#,##0")
    body(ws, "L", "M", r1, r2, h="center", fmt="0.0")
    body(ws, "N", "N", r1, r2, h="center", fmt="0")
    body(ws, "O", "O", r1, r2)
    dv_num(ws, f"D{r1}:D{r2}", 0, 24, prompt="Horas dormidas (ex.: 7,5)")
    dv_num(ws, f"E{r1}:G{r2}", 1, 5, True, prompt="1 = muito ruim · 3 = normal · 5 = excelente")
    dv_num(ws, f"H{r1}:H{r2}", 1, 5, True, prompt="1 = tranquilo · 5 = muito estressado")
    dv_num(ws, f"N{r1}:N{r2}", 1, 5, True, prompt="1 = muito desregrada · 5 = muito boa")
    dv_list(ws, f"I{r1}:I{r2}", "lst_Treino", "Deixe vazio se não treinou")
    dv_num(ws, f"J{r1}:J{r2}", 0, 600)
    dv_num(ws, f"K{r1}:K{r2}", 0, 100000, True)
    dv_num(ws, f"L{r1}:L{r2}", 20, 300)
    dv_num(ws, f"M{r1}:M{r2}", 0, 15)
    cf_diverge(ws, f"E{r1}:G{r2}", 1, 3, 5)
    cf_diverge(ws, f"N{r1}:N{r2}", 1, 3, 5)
    cf_diverge(ws, f"H{r1}:H{r2}", 1, 3, 5, reverse=True)
    ws.conditional_formatting.add(f"B{r1}:O{r2}", FormulaRule(formula=[f"$B{r1}=Hoje"], fill=fill(TODAY_BG),
                                                               font=Font(bold=True)))
    ws.conditional_formatting.add(f"B{r1}:C{r2}", FormulaRule(formula=[f'AND($B{r1}<>"",WEEKDAY($B{r1},2)>5)'],
                                                               font=Font(color=MUTED)))
    ws.freeze_panes = "D11"

    # --- resumo mensal (alimenta gráficos)
    section(ws, 9, "Q", "Z", "RESUMO POR MÊS")
    headers(ws, 10, [("Q", "Mês", "auto"), ("R", "Sono (h)", "auto"), ("S", "Humor", "auto"), ("T", "Energia", "auto"),
                     ("U", "Estresse", "auto"), ("V", "Treinos", "auto"), ("W", "Min. treino", "auto"),
                     ("X", "Passos/dia", "auto"), ("Y", "Peso (kg)", "auto"), ("Z", "Dias c/ registro", "auto")], 34)
    body(ws, "Q", "Z", 11, 23, "auto", h="center", fmt='0.0;;"–"')
    for i in range(12):
        r, m = 11 + i, i + 1
        ini, fim = month_bounds(m)
        mm = f'{Bd},">="&{ini},{Bd},"<="&{fim}'
        ws[f"Q{r}"] = MESES_ABR[i]
        for col, src in (("R", "D"), ("S", "F"), ("T", "G"), ("U", "H"), ("X", "K"), ("Y", "L")):
            ws[f"{col}{r}"] = f'=IFERROR(AVERAGEIFS({rg(src)},{mm}),"")'
        ws[f"V{r}"] = f'=COUNTIFS({rg("I")},"?*",{mm})'
        ws[f"W{r}"] = f'=SUMIFS({rg("J")},{mm})'
        ws[f"Z{r}"] = (f'=SUMPRODUCT(({Bd}>={ini})*({Bd}<={fim})*'
                       f'((({rg("D")}<>"")+({rg("F")}<>"")+({rg("I")}<>""))>0))')
        for col in "VWZ":
            ws[f"{col}{r}"].number_format = '0;;"–"'
        ws[f"X{r}"].number_format = '#,##0;;"–"'
    ws["Q23"] = "Ano"
    for col in "RSTUXY":
        src = {"R": "D", "S": "F", "T": "G", "U": "H", "X": "K", "Y": "L"}[col]
        ws[f"{col}23"] = f'=IFERROR(AVERAGE({rg(src)}),"")'
    for col in "VWZ":
        ws[f"{col}23"] = f"=SUM({col}11:{col}22)"
        ws[f"{col}23"].number_format = '#,##0;;"–"'
    ws["X23"].number_format = '#,##0;;"–"'
    for c in range(CI("Q"), CI("Z") + 1):
        ws.cell(23, c).font = F(9, True, INK)
    ws.conditional_formatting.add("Q11:Z22", FormulaRule(formula=["ROW()-10=RefMes"], fill=fill(TODAY_BG)))

    # --- consultas
    section(ws, 25, "Q", "Z", "CONSULTAS, EXAMES & CHECK-UPS", height=None)
    headers(ws, 26, [("Q", "Data", "in"), ("R", "Tipo", "in"), ("S", "Profissional / exame", "in", "T"),
                     ("U", "Status", "in"), ("V", "Próxima", "in"), ("W", "Faltam", "auto"),
                     ("X", "Notas", "in", "Z")], None)
    c1, c2 = 27, 56
    body(ws, "Q", "Q", c1, c2, fmt=DATE, h="center")
    body(ws, "R", "U", c1, c2)
    body(ws, "V", "V", c1, c2, fmt=DATE, h="center")
    body(ws, "W", "W", c1, c2, "auto", fmt='0;[Red]-0', h="center")
    body(ws, "X", "Z", c1, c2)
    merge_rows(ws, "S", "T", c1, c2)
    merge_rows(ws, "X", "Z", c1, c2)
    formulas(ws, "W", c1, c2, lambda r: f'=IF(V{r}="","",V{r}-Hoje)')
    dv_date(ws, f"Q{c1}:Q{c2}")
    dv_date(ws, f"V{c1}:V{c2}", "Quando é a próxima (retorno, novo exame)")
    dv_list(ws, f"R{c1}:R{c2}", "lst_TipoConsulta")
    dv_list(ws, f"U{c1}:U{c2}", "lst_StatusConsulta")
    ws.conditional_formatting.add(f"W{c1}:W{c2}", FormulaRule(formula=[f'AND(ISNUMBER(W{c1}),W{c1}<=30)'],
                                                               font=Font(color=CRIT_T, bold=True)))
    V = A(S.SAU, "V", c1, "V", c2)
    K["sau_prox_consulta"] = f'SUMPRODUCT(MIN(({V}<>"")*({V}>=Hoje)*{V}+(1-({V}<>"")*({V}>=Hoje))*99999))'


# =============================================================================
# Hábitos (grade anual)
# =============================================================================
HAB_R1, HAB_R2 = 17, 382
HAB_COLS = [CL(4 + i) for i in range(10)]          # D..M
HAB_HELP = [CL(17 + i) for i in range(10)]         # Q..Z (ocultas)


def build_habitos(ws):
    sheet_title(ws, "Hábitos",
                "Escreva o nome dos hábitos na linha azul e marque x no dia em que cumprir. Metas por semana definem o que é 'no alvo'.", "O")
    setw(ws, {"B": 12, "C": 9, **{c: 11.5 for c in HAB_COLS}, "N": 8, "O": 9, "P": 3})
    for c in HAB_HELP:
        ws.column_dimensions[c].hidden = True
    r1, r2 = HAB_R1, HAB_R2
    c = ws["M3"]
    c.value = f'=HYPERLINK("#\'{S.HAB}\'!B"&MAX({r1},IdxHoje+{r1 - 1}),"→ Ir para a linha de hoje")'
    c.font = F(9, True, BLUE, underline="single")
    section(ws, 5, "B", "O", None)
    ws["B5"] = '="RESUMO · "&UPPER(RotuloMes)'
    labels = [(6, "Área da vida", "in"), (7, "Meta (dias por semana)", "in"), (8, "Feitos no mês", "auto"),
              (9, "% dos dias do mês", "auto"), (10, "% da meta do mês", "auto"), (11, "Sequência atual (dias)", "auto"),
              (12, "Melhor sequência no ano", "auto"), (13, "Feitos no ano", "auto"), (14, "Situação", "auto")]
    for r, label, kind in labels:
        merge(ws, f"B{r}:C{r}", label, F(9, True, INK), fill(WHITE if kind == "in" else AUTO_BG), al(indent=1), BOX)
        body(ws, "D", "M", r, r, kind, h="center", wrap=(r == 6))
        body(ws, "N", "O", r, r, "auto", h="center")
        ws.row_dimensions[r].height = 26 if r == 6 else 18
    merge(ws, "N6:O6", "Geral", F(9, True, INK), fill(AUTO_BG), al("center"), BOX)
    hdr = "$D$16:$M$16"
    for i, col in enumerate(HAB_COLS):
        h = HAB_HELP[i]
        g = f"{col}$16"
        rng = f"{col}${r1}:{col}${r2}"
        ws[f"{col}8"] = f'=IF({g}="","",COUNTIFS({rng},"x",$B${r1}:$B${r2},">="&RefIni,$B${r1}:$B${r2},"<="&RefFim))'
        ws[f"{col}9"] = f'=IF(OR({g}="",DiasRef=0),"",{col}8/DiasRef)'
        ws[f"{col}10"] = f'=IF(OR({g}="",DiasRef=0,N({col}7)=0),"",MIN(1,{col}8/({col}7*DiasRef/7)))'
        ws[f"{col}11"] = (f'=IF({g}="","",IF(IdxHoje=0,0,IF(INDEX({rng},IdxHoje)="x",INDEX({h}${r1}:{h}${r2},IdxHoje),'
                          f'IF(IdxHoje>1,INDEX({h}${r1}:{h}${r2},IdxHoje-1),0))))')
        ws[f"{col}12"] = f'=IF({g}="","",MAX({h}${r1}:{h}${r2}))'
        ws[f"{col}13"] = f'=IF({g}="","",COUNTIF({rng},"x"))'
        ws[f"{col}14"] = f'=IF({col}10="","",IF({col}10>=0.9,"✔ No alvo",IF({col}10>=0.6,"▲ Quase","✖ Abaixo")))'
        ws[f"{col}9"].number_format = ws[f"{col}10"].number_format = PCT
        ws[f"{h}{r1}"] = f'=IF({col}{r1}="x",1,0)'
        formulas(ws, h, r1 + 1, r2, lambda r, col=col, h=h: f'=IF({col}{r}="x",{h}{r-1}+1,0)')
    merge(ws, "N8:O8", "=SUM(D8:M8)", fill_=fill(AUTO_BG), align=al("center"), border=BOX)
    merge(ws, "N9:O9", f'=IF(OR(COUNTA({hdr})=0,DiasRef=0),"",N8/(COUNTA({hdr})*DiasRef))', fill_=fill(AUTO_BG),
          align=al("center"), border=BOX, fmt=PCT)
    merge(ws, "N10:O10", '=IFERROR(AVERAGE(D10:M10),"")', fill_=fill(AUTO_BG), align=al("center"), border=BOX, fmt=PCT)
    merge(ws, "N11:O11", '=IF(COUNT(D11:M11)=0,"",MAX(D11:M11))', fill_=fill(AUTO_BG), align=al("center"), border=BOX)
    merge(ws, "N12:O12", '=IF(COUNT(D12:M12)=0,"",MAX(D12:M12))', fill_=fill(AUTO_BG), align=al("center"), border=BOX)
    merge(ws, "N13:O13", "=SUM(D13:M13)", fill_=fill(AUTO_BG), align=al("center"), border=BOX)
    merge(ws, "N14:O14", '=IF(N10="","",IF(N10>=0.9,"✔ No alvo",IF(N10>=0.6,"▲ Quase","✖ Abaixo")))',
          fill_=fill(AUTO_BG), align=al("center"), border=BOX)
    for r in range(8, 15):
        for c in range(CI("D"), CI("O") + 1):
            ws.cell(r, c).font = F(9, r in (10, 14), INK)
    cf_bar(ws, "D10:O10", 0, 1, ORANGE)
    cf_status(ws, "D14:O14")
    dv_list(ws, "D6:M6", "lst_Areas", "A qual área da vida este hábito serve? (alimenta o placar da área)")
    dv_num(ws, "D7:M7", 1, 7, True, prompt="Quantos dias por semana você quer cumprir este hábito")
    ws.row_dimensions[15].height = 8
    headers(ws, 16, [("B", "Data", "auto"), ("C", "Dia", "auto")] + [(c, "", "edit") for c in HAB_COLS] +
            [("N", "Feitos", "auto"), ("O", "% do dia", "auto")], 36)
    for c in HAB_COLS:
        ws[f"{c}16"].border = Border(left=side(BLUE), right=side(BLUE), top=side(BLUE), bottom=side(BLUE))
    body(ws, "B", "C", r1, r2, "auto", h="center")
    ws[f"B{r1}"] = "=DATE(Ano,1,1)"
    formulas(ws, "B", r1 + 1, r2, lambda r: f'=IF(B{r-1}="","",IF(YEAR(B{r-1}+1)=Ano,B{r-1}+1,""))')
    formulas(ws, "C", r1, r2, lambda r: f'=IF(B{r}="","",CHOOSE(WEEKDAY(B{r}),"Dom","Seg","Ter","Qua","Qui","Sex","Sáb"))')
    for r in range(r1, r2 + 1):
        ws[f"B{r}"].number_format = DATE
    body(ws, "D", "M", r1, r2, h="center")
    body(ws, "N", "O", r1, r2, "auto", h="center")
    formulas(ws, "N", r1, r2, lambda r: f'=IF(OR(B{r}="",B{r}>Hoje),"",COUNTIF(D{r}:M{r},"x"))')
    formulas(ws, "O", r1, r2, lambda r: f'=IF(OR(N{r}="",COUNTA({hdr})=0),"",N{r}/COUNTA({hdr}))')
    for r in range(r1, r2 + 1):
        ws[f"O{r}"].number_format = PCT
    dv = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dv.showInputMessage = True
    dv.prompt = "Digite x quando cumprir o hábito neste dia"
    dv.showErrorMessage = True
    dv.errorStyle = "warning"
    dv.error = "Use x para marcar o hábito como feito."
    ws.add_data_validation(dv)
    dv.add(f"D{r1}:M{r2}")
    ws.conditional_formatting.add(f"D{r1}:M{r2}", FormulaRule(formula=[f'D{r1}="x"'], fill=fill(ORANGE),
                                                               font=Font(color=WHITE, bold=True)))
    ws.conditional_formatting.add(f"B{r1}:O{r2}", FormulaRule(formula=[f"$B{r1}=Hoje"], fill=fill(TODAY_BG),
                                                               font=Font(bold=True)))
    ws.conditional_formatting.add(f"B{r1}:C{r2}", FormulaRule(formula=[f'AND($B{r1}<>"",WEEKDAY($B{r1},2)>5)'],
                                                               font=Font(color=MUTED)))
    cf_bar(ws, f"O{r1}:O{r2}", 0, 1, ORANGE)
    ws.freeze_panes = "D17"
    K["hab_meta"] = A(S.HAB, "N", 10)
    K["hab_dias"] = A(S.HAB, "N", 9)


# =============================================================================
# Metas
# =============================================================================
MET_R1, MET_R2 = 15, 74


def build_metas(ws):
    sheet_title(ws, "Metas · curto, médio e longo prazo",
                "Visão no topo, metas mensuráveis embaixo. Progresso, ritmo esperado e horizonte são calculados.", "T")
    setw(ws, {"B": 5, "C": 42, "D": 20, "E": 12, "F": 11, "G": 11, "H": 10, "I": 9, "J": 9, "K": 9, "L": 10,
              "M": 11, "N": 11, "O": 15, "P": 9, "Q": 10, "R": 9, "S": 34, "T": 34})
    for c in "UVW":
        ws.column_dimensions[c].hidden = True
    vis = [(5, "Minha visão (daqui a 5 anos)"), (6, "Propósito · por que faço o que faço"),
           (7, "Valores que não negocio"), (8, "Tema / palavra do ano")]
    for r, label in vis:
        merge(ws, f"B{r}:C{r}", label, F(9, True, INK), fill(AUTO_BG), al(indent=1, wrap=True), BOX)
        merge(ws, f"D{r}:T{r}", None, F(10, False, INK, italic=True), fill(WHITE), al(wrap=True, indent=1), BOX)
        ws.row_dimensions[r].height = 30
    r1, r2 = MET_R1, MET_R2
    O, U, V, M = (f"${c}${r1}:${c}${r2}" for c in "OUVM")
    kpi_strip(ws, 10, [
        ("met_and", "C", "C", "Metas em andamento",
         f'=COUNTIF({O},"✔ No ritmo")+COUNTIF({O},"▲*")+COUNTIF({O},"✖*")+COUNTIF({O},"○*")', "0",
         f'="Concluídas: "&COUNTIF({O},"✔ Concluída")&" · pausadas: "&COUNTIF({O},"‖*")'),
        ("met_ritmo", "D", "E", "No ritmo", f'=COUNTIF({O},"✔ No ritmo")', "0",
         '=IF(C11=0,"",IF(D11/C11>=0.7,"✔ ","▲ ")&ROUND(D11/C11*100,0)&"% das em andamento")'),
        ("met_atencao", "F", "H", "Em atenção", f'=COUNTIF({O},"▲*")', "0", '="progresso um pouco atrás do tempo"'),
        ("met_atras", "I", "K", "Atrasadas ou vencidas", f'=COUNTIF({O},"✖*")', "0",
         '=IF(I11>0,"✖ replaneje ou ajuste o prazo","✔ nenhuma")'),
        ("met_prog", "L", "N", "Progresso médio (em andamento)",
         f'=IFERROR(AVERAGEIFS({M},{U},1,{O},"<>✔ Concluída"),"–")', PCT, '="vs. tempo decorrido de cada meta"'),
        ("met_hor", "O", "T", "Por horizonte",
         f'="Curto "&COUNTIFS($Q${r1}:$Q${r2},"Curto",{U},1)&" · Médio "&COUNTIFS($Q${r1}:$Q${r2},"Médio",{U},1)&" · Longo "&COUNTIFS($Q${r1}:$Q${r2},"Longo",{U},1)',
         None, '="pela duração: curto ≤ 6 meses · médio ≤ 2 anos · longo > 2 anos"'),
    ])
    ws["O11"].font = F(12, True, INK)
    headers(ws, 14, [("B", "#", "auto"), ("C", "Meta (específica e mensurável)", "in"), ("D", "Área da vida", "in"),
                     ("E", "Status", "in"), ("F", "Início", "in"), ("G", "Prazo", "in"), ("H", "Unidade", "in"),
                     ("I", "Valor inicial", "in"), ("J", "Valor atual", "in"), ("K", "Valor alvo", "in"),
                     ("L", "Progresso manual", "in"), ("M", "Progresso", "auto"), ("N", "Esperado hoje", "auto"),
                     ("O", "Ritmo", "auto"), ("P", "Dias restantes", "auto"), ("Q", "Horizonte", "auto"),
                     ("R", "Tarefas abertas", "auto"), ("S", "Próximo passo concreto", "in"),
                     ("T", "Por que importa / notas", "in")], 36)
    body(ws, "B", "B", r1, r2, "auto", h="center")
    body(ws, "C", "E", r1, r2)
    body(ws, "F", "G", r1, r2, fmt=DATE, h="center")
    body(ws, "H", "K", r1, r2, h="center", fmt="#,##0.##")
    body(ws, "L", "L", r1, r2, h="center", fmt=PCT)
    body(ws, "M", "N", r1, r2, "auto", h="center", fmt=PCT)
    body(ws, "O", "R", r1, r2, "auto", h="center")
    body(ws, "S", "T", r1, r2, wrap=True)
    T_meta, T_st = A(S.TAR, "I", 34, "I", 333), A(S.TAR, "G", 34, "G", 333)
    formulas(ws, "B", r1, r2, lambda r: f'=IF(C{r}="","",ROW()-{r1 - 1})')
    formulas(ws, "M", r1, r2, lambda r: (
        f'=IF(C{r}="","",IF(E{r}="Concluída",1,IF(AND(ISNUMBER(I{r}),ISNUMBER(J{r}),ISNUMBER(K{r}),K{r}<>I{r}),'
        f'MAX(0,MIN(1,(J{r}-I{r})/(K{r}-I{r}))),IF(ISNUMBER(L{r}),MAX(0,MIN(1,L{r})),0))))'))
    formulas(ws, "N", r1, r2, lambda r: (
        f'=IF(OR(C{r}="",F{r}="",G{r}=""),"",IF(G{r}<=F{r},1,MAX(0,MIN(1,(Hoje-F{r})/(G{r}-F{r})))))'))
    formulas(ws, "O", r1, r2, lambda r: (
        f'=IF(C{r}="","",IF(OR(E{r}="Concluída",M{r}>=1),"✔ Concluída",IF(E{r}="Pausada","‖ Pausada",'
        f'IF(E{r}="Abandonada","– Abandonada",IF(N{r}="","○ Sem prazo",IF(G{r}<Hoje,"✖ Vencida",'
        f'IF(M{r}>=N{r}-0.1,"✔ No ritmo",IF(M{r}>=N{r}-0.25,"▲ Atenção","✖ Atrasada"))))))))'))
    formulas(ws, "P", r1, r2, lambda r: f'=IF(OR(C{r}="",G{r}="",O{r}="✔ Concluída"),"",G{r}-Hoje)')
    formulas(ws, "Q", r1, r2, lambda r: (
        f'=IF(OR(C{r}="",G{r}=""),"",IF(G{r}-IF(F{r}="",Hoje,F{r})<=183,"Curto",'
        f'IF(G{r}-IF(F{r}="",Hoje,F{r})<=731,"Médio","Longo")))'))
    formulas(ws, "R", r1, r2, lambda r: (
        f'=IF(C{r}="","",COUNTIFS({T_meta},C{r},{T_st},"<>Concluída",{T_st},"<>Cancelada"))'))
    formulas(ws, "U", r1, r2, lambda r: f'=IF(C{r}="",0,IF(OR(E{r}="Pausada",E{r}="Abandonada"),0,1))')
    formulas(ws, "V", r1, r2, lambda r: (
        f'=IF(U{r}=0,0,IF(OR(O{r}="✔ Concluída",O{r}="✔ No ritmo",O{r}="○ Sem prazo"),1,0))'))
    formulas(ws, "W", r1, r2, lambda r: f'=IF(AND(U{r}=1,O{r}<>"✔ Concluída",G{r}<>""),G{r}+ROW()/100000,"")')
    for r in range(r1, r2 + 1):
        ws[f"P{r}"].number_format = '0;[Red]-0'
        ws.row_dimensions[r].height = 30
    dv_list(ws, f"D{r1}:D{r2}", "lst_Areas")
    dv_list(ws, f"E{r1}:E{r2}", "lst_StatusMeta", "Deixe 'Ativa' enquanto estiver perseguindo a meta")
    dv_date(ws, f"F{r1}:G{r2}")
    dv_num(ws, f"L{r1}:L{r2}", 0, 1, prompt="Só para metas sem número (marcos). Ex.: 40%. Se preencher inicial/atual/alvo, este campo é ignorado.")
    cf_bar(ws, f"M{r1}:M{r2}", 0, 1, BLUE)
    cf_status(ws, f"O{r1}:O{r2}")
    ws.conditional_formatting.add(f"C{r1}:C{r2}", FormulaRule(formula=[f'$O{r1}="✔ Concluída"'],
                                                               font=Font(color=MUTED, strike=True)))
    ws.auto_filter.ref = f"B14:T{r2}"
    ws.freeze_panes = "D15"


# =============================================================================
# Tarefas & projetos
# =============================================================================
PRJ_R1, PRJ_R2 = 11, 30
TSK_R1, TSK_R2 = 34, 333


def build_tarefas(ws):
    sheet_title(ws, "Tarefas & projetos",
                "Projetos agrupam tarefas; tarefas podem apontar para uma meta. Alertas e prioridade (matriz de Eisenhower) são automáticos.", "M")
    setw(ws, {"B": 42, "C": 22, "D": 18, "E": 12, "F": 12, "G": 14, "H": 13, "I": 26, "J": 9, "K": 17,
              "L": 18, "M": 36})
    ws.column_dimensions["N"].hidden = True
    p1, p2, t1, t2 = PRJ_R1, PRJ_R2, TSK_R1, TSK_R2
    tr = lambda c: f"${c}${t1}:${c}${t2}"
    kpi_strip(ws, 5, [
        ("tar_abertas", "B", "B", "Tarefas abertas",
         f'=COUNTIFS({tr("B")},"?*",{tr("G")},"<>Concluída",{tr("G")},"<>Cancelada")', "0",
         f'="Alta prioridade: "&COUNTIFS({tr("E")},"Alta",{tr("K")},"<>✔*",{tr("K")},"<>–*",{tr("B")},"?*")'),
        ("tar_atras", "C", "D", "Atrasadas", f'=COUNTIF({tr("K")},"✖*")', "0",
         '=IF(C6>0,"✖ resolva ou replaneje","✔ nada atrasado")'),
        ("tar_7d", "E", "G", "Vencem em até 7 dias", f'=COUNTIF({tr("K")},"▲*")', "0", '="inclui as que vencem hoje"'),
        ("tar_concl", "H", "I", "Concluídas no mês",
         f'=COUNTIFS({tr("G")},"Concluída",{tr("H")},">="&RefIni,{tr("H")},"<="&RefFim)', "0", '="Mês: "&RotuloMes'),
        ("tar_prazo", "J", "L", "Entregues no prazo (mês)",
         f'=IFERROR(SUMPRODUCT(({tr("G")}="Concluída")*({tr("H")}>=RefIni)*({tr("H")}<=RefFim)*({tr("F")}<>"")*({tr("H")}<={tr("F")}))'
         f'/SUMPRODUCT(({tr("G")}="Concluída")*({tr("H")}>=RefIni)*({tr("H")}<=RefFim)*({tr("F")}<>"")),"–")', PCT,
         '=IF(ISNUMBER(J6),IF(J6>=0.8,"✔ ","▲ ")&"das concluídas com prazo","○ Nenhuma concluída com prazo")'),
        ("tar_proj", "M", "M", "Projetos em andamento", f'=COUNTIF($D${p1}:$D${p2},"Em andamento")', "0",
         f'="Ideias guardadas: "&COUNTIF($D${p1}:$D${p2},"Ideia")'),
    ])
    section(ws, 9, "B", "M", "PROJETOS · inclui o banco de ideias (status 'Ideia')")
    headers(ws, 10, [("B", "Projeto", "in"), ("C", "Área da vida", "in"), ("D", "Status", "in"), ("E", "Início", "in"),
                     ("F", "Prazo", "in"), ("G", "Tarefas", "auto"), ("H", "Concluídas", "auto"),
                     ("I", "Meta vinculada", "in"), ("J", "%", "auto"), ("K", "Atrasadas", "auto"),
                     ("L", "Próxima entrega", "auto"), ("M", "Descrição / notas", "in")])
    body(ws, "B", "D", p1, p2)
    body(ws, "E", "F", p1, p2, fmt=DATE, h="center")
    body(ws, "G", "H", p1, p2, "auto", h="center")
    body(ws, "I", "I", p1, p2)
    body(ws, "J", "J", p1, p2, "auto", fmt=PCT, h="center")
    body(ws, "K", "K", p1, p2, "auto", h="center", fmt='0;;"–"')
    body(ws, "L", "L", p1, p2, "auto", h="center", fmt=DATE)
    body(ws, "M", "M", p1, p2, wrap=True)
    canc = lambda r: f'COUNTIFS({tr("C")},B{r},{tr("G")},"Cancelada")'
    formulas(ws, "G", p1, p2, lambda r: f'=IF(B{r}="","",COUNTIF({tr("C")},B{r}))')
    formulas(ws, "H", p1, p2, lambda r: f'=IF(B{r}="","",COUNTIFS({tr("C")},B{r},{tr("G")},"Concluída"))')
    formulas(ws, "J", p1, p2, lambda r: (
        f'=IF(B{r}="","",IF(D{r}="Concluído",1,IF(G{r}-{canc(r)}<=0,0,H{r}/(G{r}-{canc(r)}))))'))
    formulas(ws, "K", p1, p2, lambda r: f'=IF(B{r}="","",COUNTIFS({tr("C")},B{r},{tr("K")},"✖*"))')
    cond = lambda r: (f'(({tr("C")}=B{r})*({tr("G")}<>"Concluída")*({tr("G")}<>"Cancelada")*({tr("F")}<>""))')
    formulas(ws, "L", p1, p2, lambda r: (
        f'=IF(B{r}="","",IF({smin(cond(r), tr("F"))}>=99999,"—",{smin(cond(r), tr("F"))}))'))
    dv_list(ws, f"C{p1}:C{p2}", "lst_Areas")
    dv_list(ws, f"D{p1}:D{p2}", "lst_StatusProjeto")
    dv_date(ws, f"E{p1}:F{p2}")
    dv_list(ws, f"I{p1}:I{p2}", "lst_Metas", strict=False)
    cf_bar(ws, f"J{p1}:J{p2}", 0, 1, BLUE)
    ws.conditional_formatting.add(f"K{p1}:K{p2}", FormulaRule(formula=[f"N(K{p1})>0"], font=Font(color=CRIT_T, bold=True)))
    ws.conditional_formatting.add(f"B{p1}:B{p2}", FormulaRule(formula=[f'$D{p1}="Ideia"'], font=Font(color=VIOLET, italic=True)))

    section(ws, 32, "B", "M", "TAREFAS")
    headers(ws, 33, [("B", "Tarefa", "in"), ("C", "Projeto", "in"), ("D", "Área da vida", "in"),
                     ("E", "Prioridade", "in"), ("F", "Prazo", "in"), ("G", "Status", "in"),
                     ("H", "Concluída em", "in"), ("I", "Meta vinculada", "in"), ("J", "Dias p/ prazo", "auto"),
                     ("K", "Alerta", "auto"), ("L", "Quadrante", "auto"), ("M", "Notas", "in")])
    body(ws, "B", "E", t1, t2)
    body(ws, "F", "F", t1, t2, fmt=DATE, h="center")
    body(ws, "G", "G", t1, t2)
    body(ws, "H", "H", t1, t2, fmt=DATE, h="center")
    body(ws, "I", "I", t1, t2)
    body(ws, "J", "J", t1, t2, "auto", fmt='0;[Red]-0', h="center")
    body(ws, "K", "L", t1, t2, "auto")
    body(ws, "M", "M", t1, t2)
    formulas(ws, "J", t1, t2, lambda r: f'=IF(OR(B{r}="",F{r}="",G{r}="Concluída",G{r}="Cancelada"),"",F{r}-Hoje)')
    formulas(ws, "K", t1, t2, lambda r: (
        f'=IF(B{r}="","",IF(G{r}="Concluída","✔ Feita",IF(G{r}="Cancelada","– Cancelada",IF(F{r}="","○ Sem prazo",'
        f'IF(F{r}<Hoje,"✖ Atrasada",IF(F{r}=Hoje,"▲ Vence hoje",IF(F{r}-Hoje<=7,"▲ Próximos 7 dias","○ No prazo")))))))'))
    formulas(ws, "L", t1, t2, lambda r: (
        f'=IF(OR(B{r}="",G{r}="Concluída",G{r}="Cancelada"),"",IF(E{r}="Alta",IF(AND(F{r}<>"",F{r}-Hoje<=3),'
        f'"1 · Fazer já","2 · Agendar"),IF(AND(F{r}<>"",F{r}-Hoje<=3),"3 · Resolver rápido","4 · Depois")))'))
    formulas(ws, "N", t1, t2, lambda r: f'=IF(AND(B{r}<>"",G{r}<>"Concluída",G{r}<>"Cancelada",F{r}<>""),F{r}+ROW()/100000,"")')
    dv_list(ws, f"C{t1}:C{t2}", "lst_Projetos", "Projetos cadastrados acima (opcional)", strict=False)
    dv_list(ws, f"D{t1}:D{t2}", "lst_Areas")
    dv_list(ws, f"E{t1}:E{t2}", "lst_Prioridade")
    dv_date(ws, f"F{t1}:F{t2}")
    dv_list(ws, f"G{t1}:G{t2}", "lst_StatusTarefa", "Ao concluir, preencha também 'Concluída em'")
    dv_date(ws, f"H{t1}:H{t2}", "Data em que terminou (Ctrl+; = hoje)")
    dv_list(ws, f"I{t1}:I{t2}", "lst_Metas", "Opcional: liga a tarefa a uma meta", strict=False)
    cf_status(ws, f"K{t1}:K{t2}")
    ws.conditional_formatting.add(f"B{t1}:B{t2}", FormulaRule(formula=[f'OR($G{t1}="Concluída",$G{t1}="Cancelada")'],
                                                               font=Font(color=MUTED, strike=True)))
    ws.conditional_formatting.add(f"H{t1}:H{t2}", FormulaRule(formula=[f'AND($G{t1}="Concluída",$H{t1}="")'],
                                                               fill=fill(WARN_BG)))
    ws.conditional_formatting.add(f"L{t1}:L{t2}", FormulaRule(formula=[f'LEFT($L{t1},1)="1"'],
                                                               font=Font(color=CRIT_T, bold=True)))
    ws.conditional_formatting.add(f"E{t1}:E{t2}", FormulaRule(formula=[f'$E{t1}="Alta"'], font=Font(bold=True)))
    cf_zebra(ws, f"B{t1}:M{t2}", "B")
    ws.auto_filter.ref = f"B33:M{t2}"
    ws.freeze_panes = "C5"


# =============================================================================
# Carreira
# =============================================================================
CMP_R1, CMP_R2 = 15, 29
OPP_R1, OPP_R2 = 33, 62
WIN_R1, WIN_R2 = 66, 125


def build_carreira(ws):
    sheet_title(ws, "Carreira & desenvolvimento profissional",
                "Onde você está, que competências quer subir de nível, que oportunidades persegue e o que já entregou.", "L")
    setw(ws, {"B": 34, "C": 16, "D": 12, "E": 12, "F": 11, "G": 16, "H": 30, "I": 13, "J": 14, "K": 13, "L": 30})
    section(ws, 5, "B", "E", "SITUAÇÃO ATUAL")
    section(ws, 5, "G", "L", "INDICADORES (automático)")
    prof = [(6, "Cargo atual"), (7, "Empresa"), (8, "Desde"), (9, "Remuneração líquida / mês (€)"),
            (10, "Satisfação com o trabalho (0–10)"), (11, "Objetivo para os próximos 12 meses"),
            (12, "Visão de carreira (3–5 anos)")]
    for r, label in prof:
        merge(ws, f"B{r}", label, F(9, True, INK), fill(AUTO_BG), al(indent=1), BOX)
        merge(ws, f"C{r}:E{r}", None, F(10, False, INK), fill(WHITE), al(indent=1, wrap=True), BOX)
        ws.row_dimensions[r].height = 22
    ws.row_dimensions[11].height = ws.row_dimensions[12].height = 34
    ws["C8"].number_format = DATE
    ws["C9"].number_format = '€ #,##0'
    dv_date(ws, "C8")
    dv_num(ws, "C10", 0, 10, prompt="0 = péssimo · 10 = não trocaria por nada")
    E, G, I_, D = (f"${c}${OPP_R1}:${c}${OPP_R2}" for c in "EGID")
    act = "+".join(f'COUNTIF({E},"{e}")' for e in ("Aplicado", "Entrevista", "Teste / case", "Proposta"))
    adv = "+".join(f'COUNTIF({E},"{e}")' for e in ("Entrevista", "Teste / case", "Proposta"))
    ind = [(6, "Candidaturas ativas", f"={act}", "0", '="aplicado, entrevista, teste ou proposta"'),
           (7, "Em fase avançada", f"={adv}", "0", '="entrevista, teste ou proposta"'),
           (8, "Ações pendentes", f'=COUNTIF({G},"▲*")+COUNTIF({G},"✖*")', "0",
            f'=IF(COUNTIF({G},"✖*")>0,"✖ "&COUNTIF({G},"✖*")&" ação(ões) atrasada(s)","follow-ups e próximas ações")'),
           (9, "Conquistas registradas no ano",
            f'=COUNTIFS($D${WIN_R1}:$D${WIN_R2},">="&IniAno,$D${WIN_R1}:$D${WIN_R2},"<="&FimAno)', "0",
            '="seu \'brag document\' para avaliações e entrevistas"'),
           (10, "Competências: % do nível alvo", f'=IFERROR(AVERAGE($K${CMP_R1}:$K${CMP_R2}),"–")', PCT,
            '=IF(ISNUMBER(I10),IF(I10>=0.85,"✔ ","▲ ")&"média das competências mapeadas","○ mapeie suas competências")'),
           (11, "Satisfação com o trabalho", '=IF(C10="","–",C10)', '0"/10"',
            '=IF(ISNUMBER(I11),IF(I11>=7,"✔ satisfeito",IF(I11>=5,"▲ morno","✖ insatisfeito")),"")'),
           (12, "Tempo na empresa", '=IF(C8="","–",ROUND((Hoje-C8)/365.25,1))', '0.0" anos"', '=""')]
    for r, label, fml, fmt, sub in ind:
        merge(ws, f"G{r}:H{r}", label, F(9, True, INK), fill(AUTO_BG), al(indent=1), BOX)
        merge(ws, f"I{r}", fml, F(11, True, INK), fill(AUTO_BG), al("center"), BOX, fmt)
        merge(ws, f"J{r}:L{r}", sub, F(8, False, INK2), fill(AUTO_BG), al(indent=1), BOX)
    cf_status(ws, "J6:J12")
    K["car_cand"] = A(S.CAR, "I", 6)
    K["car_comp"] = A(S.CAR, "I", 10)
    K["car_sat"] = A(S.CAR, "I", 11)
    K["car_conq"] = A(S.CAR, "I", 9)

    r1, r2 = CMP_R1, CMP_R2
    section(ws, 13, "B", "L", "COMPETÊNCIAS · nível de 1 (iniciante) a 5 (referência)")
    headers(ws, 14, [("B", "Competência", "in"), ("C", "Categoria", "in"), ("D", "Nível atual", "in"),
                     ("E", "Nível alvo", "in"), ("F", "Gap", "auto"), ("G", "Visual", "auto"),
                     ("H", "Como vou desenvolver", "in"), ("I", "Prazo", "in"), ("J", "Prioridade", "in"),
                     ("K", "% do alvo", "auto"), ("L", "Notas", "in")])
    body(ws, "B", "C", r1, r2)
    body(ws, "D", "E", r1, r2, h="center", fmt="0")
    body(ws, "F", "G", r1, r2, "auto", h="center")
    body(ws, "H", "H", r1, r2, wrap=True)
    body(ws, "I", "I", r1, r2, fmt=DATE, h="center")
    body(ws, "J", "J", r1, r2)
    body(ws, "K", "K", r1, r2, "auto", fmt=PCT, h="center")
    body(ws, "L", "L", r1, r2)
    formulas(ws, "F", r1, r2, lambda r: f'=IF(OR(B{r}="",D{r}="",E{r}=""),"",E{r}-D{r})')
    formulas(ws, "G", r1, r2, lambda r: f'=IF(OR(B{r}="",D{r}="",E{r}=""),"",REPT("●",D{r})&REPT("○",MAX(0,E{r}-D{r})))')
    formulas(ws, "K", r1, r2, lambda r: f'=IF(OR(B{r}="",N(E{r})=0),"",MIN(1,D{r}/E{r}))')
    for r in range(r1, r2 + 1):
        ws[f"G{r}"].font = F(10, False, VIOLET)
    dv_list(ws, f"C{r1}:C{r2}", "lst_CatComp")
    dv_num(ws, f"D{r1}:E{r2}", 1, 5, True, prompt="1 iniciante · 2 básico · 3 autônomo · 4 avançado · 5 referência")
    dv_date(ws, f"I{r1}:I{r2}")
    dv_list(ws, f"J{r1}:J{r2}", "lst_Prioridade")
    cf_bar(ws, f"K{r1}:K{r2}", 0, 1, VIOLET)
    ws.conditional_formatting.add(f"F{r1}:F{r2}", FormulaRule(formula=[f"N(F{r1})>=2"], font=Font(color=CRIT_T, bold=True)))

    r1, r2 = OPP_R1, OPP_R2
    section(ws, 31, "B", "L", "OPORTUNIDADES & CANDIDATURAS")
    headers(ws, 32, [("B", "Empresa / oportunidade", "in"), ("C", "Cargo", "in"), ("D", "Data", "in"),
                     ("E", "Etapa", "in"), ("F", "Dias desde", "auto"), ("G", "Alerta", "auto"),
                     ("H", "Próxima ação", "in"), ("I", "Data da ação", "in"), ("J", "Remuneração (€)", "in"),
                     ("K", "Fonte / contato", "in"), ("L", "Notas", "in")])
    body(ws, "B", "C", r1, r2)
    body(ws, "D", "D", r1, r2, fmt=DATE, h="center")
    body(ws, "E", "E", r1, r2)
    body(ws, "F", "F", r1, r2, "auto", h="center")
    body(ws, "G", "G", r1, r2, "auto")
    body(ws, "H", "H", r1, r2, wrap=True)
    body(ws, "I", "I", r1, r2, fmt=DATE, h="center")
    body(ws, "J", "J", r1, r2, fmt='€ #,##0', h="right")
    body(ws, "K", "L", r1, r2)
    formulas(ws, "F", r1, r2, lambda r: f'=IF(OR(B{r}="",D{r}=""),"",Hoje-D{r})')
    formulas(ws, "G", r1, r2, lambda r: (
        f'=IF(B{r}="","",IF(E{r}="Aceito","✔ Aceito",IF(OR(E{r}="Recusado",E{r}="Desisti"),"– Encerrado",'
        f'IF(AND(I{r}<>"",I{r}<Hoje),"✖ Ação atrasada",IF(AND(E{r}="Aplicado",N(F{r})>=FollowUpDias),"▲ Fazer follow-up",'
        f'IF(AND(I{r}<>"",I{r}-Hoje<=3),"▲ Ação em breve","○ Em andamento"))))))'))
    dv_date(ws, f"D{r1}:D{r2}")
    dv_date(ws, f"I{r1}:I{r2}")
    dv_list(ws, f"E{r1}:E{r2}", "lst_Etapa")
    cf_status(ws, f"G{r1}:G{r2}")

    r1, r2 = WIN_R1, WIN_R2
    section(ws, 64, "B", "L", "CONQUISTAS & MARCOS · registre logo que acontecem")
    headers(ws, 65, [("B", "Conquista / marco", "in"), ("C", "Tipo", "in"), ("D", "Data", "in"),
                     ("E", "Impacto (1–5)", "in"), ("F", "Resultado mensurável", "in", "G"),
                     ("H", "Evidência / contexto / quem viu", "in", "K"), ("L", "Notas", "in")])
    body(ws, "B", "C", r1, r2, wrap=True)
    body(ws, "D", "D", r1, r2, fmt=DATE, h="center")
    body(ws, "E", "E", r1, r2, h="center", fmt="0")
    body(ws, "F", "L", r1, r2, wrap=True)
    merge_rows(ws, "F", "G", r1, r2)
    merge_rows(ws, "H", "K", r1, r2)
    dv_list(ws, f"C{r1}:C{r2}", "lst_TipoMarco")
    dv_date(ws, f"D{r1}:D{r2}")
    dv_num(ws, f"E{r1}:E{r2}", 1, 5, True)
    cf_diverge(ws, f"E{r1}:E{r2}", 1, 3, 5)
    ws.freeze_panes = "A5"


# =============================================================================
# Relações
# =============================================================================
PPL_R1, PPL_R2 = 10, 69
INT_R1, INT_R2 = 10, 609


def build_relacoes(ws):
    sheet_title(ws, "Relacionamentos & vida social",
                "Cadastre as pessoas importantes e com que frequência quer falar com cada uma. Registre os contatos à direita.", "W")
    setw(ws, {"B": 22, "C": 13, "D": 10, "E": 11, "F": 12, "G": 10, "H": 17, "I": 12, "J": 12, "K": 9, "L": 10,
              "M": 34, "P": 3, "Q": 11, "R": 22, "S": 15, "T": 10, "U": 10, "V": 34, "W": 13})
    ws.column_dimensions["N"].hidden = ws.column_dimensions["O"].hidden = True
    p1, p2, i1, i2 = PPL_R1, PPL_R2, INT_R1, INT_R2
    pr = lambda c: f"${c}${p1}:${c}${p2}"
    ir = lambda c: f"${c}${i1}:${c}${i2}"
    kpi_strip(ws, 5, [
        ("rel_n", "B", "B", "Pessoas no radar", f'=COUNTIF({pr("B")},"?*")', "0",
         f'="Família "&COUNTIF({pr("C")},"Família")&" · amigos "&COUNTIF({pr("C")},"Amizade")'),
        ("rel_emdia", "C", "D", "Contatos em dia",
         f'=IFERROR((COUNTIF({pr("H")},"✔*")+COUNTIF({pr("H")},"▲*"))/COUNTIFS({pr("B")},"?*",{pr("E")},">0"),"–")', PCT,
         '=IF(ISNUMBER(C6),IF(C6>=0.8,"✔ ","▲ ")&"dentro da frequência que você definiu","○ defina a frequência de contato")'),
        ("rel_atras", "E", "F", "Contatos atrasados", f'=COUNTIF({pr("H")},"✖*")', "0",
         '=IF(E6>0,"✖ veja a lista no Dashboard","✔ ninguém esquecido")'),
        ("rel_inter", "G", "H", "Interações no mês", f'=COUNTIFS({in_month(ir("Q"))})', "0",
         f'="Tempo junto: "&FIXED(SUMIFS({ir("U")},{in_month(ir("Q"))})/60,1)&" h"'),
        ("rel_qual", "I", "J", "Qualidade média", f'=IFERROR(AVERAGEIFS({ir("T")},{in_month(ir("Q"))}),"–")', '0.0"/5"',
         '="1 = superficial · 5 = profunda"'),
        ("rel_aniv", "K", "M", "Próximo aniversário",
         f'=IFERROR(INDEX({pr("B")},MATCH(SMALL({pr("O")},1),{pr("O")},0)),"–")', None,
         f'=IFERROR(IF(INT(SMALL({pr("O")},1))<=AlertaAniv,"▲ ","")&"em "&INT(SMALL({pr("O")},1))&" dias ("&DAY(INDEX({pr("J")},MATCH(SMALL({pr("O")},1),{pr("O")},0)))&"/"&MONTH(INDEX({pr("J")},MATCH(SMALL({pr("O")},1),{pr("O")},0)))&")","")'),
    ])
    ws["K6"].font = F(13, True, INK)
    headers(ws, 9, [("B", "Nome", "in"), ("C", "Relação", "in"), ("D", "Círculo", "in"),
                    ("E", "Falar a cada (dias)", "in"), ("F", "Último contato", "auto"), ("G", "Dias sem contato", "auto"),
                    ("H", "Status", "auto"), ("I", "Aniversário", "in"), ("J", "Próximo aniversário", "auto"),
                    ("K", "Faltam (dias)", "auto"), ("L", "Contatos no ano", "auto"),
                    ("M", "Notas · interesses, presentes, assuntos", "in")], 34)
    body(ws, "B", "D", p1, p2)
    body(ws, "E", "E", p1, p2, h="center", fmt="0")
    body(ws, "F", "F", p1, p2, "auto", fmt=DATE, h="center")
    body(ws, "G", "G", p1, p2, "auto", h="center", fmt="0")
    body(ws, "H", "H", p1, p2, "auto")
    body(ws, "I", "I", p1, p2, fmt=DATE, h="center")
    body(ws, "J", "J", p1, p2, "auto", fmt=DATE, h="center")
    body(ws, "K", "L", p1, p2, "auto", h="center", fmt="0")
    body(ws, "M", "M", p1, p2, wrap=True)
    last = lambda r: f'SUMPRODUCT(MAX(({ir("R")}=B{r})*({ir("Q")}<=Hoje)*{ir("Q")}))'
    formulas(ws, "F", p1, p2, lambda r: f'=IF(B{r}="","",IF({last(r)}=0,"",{last(r)}))')
    formulas(ws, "G", p1, p2, lambda r: f'=IF(OR(B{r}="",F{r}=""),"",Hoje-F{r})')
    formulas(ws, "H", p1, p2, lambda r: (
        f'=IF(B{r}="","",IF(N(E{r})=0,"○ Sem frequência",IF(F{r}="","✖ Sem registro",'
        f'IF(G{r}>E{r},"✖ Atrasado",IF(G{r}>=0.8*E{r},"▲ Em breve","✔ Em dia")))))'))
    formulas(ws, "J", p1, p2, lambda r: (
        f'=IF(OR(B{r}="",I{r}=""),"",DATE(YEAR(Hoje)+(DATE(YEAR(Hoje),MONTH(I{r}),DAY(I{r}))<Hoje),MONTH(I{r}),DAY(I{r})))'))
    formulas(ws, "K", p1, p2, lambda r: f'=IF(J{r}="","",J{r}-Hoje)')
    formulas(ws, "L", p1, p2, lambda r: f'=IF(B{r}="","",COUNTIFS({ir("R")},B{r},{ir("Q")},">="&IniAno,{ir("Q")},"<="&FimAno))')
    formulas(ws, "N", p1, p2, lambda r: (
        f'=IF(AND(B{r}<>"",N(E{r})>0),IF(F{r}="",99,G{r}/E{r})+ROW()/1000000,"")'))
    formulas(ws, "O", p1, p2, lambda r: f'=IF(K{r}="","",K{r}+ROW()/1000000)')
    dv_list(ws, f"C{p1}:C{p2}", "lst_Relacao")
    dv_list(ws, f"D{p1}:D{p2}", "lst_Circulo")
    dv_num(ws, f"E{p1}:E{p2}", 1, 3650, True, prompt="Ex.: 7 = toda semana · 30 = todo mês · 90 = a cada 3 meses")
    dv_date(ws, f"I{p1}:I{p2}", "Data de nascimento (o ano pode ser qualquer um)")
    cf_status(ws, f"H{p1}:H{p2}")
    ws.conditional_formatting.add(f"K{p1}:K{p2}", FormulaRule(formula=[f"AND(ISNUMBER(K{p1}),K{p1}<=AlertaAniv)"],
                                                               font=Font(color=WARN_T, bold=True), fill=fill(WARN_BG)))
    section(ws, 8, "Q", "W", "REGISTRO DE CONTATOS")
    section(ws, 8, "B", "M", "PESSOAS")
    headers(ws, 9, [("Q", "Data", "in"), ("R", "Pessoa", "in"), ("S", "Tipo", "in"), ("T", "Qualidade (1–5)", "in"),
                    ("U", "Duração (min)", "in"), ("V", "Notas · o que conversamos", "in"), ("W", "Relação", "auto")], 34)
    body(ws, "Q", "Q", i1, i2, fmt=DATE, h="center")
    body(ws, "R", "S", i1, i2)
    body(ws, "T", "U", i1, i2, h="center", fmt="0")
    body(ws, "V", "V", i1, i2)
    body(ws, "W", "W", i1, i2, "auto")
    formulas(ws, "W", i1, i2, lambda r: f'=IF(R{r}="","",IFERROR(INDEX({pr("C")},MATCH(R{r},{pr("B")},0)),"—"))')
    dv_date(ws, f"Q{i1}:Q{i2}")
    dv_list(ws, f"R{i1}:R{i2}", "lst_Pessoas", "Pessoas cadastradas à esquerda")
    dv_list(ws, f"S{i1}:S{i2}", "lst_TipoInter")
    dv_num(ws, f"T{i1}:T{i2}", 1, 5, True, prompt="1 = superficial · 5 = conversa profunda, presença real")
    dv_num(ws, f"U{i1}:U{i2}", 0, 1440, True)
    cf_diverge(ws, f"T{i1}:T{i2}", 1, 3, 5)
    cf_zebra(ws, f"Q{i1}:W{i2}", "Q")
    ws.freeze_panes = "C10"


# =============================================================================
# Aprendizado
# =============================================================================
APR_R1, APR_R2 = 10, 89
LOG_R1, LOG_R2 = 10, 509


def build_aprendizado(ws):
    sheet_title(ws, "Aprendizado & desenvolvimento pessoal",
                "Livros, cursos, idiomas e certificações à esquerda; cada sessão de estudo à direita (as horas somam sozinhas).", "U")
    setw(ws, {"B": 36, "C": 13, "D": 18, "E": 14, "F": 11, "G": 11, "H": 9, "I": 9, "J": 11, "K": 9, "L": 8,
              "M": 36, "N": 24, "P": 3, "Q": 11, "R": 30, "S": 8, "T": 13, "U": 36})
    ws.column_dimensions["O"].hidden = True
    a1, a2, l1, l2 = APR_R1, APR_R2, LOG_R1, LOG_R2
    ar = lambda c: f"${c}${a1}:${c}${a2}"
    lr = lambda c: f"${c}${l1}:${c}${l2}"
    last = f'SUMPRODUCT(MAX(({lr("Q")}<=Corte)*{lr("Q")}))'
    kpi_strip(ws, 5, [
        ("apr_h", "B", "B", "Horas de estudo no mês", f'=SUMIFS({lr("S")},{in_month(lr("Q"))})', '0.0" h"',
         '=IF(DiasRef=0,"",IF(B6>=MetaEstudoMes*DiasRef/DiasMes,"✔ ","▲ ")&"meta "&MetaEstudoMes&" h/mês")'),
        ("apr_hano", "C", "D", "Horas no ano", f'=SUMIFS({lr("S")},{lr("Q")},">="&IniAno,{lr("Q")},"<="&Corte)',
         '0" h"', '=IF(DiasAno=0,"","≈ "&FIXED(C6/DiasAno*7,1)&" h por semana")'),
        ("apr_livros", "E", "F", "Livros concluídos no ano",
         f'=COUNTIFS({ar("C")},"Livro",{ar("E")},"Concluído",{ar("G")},">="&IniAno,{ar("G")},"<="&FimAno)', "0",
         '=IF(DiasAno=0,"",IF(E6>=MetaLivrosAno*DiasAno/365*0.9,"✔ ","▲ ")&"meta "&MetaLivrosAno&" · esperado hoje "&ROUND(MetaLivrosAno*DiasAno/365,0))'),
        ("apr_and", "G", "I", "Em andamento", f'=COUNTIF({ar("E")},"Em andamento")', "0",
         '=IF(G6>3,"▲ muitos ao mesmo tempo dispersam","✔ foco saudável")'),
        ("apr_concl", "J", "L", "Concluídos no ano (tudo)",
         f'=COUNTIFS({ar("E")},"Concluído",{ar("G")},">="&IniAno,{ar("G")},"<="&FimAno)', "0",
         f'=IFERROR("nota média "&FIXED(AVERAGEIFS({ar("L")},{ar("E")},"Concluído"),1)&"/5","")'),
        ("apr_ult", "M", "N", "Dias desde o último estudo", f'=IF({last}=0,"–",Corte-{last})', "0",
         '=IF(ISNUMBER(M6),IF(M6<=2,"✔ constância","▲ retome hoje, nem que seja 15 min"),"")'),
    ])
    section(ws, 8, "B", "N", "O QUE ESTOU APRENDENDO")
    section(ws, 8, "Q", "U", "SESSÕES DE ESTUDO")
    headers(ws, 9, [("B", "Título", "in"), ("C", "Tipo", "in"), ("D", "Área da vida", "in"), ("E", "Status", "in"),
                    ("F", "Início", "in"), ("G", "Conclusão", "in"), ("H", "Total (pág./aulas)", "in"),
                    ("I", "Onde estou", "in"), ("J", "Progresso", "auto"), ("K", "Horas", "auto"),
                    ("L", "Nota (1–5)", "in"), ("M", "Principais aprendizados", "in"), ("N", "Próximo passo / notas", "in"),
                    ("Q", "Data", "in"), ("R", "Item / assunto", "in"), ("S", "Horas", "in"), ("T", "Atividade", "in"),
                    ("U", "O que aprendi / notas", "in")], 34)
    body(ws, "B", "E", a1, a2)
    body(ws, "F", "G", a1, a2, fmt=DATE, h="center")
    body(ws, "H", "I", a1, a2, h="center", fmt="0")
    body(ws, "J", "J", a1, a2, "auto", fmt=PCT, h="center")
    body(ws, "K", "K", a1, a2, "auto", fmt='0.0;;"–"', h="center")
    body(ws, "L", "L", a1, a2, h="center", fmt="0")
    body(ws, "M", "N", a1, a2, wrap=True)
    formulas(ws, "J", a1, a2, lambda r: (
        f'=IF(B{r}="","",IF(E{r}="Concluído",1,IF(AND(N(H{r})>0,ISNUMBER(I{r})),MIN(1,I{r}/H{r}),0)))'))
    formulas(ws, "K", a1, a2, lambda r: f'=IF(B{r}="","",SUMIF({lr("R")},B{r},{lr("S")}))')
    formulas(ws, "O", a1, a2, lambda r: f'=IF(AND(B{r}<>"",E{r}="Em andamento"),ROW(),"")')
    dv_list(ws, f"C{a1}:C{a2}", "lst_TipoApr")
    dv_list(ws, f"D{a1}:D{a2}", "lst_Areas")
    dv_list(ws, f"E{a1}:E{a2}", "lst_StatusApr")
    dv_date(ws, f"F{a1}:G{a2}")
    dv_num(ws, f"L{a1}:L{a2}", 1, 5, True)
    cf_bar(ws, f"J{a1}:J{a2}", 0, 1, VIOLET)
    ws.conditional_formatting.add(f"E{a1}:E{a2}", FormulaRule(formula=[f'$E{a1}="Em andamento"'], font=Font(color=VIOLET, bold=True)))
    ws.conditional_formatting.add(f"E{a1}:E{a2}", FormulaRule(formula=[f'$E{a1}="Concluído"'], font=Font(color=GOOD_T, bold=True)))
    body(ws, "Q", "Q", l1, l2, fmt=DATE, h="center")
    body(ws, "R", "R", l1, l2)
    body(ws, "S", "S", l1, l2, h="center", fmt="0.0")
    body(ws, "T", "U", l1, l2)
    dv_date(ws, f"Q{l1}:Q{l2}")
    dv_list(ws, f"R{l1}:R{l2}", "lst_ItensApr", "Escolha um item da lista ou escreva um assunto livre", strict=False)
    dv_num(ws, f"S{l1}:S{l2}", 0, 24, prompt="Horas (ex.: 0,5 = 30 min)")
    dv_list(ws, f"T{l1}:T{l2}", "lst_AtivEstudo")
    cf_zebra(ws, f"Q{l1}:U{l2}", "Q")
    ws.freeze_panes = "C10"


# =============================================================================
# Lazer
# =============================================================================
LAZ_R1, LAZ_R2 = 10, 409
BKT_R1, BKT_R2 = 10, 59


def build_lazer(ws):
    sheet_title(ws, "Lazer & diversão",
                "Tempo livre também se cuida: registre o que fez e como se sentiu. À direita, a lista do que ainda quer viver.", "R")
    setw(ws, {"B": 11, "C": 32, "D": 18, "E": 16, "F": 10, "G": 10, "H": 11, "I": 10, "J": 30, "K": 3,
              "L": 34, "M": 16, "N": 12, "O": 12, "P": 13, "Q": 11, "R": 30})
    l1, l2, b1, b2 = LAZ_R1, LAZ_R2, BKT_R1, BKT_R2
    lr = lambda c: f"${c}${l1}:${c}${l2}"
    br = lambda c: f"${c}${b1}:${c}${b2}"
    last = f'SUMPRODUCT(MAX(({lr("B")}<=Corte)*{lr("B")}))'
    kpi_strip(ws, 5, [
        ("laz_h", "B", "C", "Horas de lazer no mês", f'=SUMIFS({lr("F")},{in_month(lr("B"))})', '0.0" h"',
         '=IF(DiasRef=0,"",IF(B6>=MetaLazerSem*DiasRef/7,"✔ ","▲ ")&"meta "&ROUND(MetaLazerSem*DiasRef/7,0)&" h até agora")'),
        ("laz_n", "D", "D", "Atividades no mês", f'=COUNTIFS({in_month(lr("B"))})', "0",
         f'="Com outras pessoas: "&COUNTIFS({lr("E")},"?*",{in_month(lr("B"))})'),
        ("laz_sat", "E", "F", "Satisfação média", f'=IFERROR(AVERAGEIFS({lr("H")},{in_month(lr("B"))}),"–")', '0.0"/5"',
         '=IF(ISNUMBER(E6),IF(E6>=4,"✔ recarregando","▲ escolha atividades que renovam"),"")'),
        ("laz_custo", "G", "H", "Gasto com lazer no mês", f'=SUMIFS({lr("G")},{in_month(lr("B"))})', EUR0,
         '=IFERROR("€ "&FIXED(G6/B6,0)&" por hora de lazer","")'),
        ("laz_ult", "I", "J", "Dias desde o último lazer", f'=IF({last}=0,"–",Corte-{last})', "0",
         '=IF(ISNUMBER(I6),IF(I6<=4,"✔ ","▲ ")&"lazer não é prêmio, é manutenção","")'),
        ("laz_sonho", "L", "N", "Sonhos realizados", f'=COUNTIF({br("P")},"Realizado")', "0",
         f'="de "&COUNTIF({br("L")},"?*")&" na lista · "&COUNTIF({br("P")},"Agendado")&" agendados"'),
    ])
    section(ws, 8, "B", "J", "REGISTRO DE LAZER")
    section(ws, 8, "L", "R", "LISTA DE SONHOS & EXPERIÊNCIAS")
    headers(ws, 9, [("B", "Data", "in"), ("C", "Atividade", "in"), ("D", "Categoria", "in"), ("E", "Com quem", "in"),
                    ("F", "Duração (h)", "in"), ("G", "Custo (€)", "in"), ("H", "Satisfação (1–5)", "in"),
                    ("I", "Repetiria?", "in"), ("J", "Notas", "in"),
                    ("L", "Sonho / experiência", "in"), ("M", "Categoria", "in"), ("N", "Custo estimado (€)", "in"),
                    ("O", "Quando (alvo)", "in"), ("P", "Status", "in"), ("Q", "Prioridade", "in"), ("R", "Notas", "in")], 34)
    body(ws, "B", "B", l1, l2, fmt=DATE, h="center")
    body(ws, "C", "E", l1, l2)
    body(ws, "F", "F", l1, l2, h="center", fmt="0.0")
    body(ws, "G", "G", l1, l2, h="right", fmt="#,##0.00")
    body(ws, "H", "I", l1, l2, h="center", fmt="0")
    body(ws, "J", "J", l1, l2)
    dv_date(ws, f"B{l1}:B{l2}")
    dv_list(ws, f"D{l1}:D{l2}", "lst_CatLazer")
    dv_num(ws, f"F{l1}:F{l2}", 0, 72)
    dv_num(ws, f"G{l1}:G{l2}", 0, 1000000)
    dv_num(ws, f"H{l1}:H{l2}", 1, 5, True, prompt="Quanto isso te recarregou? 1 = nada · 5 = muito")
    dv_list(ws, f"I{l1}:I{l2}", "lst_SimNao")
    cf_diverge(ws, f"H{l1}:H{l2}", 1, 3, 5)
    cf_zebra(ws, f"B{l1}:J{l2}", "B")
    body(ws, "L", "M", b1, b2)
    body(ws, "N", "N", b1, b2, h="right", fmt="#,##0")
    body(ws, "O", "O", b1, b2, fmt=DATE, h="center")
    body(ws, "P", "Q", b1, b2)
    body(ws, "R", "R", b1, b2)
    dv_list(ws, f"M{b1}:M{b2}", "lst_CatLazer")
    dv_date(ws, f"O{b1}:O{b2}")
    dv_list(ws, f"P{b1}:P{b2}", "lst_StatusSonho")
    dv_list(ws, f"Q{b1}:Q{b2}", "lst_Prioridade")
    ws.conditional_formatting.add(f"L{b1}:L{b2}", FormulaRule(formula=[f'$P{b1}="Realizado"'], font=Font(color=GOOD_T, strike=True)))
    ws.conditional_formatting.add(f"P{b1}:P{b2}", FormulaRule(formula=[f'$P{b1}="Agendado"'], font=Font(color=MAGENTA, bold=True)))
    ws.freeze_panes = "A10"


# =============================================================================
# Casa & documentos
# =============================================================================
DOC_R1, DOC_R2 = 11, 35
ROT_R1, ROT_R2 = 39, 73
ASS_R1, ASS_R2 = 77, 101


def build_casa(ws):
    sheet_title(ws, "Casa, documentos & rotinas",
                "Vencimentos de documentos, manutenção recorrente da casa/vida digital e assinaturas. Tudo avisa antes de vencer.", "L")
    setw(ws, {"B": 34, "C": 18, "D": 14, "E": 13, "F": 13, "G": 12, "H": 17, "I": 22, "J": 14, "K": 26, "L": 28})
    ws.column_dimensions["N"].hidden = ws.column_dimensions["O"].hidden = True
    d1, d2, r1, r2, s1, s2 = DOC_R1, DOC_R2, ROT_R1, ROT_R2, ASS_R1, ASS_R2
    kpi_strip(ws, 5, [
        ("casa_venc", "B", "B", "Documentos vencidos", f'=COUNTIF($H${d1}:$H${d2},"✖*")', "0",
         '=IF(B6>0,"✖ regularize","✔ nenhum")'),
        ("casa_renov", "C", "D", "Documentos a renovar", f'=COUNTIF($H${d1}:$H${d2},"▲*")', "0",
         '="vencem em até "&AlertaDocs&" dias"'),
        ("casa_rotatr", "E", "F", "Rotinas atrasadas", f'=COUNTIF($H${r1}:$H${r2},"✖*")', "0",
         '=IF(E6>0,"✖ veja a lista abaixo","✔ tudo em dia")'),
        ("casa_rotbreve", "G", "H", "Rotinas nos próximos 7 dias", f'=COUNTIF($H${r1}:$H${r2},"▲*")', "0", '=""'),
        ("casa_ass", "I", "J", "Assinaturas por mês", f"=SUM($F${s1}:$F${s2})", EUR0,
         f'=COUNTIF($B${s1}:$B${s2},"?*")&" serviços · € "&FIXED(SUM($G${s1}:$G${s2}),0)&" por ano"'),
        ("casa_econ", "K", "L", "Economia possível", f'=SUMIF($J${s1}:$J${s2},"Baixo",$F${s1}:$F${s2})', EUR0,
         '=IF(K6>0,"▲ por mês em serviços de pouco uso","✔ nada sobrando")'),
    ])
    section(ws, 9, "B", "L", "DOCUMENTOS & BUROCRACIA")
    headers(ws, 10, [("B", "Documento", "in"), ("C", "Titular", "in"), ("D", "Número / ref.", "in"), ("E", "Emissão", "in"),
                     ("F", "Validade", "in"), ("G", "Dias p/ vencer", "auto"), ("H", "Status", "auto"),
                     ("I", "Próxima ação", "in"), ("J", "Notas", "in", "L")])
    body(ws, "B", "D", d1, d2)
    body(ws, "E", "F", d1, d2, fmt=DATE, h="center")
    body(ws, "G", "G", d1, d2, "auto", fmt='0;[Red]-0', h="center")
    body(ws, "H", "H", d1, d2, "auto")
    body(ws, "I", "L", d1, d2, wrap=True)
    merge_rows(ws, "J", "L", d1, d2)
    formulas(ws, "G", d1, d2, lambda r: f'=IF(OR(B{r}="",F{r}=""),"",F{r}-Hoje)')
    formulas(ws, "H", d1, d2, lambda r: (
        f'=IF(B{r}="","",IF(F{r}="","○ Sem validade",IF(G{r}<0,"✖ Vencido",IF(G{r}<=AlertaDocs,"▲ Renovar","✔ Válido"))))'))
    formulas(ws, "N", d1, d2, lambda r: f'=IF(AND(B{r}<>"",F{r}<>""),IF(G{r}<=AlertaDocs,G{r}+ROW()/1000000,""),"")')
    formulas(ws, "O", d1, d2, lambda r: f'=IF(N{r}="","","Documento")')
    dv_date(ws, f"E{d1}:F{d2}")
    cf_status(ws, f"H{d1}:H{d2}")

    section(ws, 37, "B", "L", "ROTINAS & MANUTENÇÃO RECORRENTE · casa, carro, vida digital, finanças")
    headers(ws, 38, [("B", "Rotina", "in"), ("C", "Categoria", "in"), ("D", "A cada (dias)", "in"), ("E", "Última vez", "in"),
                     ("F", "Próxima", "auto"), ("G", "Dias", "auto"), ("H", "Status", "auto"), ("I", "Responsável", "in"),
                     ("J", "Duração (min)", "in"), ("K", "Notas", "in", "L")])
    body(ws, "B", "C", r1, r2)
    body(ws, "D", "D", r1, r2, h="center", fmt="0")
    body(ws, "E", "E", r1, r2, fmt=DATE, h="center")
    body(ws, "F", "F", r1, r2, "auto", fmt=DATE, h="center")
    body(ws, "G", "G", r1, r2, "auto", fmt='0;[Red]-0', h="center")
    body(ws, "H", "H", r1, r2, "auto")
    body(ws, "I", "I", r1, r2)
    body(ws, "J", "J", r1, r2, h="center", fmt="0")
    body(ws, "K", "L", r1, r2)
    merge_rows(ws, "K", "L", r1, r2)
    formulas(ws, "F", r1, r2, lambda r: f'=IF(OR(B{r}="",E{r}="",N(D{r})=0),"",E{r}+D{r})')
    formulas(ws, "G", r1, r2, lambda r: f'=IF(F{r}="","",F{r}-Hoje)')
    formulas(ws, "H", r1, r2, lambda r: (
        f'=IF(B{r}="","",IF(F{r}="","○ Sem registro",IF(G{r}<0,"✖ Atrasada",IF(G{r}<=7,"▲ Em breve","✔ Em dia"))))'))
    formulas(ws, "N", r1, r2, lambda r: f'=IF(F{r}="","",IF(G{r}<=7,G{r}+ROW()/1000000,""))')
    formulas(ws, "O", r1, r2, lambda r: f'=IF(N{r}="","","Rotina")')
    dv_list(ws, f"C{r1}:C{r2}", "lst_CatManut")
    dv_num(ws, f"D{r1}:D{r2}", 1, 3650, True, prompt="7 = semanal · 30 = mensal · 90 = trimestral · 365 = anual")
    dv_date(ws, f"E{r1}:E{r2}", "Quando fez pela última vez — atualize ao fazer de novo")
    cf_status(ws, f"H{r1}:H{r2}")

    section(ws, 75, "B", "L", "ASSINATURAS & CONTRATOS")
    headers(ws, 76, [("B", "Serviço", "in"), ("C", "Categoria", "in"), ("D", "Valor (€)", "in"), ("E", "Periodicidade", "in"),
                     ("F", "Custo mensal", "auto"), ("G", "Custo anual", "auto"), ("H", "Próxima renovação", "in"),
                     ("I", "Dias p/ renovar", "auto"), ("J", "Uso real", "in"), ("K", "Sugestão", "auto"), ("L", "Notas", "in")])
    body(ws, "B", "C", s1, s2)
    body(ws, "D", "D", s1, s2, fmt="#,##0.00", h="right")
    body(ws, "E", "E", s1, s2)
    body(ws, "F", "G", s1, s2, "auto", fmt='#,##0.00;;"–"', h="right")
    body(ws, "H", "H", s1, s2, fmt=DATE, h="center")
    body(ws, "I", "I", s1, s2, "auto", fmt='0;[Red]-0', h="center")
    body(ws, "J", "J", s1, s2)
    body(ws, "K", "K", s1, s2, "auto")
    body(ws, "L", "L", s1, s2)
    per = A(S.CFG, LIST_COL["lst_Periodo"], 6, LIST_COL["lst_Periodo"], 40)
    perm = A(S.CFG, LIST_COL["lst_PeriodoMeses"], 6, LIST_COL["lst_PeriodoMeses"], 40)
    formulas(ws, "F", s1, s2, lambda r: f'=IF(OR(B{r}="",D{r}=""),"",D{r}/IFERROR(INDEX({perm},MATCH(E{r},{per},0)),1))')
    formulas(ws, "G", s1, s2, lambda r: f'=IF(F{r}="","",F{r}*12)')
    formulas(ws, "I", s1, s2, lambda r: f'=IF(OR(B{r}="",H{r}=""),"",H{r}-Hoje)')
    formulas(ws, "K", s1, s2, lambda r: (
        f'=IF(B{r}="","",IF(J{r}="Baixo","✖ Avaliar cancelamento",IF(AND(I{r}<>"",N(I{r})<=15),"▲ Renova em breve","✔ Mantém")))'))
    formulas(ws, "N", s1, s2, lambda r: f'=IF(I{r}="","",IF(I{r}<=15,I{r}+ROW()/1000000,""))')
    formulas(ws, "O", s1, s2, lambda r: f'=IF(N{r}="","","Assinatura")')
    dv_list(ws, f"C{s1}:C{s2}", "lst_CatDesp", strict=False)
    dv_num(ws, f"D{s1}:D{s2}", 0, 1000000)
    dv_list(ws, f"E{s1}:E{s2}", "lst_Periodo")
    dv_date(ws, f"H{s1}:H{s2}")
    dv_list(ws, f"J{s1}:J{s2}", "lst_Uso", "Seja honesto: quanto você realmente usa?")
    cf_status(ws, f"K{s1}:K{s2}")
    ws.freeze_panes = "A5"


# =============================================================================
# Roda da Vida & revisão mensal
# =============================================================================
RODA_R1, RODA_R2 = 7, 17
REV_R1, REV_R2 = 22, 33
RCOLS = [CL(3 + i) for i in range(12)]   # C..N


def build_roda(ws):
    sheet_title(ws, "Roda da Vida & revisão mensal",
                "Uma vez por mês: dê uma nota de 0 a 10 para cada área e escreva a revisão. O Dashboard compara sua percepção com os dados.", "V")
    setw(ws, {"B": 28, **{c: 7.5 for c in RCOLS}, "O": 8, "P": 8, "Q": 9, "R": 7, "S": 7, "T": 8, "U": 8, "V": 11})
    ws.column_dimensions["W"].hidden = ws.column_dimensions["X"].hidden = True
    r1, r2 = RODA_R1, RODA_R2
    section(ws, 5, "B", "V", "NOTAS MENSAIS · de 0 a 10, como você está em cada área")
    headers(ws, 6, [("B", "Área da vida", "in")] + [(c, MESES_ABR[i], "in") for i, c in enumerate(RCOLS)] +
            [("O", "Alvo", "in"), ("P", "Atual", "auto"), ("Q", "Anterior", "auto"), ("R", "Δ", "auto"),
             ("S", "Tend.", "auto"), ("T", "Falta p/ alvo", "auto"), ("U", "Dados", "auto"),
             ("V", "Percepção − dados", "auto")], 34)
    body(ws, "B", "B", r1, r2 + 1, "auto")
    body(ws, "C", "O", r1, r2, h="center", fmt="0")
    body(ws, "P", "V", r1, r2 + 1, "auto", h="center", fmt='0.0;[Red]-0.0;"0"')
    for i, r in enumerate(range(r1, r2 + 1)):
        ws[f"B{r}"] = f"={A(S.CFG, LIST_COL['lst_Areas'], LIST_FIRST + i)}"
        ws[f"B{r}"].font = F(9, True, INK)
        cols = f"$C{r}:$N{r}"
        idx = f"(COLUMN({cols})-COLUMN($C{r})+1)"
        ws[f"W{r}"] = f'=SUMPRODUCT(MAX(({cols}<>"")*({idx}<=RefMes)*{idx}))'
        ws[f"X{r}"] = f'=SUMPRODUCT(MAX(({cols}<>"")*({idx}<W{r})*{idx}))'
        ws[f"P{r}"] = f'=IF(W{r}=0,"",INDEX({cols},W{r}))'
        ws[f"Q{r}"] = f'=IF(X{r}=0,"",INDEX({cols},X{r}))'
        ws[f"R{r}"] = f'=IF(OR(P{r}="",Q{r}=""),"",P{r}-Q{r})'
        ws[f"S{r}"] = f'=IF(R{r}="","",IF(R{r}>0,"↑",IF(R{r}<0,"↓","→")))'
        ws[f"T{r}"] = f'=IF(OR(O{r}="",P{r}=""),"",MAX(0,O{r}-P{r}))'
        ws[f"U{r}"] = f"={A(S.CALC, 'H', 5 + i)}"
        ws[f"V{r}"] = f'=IF(OR(P{r}="",U{r}=""),"",P{r}-U{r})'
    t = r2 + 1
    ws[f"B{t}"] = "Média"
    for c in RCOLS + ["O", "P", "Q", "R", "T", "U", "V"]:
        ws[f"{c}{t}"] = f'=IFERROR(AVERAGE({c}{r1}:{c}{r2}),"")'
        ws[f"{c}{t}"].number_format = "0.0"
    body(ws, "C", "O", t, t, "auto", h="center", fmt="0.0")
    for c in range(2, 23):
        ws.cell(t, c).font = F(9, True, INK)
        ws.cell(t, c).border = Border(top=side(INK), bottom=side(INK))
    dv_num(ws, f"C{r1}:O{r2}", 0, 10, prompt="0 = péssimo · 5 = ok · 10 = plenamente satisfeito")
    cf_diverge(ws, f"C{r1}:O{r2}", 0, 5, 10)
    ws.conditional_formatting.add(f"S{r1}:S{r2}", FormulaRule(formula=[f'S{r1}="↑"'], font=Font(color=GOOD_T, bold=True)))
    ws.conditional_formatting.add(f"S{r1}:S{r2}", FormulaRule(formula=[f'S{r1}="↓"'], font=Font(color=CRIT_T, bold=True)))
    ws.conditional_formatting.add(f"V{r1}:V{r2}", FormulaRule(formula=[f'AND(ISNUMBER(V{r1}),ABS(V{r1})>=2)'],
                                                               fill=fill(WARN_BG), font=Font(color=WARN_T, bold=True)))
    ws.conditional_formatting.add(f"C{r1}:N{r2}", FormulaRule(formula=[f"COLUMN(C{r1})-2=RefMes"],
                                                               border=Border(left=side(INK, "medium"), right=side(INK, "medium"))))
    note = ws[f"B{t+1}"]
    note.value = ("Dados = placar objetivo da área (0–10), vindo de metas, hábitos e métricas das abas. "
                  "Diferença de 2 pontos ou mais (amarelo) merece uma pergunta: o que estou vendo que os números não mostram — ou o contrário?")
    note.font = F(8, False, MUTED, italic=True)

    section(ws, 20, "B", "V", "REVISÃO MENSAL · 15 minutos no último dia do mês")
    headers(ws, 21, [("B", "Mês", "in"), ("C", "Nota do mês", "in"), ("D", "Vitórias & destaques", "in", "F"),
                     ("G", "O que não funcionou", "in", "I"), ("J", "O que aprendi", "in", "L"),
                     ("M", "Foco do próximo mês", "in", "O"), ("P", "Gratidão · notas livres", "in", "V")], 30)
    q1, q2 = REV_R1, REV_R2
    body(ws, "B", "B", q1, q2, "auto")
    body(ws, "C", "C", q1, q2, h="center", fmt="0")
    body(ws, "D", "V", q1, q2, wrap=True, size=9)
    for c1, c2 in (("D", "F"), ("G", "I"), ("J", "L"), ("M", "O"), ("P", "V")):
        merge_rows(ws, c1, c2, q1, q2)
    for i, r in enumerate(range(q1, q2 + 1)):
        ws[f"B{r}"] = f"=INDEX(Meses,{i + 1})"
        ws[f"B{r}"].font = F(9, True, INK)
        ws.row_dimensions[r].height = 58
        for c in ("D", "G", "J", "M", "P"):
            ws[f"{c}{r}"].alignment = al("left", "top", True, 1)
    dv_num(ws, f"C{q1}:C{q2}", 0, 10, True)
    cf_diverge(ws, f"C{q1}:C{q2}", 0, 5, 10)
    ws.conditional_formatting.add(f"B{q1}:B{q2}", FormulaRule(formula=[f"ROW(B{q1})-{q1 - 1}=RefMes"], fill=fill(TODAY_BG)))
    ws.freeze_panes = "C7"


# =============================================================================
# Cálculos (motor transparente que alimenta o Dashboard)
# =============================================================================
CALC = {}   # endereços úteis para o Dashboard


def build_calculos(ws):
    sheet_title(ws, "Cálculos · motor do Dashboard",
                "Aba técnica e transparente: mostra de onde vem cada número do Dashboard. Não precisa editar nada aqui.", "O",
                legend=False)
    setw(ws, {"B": 30, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 12, "I": 12, "J": 10, "K": 11, "L": 22,
              "M": 52, "N": 60, "O": 4})
    rel = lambda i: A(S.CFG, LIST_COL["lst_Relacao"], LIST_FIRST + i)
    Rp = lambda c: A(S.REL, c, PPL_R1, c, PPL_R2)
    Ri = lambda c: A(S.REL, c, INT_R1, c, INT_R2)
    U = A(S.ORC, "U", ORC_DESP1, "U", ORC_DESP2)
    share = lambda rng: (f'IFERROR(COUNTIF({rng},"✔*")/(COUNTIF({rng},"✔*")+COUNTIF({rng},"▲*")+COUNTIF({rng},"✖*")),"")')
    docs, rots = A(S.CASA, "H", DOC_R1, "H", DOC_R2), A(S.CASA, "H", ROT_R1, "H", ROT_R2)

    # ---------------- bloco 1: áreas
    section(ws, 3, "B", "N", "1 · PLACAR POR ÁREA (0–1 por componente; Dados = média × 10)")
    headers(ws, 4, [("B", "Área", "auto"), ("C", "Metas no ritmo", "auto"), ("D", "Hábitos (% meta)", "auto"),
                    ("E", "Métrica 1", "auto"), ("F", "Métrica 2", "auto"), ("G", "Métrica 3", "auto"),
                    ("H", "Dados", "auto"), ("I", "Percepção", "auto"), ("J", "Alvo", "auto"),
                    ("K", "Percepção − dados", "auto"), ("L", "Status", "auto"), ("M", "Indicador-chave", "auto"),
                    ("N", "Como é calculado", "auto")], 32)
    t, sn, ps, hu = K["sau_treinos"], K["sau_sono"], K["sau_passos"], K["sau_humor"]
    tx, rs = K["fin_taxa"], K["reserva_meses"]

    def rel_metrics(i):
        R = rel(i)
        ok = f'(COUNTIFS({Rp("C")},{R},{Rp("H")},"✔*")+COUNTIFS({Rp("C")},{R},{Rp("H")},"▲*"))'
        tot = f'COUNTIFS({Rp("C")},{R},{Rp("E")},">0")'
        mes = f'COUNTIFS({Ri("W")},{R},{Ri("Q")},">="&RefIni,{Ri("Q")},"<="&RefFim)'
        return (f'=IFERROR({ok}/{tot},"")',
                f'=IFERROR((AVERAGEIFS({Ri("T")},{Ri("W")},{R},{Ri("Q")},">="&RefIni,{Ri("Q")},"<="&RefFim)-1)/4,"")',
                '=""',
                f'={ok}&" de "&{tot}&" em dia · "&{mes}&" contatos no mês"',
                "% das pessoas desse tipo dentro da frequência desejada; qualidade média dos contatos no mês.")

    generic_ind = lambda r: (f'=IF(AND(C{r}="",D{r}=""),"Ligue hábitos e metas a esta área",'
                             f'IF(D{r}="","",ROUND(D{r}*100,0)&"% da meta dos hábitos")&IF(AND(C{r}<>"",D{r}<>"")," · ","")'
                             f'&IF(C{r}="","",ROUND(C{r}*100,0)&"% das metas no ritmo"))')
    spec = [
        (f'=IF(OR(DiasRef=0,N({K["sau_dias"]})=0),"",MIN(1,{t}/(MetaTreinosSem*DiasRef/7)))', f'=IF(ISNUMBER({sn}),MIN(1,{sn}/MetaSono),"")',
         f'=IF(ISNUMBER({ps}),MIN(1,{ps}/MetaPassos),"")',
         f'=IF(ISNUMBER({sn}),"Sono "&FIXED({sn},1)&" h · ","")&{t}&" treinos"&IF(ISNUMBER({ps})," · "&FIXED({ps},0)&" passos/dia","")',
         "Treinos vs. meta semanal · sono médio vs. meta · passos vs. meta (mês de referência)."),
        (f'=IF(ISNUMBER({hu}),({hu}-1)/4,"")', '=IF(ISNUMBER(C153),(5-C153)/4,"")', '=IF(ISNUMBER(C154),(C154-1)/4,"")',
         f'=IF(ISNUMBER({hu}),"Humor "&FIXED({hu},1)&" · estresse "&IF(ISNUMBER(C153),FIXED(C153,1),"–")&" (1–5)","Sem registros de humor no mês")',
         "Humor médio, estresse médio (invertido) e energia média do registro diário, em escala 0–1."),
        (f'=IF(ISNUMBER({tx}),MIN(1,MAX(0,{tx})/MetaPoup),"")', f'=IF(ISNUMBER({rs}),MIN(1,{rs}/MetaReserva),"")',
         f"={share(U)}",
         f'=IF(ISNUMBER({tx}),"Poupança "&ROUND({tx}*100,0)&"%","Sem receitas no mês")&IF(ISNUMBER({rs}),"; reserva "&FIXED({rs},1)&" meses","")',
         "Taxa de poupança vs. meta · reserva de emergência vs. meta · % das categorias dentro do orçamento."),
        (f'=IF(ISNUMBER({K["car_comp"]}),{K["car_comp"]},"")', f'=IF(ISNUMBER({K["car_sat"]}),{K["car_sat"]}/10,"")', '=""',
         f'={K["car_cand"]}&" candidaturas ativas"&IF(ISNUMBER({K["car_comp"]})," · competências "&ROUND({K["car_comp"]}*100,0)&"% do alvo","")',
         "Competências: nível atual ÷ alvo · satisfação com o trabalho ÷ 10."),
        (f'=IF(OR(DiasRef=0,COUNT({A(S.APR, "Q", LOG_R1, "Q", LOG_R2)})=0),"",MIN(1,{K["apr_h"]}/(MetaEstudoMes*DiasRef/DiasMes)))',
         f'=IF(OR(DiasAno=0,COUNTIF({A(S.APR, "C", APR_R1, "C", APR_R2)},"Livro")=0),"",MIN(1,{K["apr_livros"]}/MAX(0.5,MetaLivrosAno*DiasAno/365)))', '=""',
         f'=FIXED({K["apr_h"]},1)&" h de estudo no mês · "&{K["apr_livros"]}&" livros no ano"',
         "Horas de estudo vs. meta proporcional ao mês · livros concluídos vs. ritmo da meta anual."),
        rel_metrics(0), rel_metrics(1), rel_metrics(2),
        (f'=IF(OR(DiasRef=0,COUNT({A(S.LAZ, "B", LAZ_R1, "B", LAZ_R2)})=0),"",MIN(1,{K["laz_h"]}/(MetaLazerSem*DiasRef/7)))',
         f'=IF(ISNUMBER({K["laz_sat"]}),({K["laz_sat"]}-1)/4,"")', '=""',
         f'=FIXED({K["laz_h"]},1)&" h de lazer · "&{K["laz_n"]}&" atividades"&IF(ISNUMBER({K["laz_sat"]})," · satisfação "&FIXED({K["laz_sat"]},1),"")',
         "Horas de lazer vs. meta semanal · satisfação média das atividades."),
        ('=""', '=""', '=""', None, "Só metas e hábitos ligados a esta área (escolha a área ao criar metas/hábitos)."),
        (f"={share(docs)}", f"={share(rots)}", '=""',
         f'=({K["casa_venc"]}+{K["casa_renov"]})&" doc(s) p/ renovar · "&{K["casa_rotatr"]}&" rotina(s) atrasada(s)"',
         "% dos documentos válidos · % das rotinas recorrentes em dia."),
    ]
    body(ws, "B", "B", 5, 16, "auto")
    body(ws, "C", "G", 5, 15, "auto", fmt='0%;;0%', h="center")
    body(ws, "H", "K", 5, 16, "auto", fmt='0.0;[Red]-0.0;0.0', h="center")
    body(ws, "L", "N", 5, 15, "auto", wrap=True, size=8)
    MET = lambda c: A(S.METAS, c, MET_R1, c, MET_R2)
    for i in range(11):
        r = 5 + i
        e, f_, g, ind, how = spec[i]
        ws[f"B{r}"] = f"={A(S.CFG, LIST_COL['lst_Areas'], LIST_FIRST + i)}"
        ws[f"C{r}"] = f'=IFERROR(SUMIFS({MET("V")},{MET("D")},B{r})/SUMIFS({MET("U")},{MET("D")},B{r}),"")'
        ws[f"D{r}"] = f'=IFERROR(AVERAGEIFS({A(S.HAB, "D", 10, "M", 10)},{A(S.HAB, "D", 6, "M", 6)},B{r}),"")'
        ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"] = e, f_, g
        ws[f"H{r}"] = f'=IFERROR(ROUND(AVERAGE(C{r}:G{r})*10,1),"")'
        ws[f"I{r}"] = f'=IF({A(S.RODA, "P", RODA_R1 + i)}="","",{A(S.RODA, "P", RODA_R1 + i)})'
        ws[f"J{r}"] = f'=IF({A(S.RODA, "O", RODA_R1 + i)}="","",{A(S.RODA, "O", RODA_R1 + i)})'
        ws[f"K{r}"] = f'=IF(OR(H{r}="",I{r}=""),"",I{r}-H{r})'
        ws[f"L{r}"] = f'=IF(H{r}="","○ Sem dados",IF(H{r}>=7,"✔ Bem",IF(H{r}>=5,"▲ Atenção","✖ Precisa de cuidado")))'
        ws[f"M{r}"] = ind if ind else generic_ind(r)
        ws[f"N{r}"] = "Metas da área no ritmo + hábitos da área (% da meta) + " + how
        ws.row_dimensions[r].height = 30
        ws[f"B{r}"].font = F(9, True, INK)
    ws["B16"] = "ÍNDICE DE VIDA (0–100) · percepção média (0–10)"
    ws["H16"] = '=IFERROR(ROUND(AVERAGE(H5:H15)*10,0),"")'
    ws["I16"] = '=IFERROR(ROUND(AVERAGE(I5:I15),1),"")'
    ws["H16"].number_format = "0"
    for c in "BHI":
        ws[f"{c}16"].font = F(10, True, INK)
    cf_status(ws, "L5:L15")
    CALC["indice"], CALC["percep"] = A(S.CALC, "H", 16), A(S.CALC, "I", 16)

    # ---------------- bloco 2: série mensal
    section(ws, 18, "B", "O", "2 · SÉRIE MENSAL (alimenta os gráficos)")
    hdrs = ["Mês", "Receitas", "Despesas", "Aportes", "Resultado", "Taxa de poupança", "Patrimônio líquido",
            "Sono médio (h)", "Treinos", "Humor médio", "Hábitos cumpridos", "Estudo (h)", "Lazer (h)", "Roda da Vida (média)"]
    headers(ws, 19, [(CL(2 + i), h, "auto") for i, h in enumerate(hdrs)], 32)
    body(ws, "B", "O", 20, 31, "auto", h="center", fmt='#,##0.0;[Red]-#,##0.0;0')
    for i in range(12):
        r, m, mc = 20 + i, i + 1, MCOLS[i]
        ini, fim = month_bounds(m)
        ws[f"B{r}"] = MESES_ABR[i]
        ws[f"C{r}"] = f"={A(S.ORC, mc, 7)}"
        ws[f"D{r}"] = f"={A(S.ORC, mc, 8)}"
        ws[f"E{r}"] = f"={A(S.ORC, mc, 10)}"
        ws[f"F{r}"] = f"={A(S.ORC, mc, 9)}"
        ws[f"G{r}"] = f'=IF({A(S.ORC, mc, 12)}="","",{A(S.ORC, mc, 12)})'
        ws[f"H{r}"] = f'=IF({A(S.ORC, mc, 84)}="","",{A(S.ORC, mc, 84)})'
        ws[f"I{r}"] = f"={A(S.SAU, 'R', 10 + m)}"
        ws[f"J{r}"] = f"={A(S.SAU, 'V', 10 + m)}"
        ws[f"K{r}"] = f"={A(S.SAU, 'S', 10 + m)}"
        ws[f"L{r}"] = (f'=IFERROR(SUMIFS({A(S.HAB, "N", HAB_R1, "N", HAB_R2)},{A(S.HAB, "B", HAB_R1, "B", HAB_R2)},">="&{ini},'
                       f'{A(S.HAB, "B", HAB_R1, "B", HAB_R2)},"<="&{fim})/(COUNTA({A(S.HAB, "D", 16, "M", 16)})*'
                       f'MAX(0,MIN({fim},Hoje)-{ini}+1)),"")')
        ws[f"M{r}"] = f'=SUMIFS({A(S.APR, "S", LOG_R1, "S", LOG_R2)},{A(S.APR, "Q", LOG_R1, "Q", LOG_R2)},">="&{ini},{A(S.APR, "Q", LOG_R1, "Q", LOG_R2)},"<="&{fim})'
        ws[f"N{r}"] = f'=SUMIFS({A(S.LAZ, "F", LAZ_R1, "F", LAZ_R2)},{A(S.LAZ, "B", LAZ_R1, "B", LAZ_R2)},">="&{ini},{A(S.LAZ, "B", LAZ_R1, "B", LAZ_R2)},"<="&{fim})'
        ws[f"O{r}"] = f'=IFERROR(AVERAGE({A(S.RODA, RCOLS[i], RODA_R1, RCOLS[i], RODA_R2)}),"")'
        for c in "GL":
            ws[f"{c}{r}"].number_format = '0%;[Red]-0%;0%'
        for c in "CDEFH":
            ws[f"{c}{r}"].number_format = '#,##0;[Red]-#,##0;0'
    ws.conditional_formatting.add("B20:O31", FormulaRule(formula=["ROW()-19=RefMes"], fill=fill(TODAY_BG)))

    # ---------------- bloco 3: despesas do mês ordenadas
    section(ws, 34, "B", "I", "3 · DESPESAS DO MÊS DE REFERÊNCIA, ORDENADAS")
    headers(ws, 35, [("B", "Categoria", "auto"), ("C", "Valor", "auto"), ("D", "Chave", "auto"), ("E", "", "auto"),
                     ("F", "Posição", "auto"), ("G", "Categoria (ordem)", "auto"), ("H", "Chave (ordem)", "auto"),
                     ("I", "Valor (ordem)", "auto")])
    body(ws, "B", "I", 36, 60, "auto", h="center", fmt="#,##0.00")
    for i in range(25):
        r = 36 + i
        ws[f"B{r}"] = f'={A(S.ORC, "B", ORC_DESP1 + i)}&""'
        ws[f"C{r}"] = f'=IF(B{r}="","",N({A(S.ORC, "S", ORC_DESP1 + i)}))'
        ws[f"D{r}"] = f'=IF(N(C{r})>0,C{r}+ROW()/1000000,"")'
        ws[f"F{r}"] = i + 1
        ws[f"H{r}"] = f'=IFERROR(LARGE($D$36:$D$60,F{r}),"")'
        ws[f"G{r}"] = f'=IF(H{r}="","",INDEX($B$36:$B$60,MATCH(H{r},$D$36:$D$60,0)))'
        ws[f"I{r}"] = f'=IF(H{r}="","",INDEX($C$36:$C$60,MATCH(H{r},$D$36:$D$60,0)))'

    # ---------------- bloco 4: hábitos no mês
    section(ws, 63, "B", "D", "4 · HÁBITOS NO MÊS")
    headers(ws, 64, [("B", "Hábito", "auto"), ("C", "% da meta", "auto"), ("D", "% dos dias", "auto")])
    body(ws, "B", "D", 65, 74, "auto", h="center", fmt=PCT)
    for i, c in enumerate(HAB_COLS):
        r = 65 + i
        ws[f"B{r}"] = f'=IF({A(S.HAB, c, 16)}="","",{A(S.HAB, c, 16)})'
        ws[f"C{r}"] = f'=IF(OR(B{r}="",NOT(ISNUMBER({A(S.HAB, c, 10)}))),"",{A(S.HAB, c, 10)})'
        ws[f"D{r}"] = f'=IF(OR(B{r}="",NOT(ISNUMBER({A(S.HAB, c, 9)}))),"",{A(S.HAB, c, 9)})'

    # ---------------- bloco 5: listas "top N"
    def top_list(sec_row, title, key_rng, fields, largest=False, min_key=None):
        section(ws, sec_row, "B", "I", title)
        headers(ws, sec_row + 1, [("B", "Chave", "auto")] + [(CL(3 + j), h, "auto") for j, (h, _, _) in enumerate(fields)])
        r1 = sec_row + 2
        body(ws, "B", CL(2 + len(fields)), r1, r1 + 7, "auto", h="center")
        for k in range(8):
            r = r1 + k
            fn = "LARGE" if largest else "SMALL"
            if min_key is None:
                ws[f"B{r}"] = f'=IFERROR({fn}({key_rng},{k + 1}),"")'
            else:
                ws[f"B{r}"] = f'=IFERROR(IF({fn}({key_rng},{k + 1})>={min_key},{fn}({key_rng},{k + 1}),""),"")'
            for j, (h, src, fmt) in enumerate(fields):
                c = CL(3 + j)
                if callable(src):
                    ws[f"{c}{r}"] = src(r)
                else:
                    ws[f"{c}{r}"] = f'=IF($B{r}="","",INDEX({src},MATCH($B{r},{key_rng},0)))'
                if fmt:
                    ws[f"{c}{r}"].number_format = fmt
        return r1

    T = lambda c: A(S.TAR, c, TSK_R1, c, TSK_R2)
    CALC["tar"] = top_list(77, "5a · TAREFAS ABERTAS POR PRAZO", T("N"), [
        ("Tarefa", T("B"), None), ("Prazo", T("F"), DATE_S), ("Alerta", T("K"), None),
        ("Projeto", lambda r: f'=IF($B{r}="","",INDEX({T("C")},MATCH($B{r},{T("N")},0))&"")', None)])
    Mt = lambda c: A(S.METAS, c, MET_R1, c, MET_R2)
    CALC["met"] = top_list(89, "5b · METAS EM ANDAMENTO POR PRAZO", Mt("W"), [
        ("Meta", Mt("C"), None), ("Área", Mt("D"), None), ("Progresso", Mt("M"), PCT), ("Esperado", Mt("N"), PCT),
        ("Ritmo", Mt("O"), None), ("Prazo", Mt("G"), DATE_S)])
    CALC["cont"] = top_list(101, "5c · CONTATOS A FAZER (maior atraso relativo)", Rp("N"), [
        ("Nome", Rp("B"), None),
        ("Sem contato", lambda r: f'=IF($B{r}="","",IF(INDEX({Rp("F")},MATCH($B{r},{Rp("N")},0))="","nunca",INDEX({Rp("G")},MATCH($B{r},{Rp("N")},0))&" dias"))', None),
        ("Status", Rp("H"), None), ("Relação", Rp("C"), None)], largest=True, min_key=0.8)
    CALC["aniv"] = top_list(113, "5d · PRÓXIMOS ANIVERSÁRIOS", Rp("O"), [
        ("Nome", Rp("B"), None), ("Data", Rp("J"), DATE_S),
        ("Faltam", lambda r: f'=IF($B{r}="","",IF(INT($B{r})=0,"▲ hoje",IF(INT($B{r})<=AlertaAniv,"▲ ","")&"em "&INT($B{r})&" d"))', None)])
    Cs = lambda c: A(S.CASA, c, DOC_R1, c, ASS_R2)
    CALC["venc"] = top_list(125, "5e · VENCIMENTOS: DOCUMENTOS, ROTINAS E ASSINATURAS", Cs("N"), [
        ("Item", lambda r: f'=IF($B{r}="","",INDEX({Cs("B")},MATCH($B{r},{Cs("N")},0))&" ("&LOWER(INDEX({Cs("O")},MATCH($B{r},{Cs("N")},0)))&")")', None),
        ("Quando", lambda r: f'=IF($B{r}="","",IF(INT($B{r})<0,"✖ venceu há "&-INT($B{r})&" d",IF(INT($B{r})=0,"▲ hoje","▲ em "&INT($B{r})&" d")))', None)])
    Ap = lambda c: A(S.APR, c, APR_R1, c, APR_R2)
    CALC["apr"] = top_list(137, "5f · APRENDIZADO EM ANDAMENTO", Ap("O"), [
        ("Título", Ap("B"), None), ("Tipo", Ap("C"), None), ("Progresso", Ap("J"), PCT)])

    # ---------------- bloco 6: auxiliares
    section(ws, 149, "B", "N", "6 · AUXILIARES")
    Sd = lambda c: A(S.SAU, c, SAU_R1, c, SAU_R2)
    rowidx = f"(ROW({Sd('L')})-{SAU_R1 - 1})"
    first_w = f'SUMPRODUCT(MIN(({Sd("L")}<>"")*{rowidx}+({Sd("L")}="")*99999))'
    last_note = f'SUMPRODUCT(MAX(({Sd("O")}<>"")*({Sd("B")}<=Corte)*{rowidx}))'
    RD = lambda c: A(S.RODA, c, REV_R1, c, REV_R2)
    aux = [
        (150, "Hoje (conferência)", "=Hoje", DATE),
        (151, "Linha do peso mais recente (até o corte)", f'=SUMPRODUCT(MAX(({Sd("L")}<>"")*({Sd("B")}<=Corte)*{rowidx}))', "0"),
        (152, "Linha da primeira pesagem do ano", f"=IF({first_w}>=99999,0,{first_w})", "0"),
        (153, "Estresse médio no mês", f'=IFERROR(AVERAGEIFS({Sd("H")},{in_month(Sd("B"))}),"")', "0.0"),
        (154, "Energia média no mês", f'=IFERROR(AVERAGEIFS({Sd("G")},{in_month(Sd("B"))}),"")', "0.0"),
        (155, "Linha do último registro com notas", f"={last_note}", "0"),
        (156, "Data do último registro com notas", f'=IF(C155=0,"",INDEX({Sd("B")},C155))', DATE),
        (157, "Texto do último registro", f'=IF(C155=0,"",INDEX({Sd("O")},C155)&"")', None),
        (158, "Maior sequência atual de hábito (dias)", f'=IF(COUNT({A(S.HAB, "D", 11, "M", 11)})=0,0,MAX({A(S.HAB, "D", 11, "M", 11)}))', "0"),
        (159, "Hábito com a maior sequência", f'=IF(C158=0,"",INDEX({A(S.HAB, "D", 16, "M", 16)},MATCH(C158,{A(S.HAB, "D", 11, "M", 11)},0)))', None),
        (160, "Foco do mês (da revisão do mês anterior)", f'=IF(RefMes>1,INDEX({RD("M")},RefMes-1)&"","")', None),
        (161, "Vitórias do mês de referência", f'=INDEX({RD("D")},RefMes)&""', None),
        (162, "Aprendizado do mês de referência", f'=INDEX({RD("J")},RefMes)&""', None),
        (163, "O que não funcionou no mês", f'=INDEX({RD("G")},RefMes)&""', None),
        (164, "Visão de 5 anos", f'={A(S.METAS, "D", 5)}&""', None),
        (165, "Tema do ano", f'={A(S.METAS, "D", 8)}&""', None),
        (166, "Próxima consulta de saúde", f'=IF({K["sau_prox_consulta"]}>=99999,"",{K["sau_prox_consulta"]})', DATE),
        (167, "Alertas de casa e documentos", f'={K["casa_venc"]}+{K["casa_renov"]}+{K["casa_rotatr"]}', "0"),
    ]
    for r, label, fml, fmt in aux:
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = F(9, False, INK2)
        merge(ws, f"C{r}:N{r}" if fmt is None else f"C{r}", fml, F(9, False, INK2), fill(AUTO_BG), al(wrap=fmt is None), BOX, fmt)
    for r in (157, 160, 161, 162, 163, 164):
        ws.row_dimensions[r].height = 30
    CALC["aux"] = {r: A(S.CALC, "C", r) for r, *_ in aux}

    # ---------------- bloco 7: faixas dos últimos 28 dias
    section(ws, 170, "B", "D", "7 · ÚLTIMOS 28 DIAS (até a data de corte)")
    headers(ws, 171, [("B", "Data", "auto"), ("C", "Humor", "auto"), ("D", "Hábitos (% do dia)", "auto")])
    body(ws, "B", "D", 172, 199, "auto", h="center")
    for k in range(28):
        r = 172 + k
        ws[f"B{r}"] = f"=Corte-{27 - k}"
        ws[f"B{r}"].number_format = DATE
        for c, src in (("C", A(S.SAU, "F", SAU_R1, "F", SAU_R2)), ("D", A(S.HAB, "O", HAB_R1, "O", HAB_R2))):
            ws[f"{c}{r}"] = (f'=IF(AND(B{r}>=IniAno,B{r}<=FimAno),IF(INDEX({src},B{r}-IniAno+1)="","",'
                             f'INDEX({src},B{r}-IniAno+1)),"")')
        ws[f"D{r}"].number_format = PCT
    CALC["strip"] = 172
    ws.freeze_panes = "A5"


# =============================================================================
# Dashboard
# =============================================================================
DCOLS = [CL(2 + i) for i in range(28)]   # B..AC
CARD_STARTS = [2, 9, 16, 23]              # B, I, P, W (6 colunas cada + 1 de respiro)


def axis_text(size=800, color=MUTED):
    cp = CharacterProperties(sz=size, solidFill=color, latin=DFont(typeface="Arial"))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text):
    cp = CharacterProperties(sz=1000, b=True, solidFill=INK, latin=DFont(typeface="Arial"))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def style_chart(ch, title, legend=True, pct=False, ymax=None):
    ch.title = chart_title(title)
    ch.style = 2
    ch.width, ch.height = 13.6, 7.6
    if legend:
        ch.legend.position = "t"
        ch.legend.txPr = axis_text(800, INK2)
    else:
        ch.legend = None
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.txPr = axis_text()
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRIDC))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    if pct:
        ch.y_axis.number_format = "0%"
        ch.y_axis.numFmt.sourceLinked = False
    if ymax is not None:
        ch.y_axis.scaling.max = ymax
        ch.y_axis.scaling.min = 0
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def color_series(ch, colors):
    for s, c in zip(ch.series, colors):
        s.graphicalProperties.solidFill = c
        s.graphicalProperties.line.solidFill = c


def col_chart(title, cats, data, colors, titles_from_data=True, pct=False, horizontal=False, ymax=None,
              legend=None, labels=False, fmt=None):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.gapWidth = 60
    ch.overlap = -10
    for d in data:
        if isinstance(d, tuple):
            ref, name = d
            s = Series(ref, title=name)
            ch.series.append(s)
        else:
            ch.add_data(d, titles_from_data=titles_from_data)
    ch.set_categories(cats)
    style_chart(ch, title, legend if legend is not None else len(ch.series) > 1, pct, ymax)
    color_series(ch, colors)
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.crosses = "max"
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showCatName = False
        ch.dataLabels.showSerName = False
        ch.dataLabels.showLegendKey = False
        ch.dataLabels.showPercent = False
        ch.dataLabels.txPr = axis_text(750, INK2)
        if fmt:
            ch.dataLabels.numFmt = fmt
    return ch


def place(ws, ch, top_left, bottom_right):
    """Ancora o gráfico a um intervalo de células: ocupa exatamente o espaço em qualquer programa."""
    def split(ref):
        col = "".join(c for c in ref if c.isalpha())
        return CI(col) - 1, int("".join(c for c in ref if c.isdigit())) - 1
    (c1, r1), (c2, r2) = split(top_left), split(bottom_right)
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=c1, row=r1, colOff=60000, rowOff=40000)
    a.to = AnchorMarker(col=c2, row=r2, colOff=0, rowOff=0)
    ch.anchor = a
    ws.add_chart(ch)


def build_dashboard(wb, ws, sample):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = DARK
    ws.column_dimensions["A"].width = 2
    for c in DCOLS:
        ws.column_dimensions[c].width = 4.9
    ws.column_dimensions["AD"].width = 2
    MAXR = 176
    for r in range(1, MAXR + 1):
        for c in range(1, 31):
            ws.cell(r, c).fill = fill(PAGE)
    calc = lambda c, r: A(S.CALC, c, r)
    X = CALC["aux"]

    # ---- cabeçalho
    for r in (1, 2):
        for c in range(1, 31):
            ws.cell(r, c).fill = fill(DARK)
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 20
    merge(ws, "B1:S1", '=IF(Nome="","PAINEL DA VIDA","PAINEL DA VIDA · "&UPPER(Nome))', F(18, True, WHITE), None, al(v="center"))
    merge(ws, "T1:Y1", "MÊS DE REFERÊNCIA ▸", F(8, True, "C3C2B7"), None, al("right", "center"))
    merge(ws, "Z1:AC1", "Setembro" if sample else "Mês atual", F(11, True, INK), fill(WHITE), al("center", "center"),
                Border(left=side(BLUE, "medium"), right=side(BLUE, "medium"), top=side(BLUE, "medium"), bottom=side(BLUE, "medium")))
    dv_list(ws, "Z1", "lst_MesSel", "Escolha o mês que o painel deve mostrar. 'Mês atual' acompanha a data de hoje.")
    merge(ws, "B2:S2", '="Visão integrada de "&RotuloMes&"  ·  hoje: "&DAY(Hoje)&"/"&MONTH(Hoje)&"/"&YEAR(Hoje)&"  ·  dados até "&DAY(Corte)&"/"&MONTH(Corte)',
          F(9, False, "C3C2B7", italic=True), None, al(v="center"))
    merge(ws, "T2:AC2", '=IF(' + X[165] + '="","",'+'"Tema do ano: "&' + X[165] + ')', F(9, True, "FAB219"), None, al("right", "center"))
    # navegação
    nav = [(S.RODA, "Roda"), (S.METAS, "Metas"), (S.TAR, "Tarefas"), (S.FIN, "Finanças"), (S.ORC, "Orçamento"),
           (S.SAU, "Saúde"), (S.HAB, "Hábitos"), (S.CAR, "Carreira"), (S.REL, "Relações"), (S.APR, "Estudo"),
           (S.LAZ, "Lazer"), (S.CASA, "Casa"), (S.GUIA, "Guia"), (S.CFG, "Config")]
    ws.row_dimensions[3].height = 20
    for i, (sh, label) in enumerate(nav):
        c1, c2 = DCOLS[2 * i], DCOLS[2 * i + 1]
        cell = merge(ws, f"{c1}3:{c2}3", None, None, fill(WHITE), al("center", "center"),
                     Border(left=side(PAGE, "medium"), right=side(PAGE, "medium"), bottom=side(DOMAIN[sh], "thick")))
        link(cell, sh, label, DARK, 8, True)
        cell.font = F(8, True, DARK)
    ws.row_dimensions[4].height = 10

    def sec(row, text):
        section(ws, row, "B", "AC", text, DARK)
        for c in range(2, 30):
            ws.cell(row, c).fill = fill(PAGE)
        ws[f"B{row}"].font = F(11, True, DARK)

    def card(row, idx, color, label, value, fmt, sub, value_size=20):
        c1, c2 = CL(CARD_STARTS[idx]), CL(CARD_STARTS[idx] + 5)
        left = side(color, "thick")
        merge(ws, f"{c1}{row}:{c2}{row}", label, F(7.5, True, MUTED), fill(WHITE), al("left", "bottom", indent=1))
        merge(ws, f"{c1}{row+1}:{c2}{row+1}", value, F(value_size, True, INK), fill(WHITE),
              al("left", "center", wrap=value_size < 14, indent=1), fmt=fmt)
        merge(ws, f"{c1}{row+2}:{c2}{row+2}", sub, F(8, False, INK2), fill(WHITE), al("left", "top", wrap=True, indent=1))
        for r in range(row, row + 3):
            for c in range(CI(c1), CI(c2) + 1):
                ws.cell(r, c).border = Border(left=left if c == CI(c1) else None,
                                              right=side(LINE) if c == CI(c2) else None,
                                              top=side(LINE) if r == row else None,
                                              bottom=side(LINE) if r == row + 2 else None)

    def dash(ref, alt="–"):
        return f'=IF(ISNUMBER({ref}),{ref},"{alt}")'

    # ---- visão geral
    sec(5, '="VISÃO GERAL · "&UPPER(RotuloMes)')
    ws["B5"] = '="VISÃO GERAL · "&UPPER(RotuloMes)'
    idx, per = CALC["indice"], CALC["percep"]
    rows_cards = [6, 10, 14, 18]
    for r in rows_cards:
        ws.row_dimensions[r].height = 15
        ws.row_dimensions[r + 1].height = 30
        ws.row_dimensions[r + 2].height = 26
        ws.row_dimensions[r + 3].height = 9
    met_and, met_ritmo, met_atras, met_prog = K["met_and"], K["met_ritmo"], K["met_atras"], K["met_prog"]
    an7 = f'COUNTIFS({A(S.REL, "K", PPL_R1, "K", PPL_R2)},">=0",{A(S.REL, "K", PPL_R1, "K", PPL_R2)},"<=7")'
    ve7 = f'COUNTIFS({A(S.CASA, "N", DOC_R1, "N", ASS_R2)},">=0",{A(S.CASA, "N", DOC_R1, "N", ASS_R2)},"<8")'
    co7 = f'COUNTIFS({A(S.SAU, "W", 27, "W", 56)},">=0",{A(S.SAU, "W", 27, "W", 56)},"<=7")'
    nx7 = f'{K["tar_7d"]}+{an7}+{ve7}+{co7}'
    cards = [
        # linha A · visão
        (6, 0, DARK, "ÍNDICE DE VIDA · DADOS", dash(idx), '0" /100"',
         f'=IF(AND(ISNUMBER({idx}),ISNUMBER({per})),IF(ABS({per}*10-{idx})<=10,"✔ ","▲ "),"")&"Percepção "&IF(ISNUMBER({per}),FIXED({per},1)&"/10","–")'
         f'&IF(AND(ISNUMBER({idx}),ISNUMBER({per})),IF(ABS({per}*10-{idx})<=10," · alinhada aos dados",IF({per}*10>{idx}," · acima dos dados"," · abaixo dos dados")),"")'),
        (6, 1, BLUE, "METAS NO RITMO", f'=IF({met_and}=0,"–",{met_ritmo}&" de "&{met_and})', None,
         f'=IF({met_and}=0,"○ Cadastre suas metas na aba Metas",IF({met_atras}>0,"▲ "&{met_atras}&" atrasada(s) · ","✔ ")&"progresso médio "&IF(ISNUMBER({met_prog}),ROUND({met_prog}*100,0)&"%","–"))'),
        (6, 2, BLUE, "TAREFAS ATRASADAS", f"={K['tar_atras']}", "0",
         f'=IF({K["tar_atras"]}>0,"✖ ","✔ ")&{K["tar_abertas"]}&" abertas · "&{K["tar_concl"]}&" concluídas no mês"'),
        (6, 3, BLUE, "PRÓXIMOS 7 DIAS", f"={nx7}", '0" item(ns)"',
         f'={K["tar_7d"]}&" tarefa(s) · "&{an7}&" aniversário(s) · "&{ve7}&" vencimento(s) · "&{co7}&" consulta(s)"'),
        # linha B · finanças
        (10, 0, AQUA, "RESULTADO DO MÊS", f"={K['fin_res']}", EUR0,
         f'=IF({K["fin_rec"]}+{K["fin_desp"]}=0,"○ Sem lançamentos no mês",IF({K["fin_res"]}>=0,"✔ ","✖ ")&"Receitas € "&FIXED({K["fin_rec"]},0)&" · despesas € "&FIXED({K["fin_desp"]},0))'),
        (10, 1, AQUA, "TAXA DE POUPANÇA", f"={K['fin_taxa']}", PCT, f"={K['fin_taxa_sub']}"),
        (10, 2, AQUA, "PATRIMÔNIO LÍQUIDO", dash(K["pl_ref"]), EUR0,
         f'=IF(ISNUMBER({K["pl_var"]}),IF({K["pl_var"]}>=0,"✔ +€ ","✖ −€ ")&FIXED(ABS({K["pl_var"]}),0)&" no mês","○ Atualize os saldos em Orçamento")'),
        (10, 3, AQUA, "RESERVA DE EMERGÊNCIA", dash(K["reserva_meses"]), '0.0" meses"',
         f'=IF(ISNUMBER({K["reserva_meses"]}),IF({K["reserva_meses"]}>=MetaReserva,"✔ ","▲ ")&"meta "&MetaReserva&" meses de despesas","○ sem dados")'),
        # linha C · saúde & hábitos
        (14, 0, ORANGE, "HÁBITOS · % DA META", dash(K["hab_meta"]), PCT,
         f'=IF({X[158]}>0,"Maior sequência: "&{X[159]}&" ("&{X[158]}&" dias)","Marque seus hábitos na aba Hábitos")'),
        (14, 1, ORANGE, "SONO MÉDIO", f"={K['sau_sono']}", '0.0" h"', f"={K['sau_sono_sub']}"),
        (14, 2, ORANGE, "TREINOS NO MÊS", f"={K['sau_treinos']}", "0", f"={K['sau_treinos_sub']}"),
        (14, 3, ORANGE, "HUMOR MÉDIO", f"={K['sau_humor']}", '0.0"/5"', f"={K['sau_humor_sub']}"),
        # linha D · pessoas, crescimento, lazer, casa
        (18, 0, MAGENTA, "RELAÇÕES EM DIA", f"={K['rel_emdia']}", PCT,
         f'=IF({K["rel_n"]}=0,"○ Cadastre pessoas na aba Relações",IF({K["rel_atras"]}>0,"▲ ","✔ ")&{K["rel_atras"]}&" contato(s) atrasado(s) · "&{K["rel_inter"]}&" no mês")'),
        (18, 1, VIOLET, "ESTUDO NO MÊS", f"={K['apr_h']}", '0.0" h"', f"={K['apr_h_sub']}"),
        (18, 2, MAGENTA, "LAZER NO MÊS", f"={K['laz_h']}", '0.0" h"', f"={K['laz_h_sub']}"),
        (18, 3, YELLOW, "CASA & DOCUMENTOS", f"={X[167]}", '0" alerta(s)"',
         f'=IF({X[167]}>0,"▲ ","✔ ")&({K["casa_venc"]}+{K["casa_renov"]})&" doc(s) · "&{K["casa_rotatr"]}&" rotina(s) atrasada(s)"'),
    ]
    for r, i, color, label, val, fmt, sub in cards:
        card(r, i, color, label, val, fmt, sub)
    ws.conditional_formatting.add("B7:G7", FormulaRule(formula=["AND(ISNUMBER($B$7),$B$7>=70)"], font=Font(color=GOOD_T)))
    ws.conditional_formatting.add("B7:G7", FormulaRule(formula=["AND(ISNUMBER($B$7),$B$7<50)"], font=Font(color=CRIT_T)))
    ws.conditional_formatting.add("B11:G11", FormulaRule(formula=["AND(ISNUMBER($B$11),$B$11<0)"], font=Font(color=CRIT_T)))

    # ---- áreas da vida
    sec(22, "ÁREAS DA VIDA · percepção (sua nota) × dados (o que os registros mostram)")
    hdr = [("B", "F", "Área"), ("G", "H", "Percepção"), ("I", "J", "Dados"), ("K", "K", "Tend."), ("L", "R", "Indicador-chave do mês")]
    for c1, c2, t in hdr:
        merge(ws, f"{c1}23:{c2}23", t, F(8, True, WHITE), fill(DARK), al("center" if c1 != "L" else "left", "center", indent=0 if c1 != "L" else 1))
    ws.row_dimensions[23].height = 20
    for i in range(11):
        r = 24 + i
        cr = 5 + i
        ws.row_dimensions[r].height = 21
        merge(ws, f"B{r}:F{r}", f"={calc('B', cr)}", F(9, True, INK), fill(WHITE), al(indent=1), Border(bottom=side(GRIDC)))
        merge(ws, f"G{r}:H{r}", f"={calc('I', cr)}", F(9, True, INK), fill(WHITE), al("center"), Border(bottom=side(GRIDC)),
              '[>=7]"✔ "0.0;[>=5]"▲ "0.0;"✖ "0.0')
        merge(ws, f"I{r}:J{r}", f"={calc('H', cr)}", F(9, True, INK), fill(WHITE), al("center"), Border(bottom=side(GRIDC)),
              '[>=7]"✔ "0.0;[>=5]"▲ "0.0;"✖ "0.0')
        merge(ws, f"K{r}", f"={A(S.RODA, 'S', RODA_R1 + i)}", F(11, True, INK), fill(WHITE), al("center"), Border(bottom=side(GRIDC)))
        merge(ws, f"L{r}:R{r}", f"={calc('M', cr)}", F(8, False, INK2), fill(WHITE), al(indent=1, wrap=True), Border(bottom=side(GRIDC)))
    for rng in ("G24:H34", "I24:J34"):
        tl = rng.split(":")[0]
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({tl}),{tl}>=7)"], font=Font(color=GOOD_T), fill=fill(GOOD_BG)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({tl}),{tl}>=5,{tl}<7)"], font=Font(color=WARN_T), fill=fill(WARN_BG)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({tl}),{tl}<5)"], font=Font(color=CRIT_T), fill=fill(CRIT_BG)))
    ws.conditional_formatting.add("K24:K34", FormulaRule(formula=['K24="↑"'], font=Font(color=GOOD_T)))
    ws.conditional_formatting.add("K24:K34", FormulaRule(formula=['K24="↓"'], font=Font(color=CRIT_T)))
    merge(ws, "B35:R35", "Percepção = sua nota na Roda da Vida · Dados = placar calculado (metas, hábitos e métricas da área). "
                         "Diferença grande entre os dois merece atenção. Detalhe na aba Cálculos.", F(7.5, False, MUTED, italic=True),
          None, al(wrap=True, v="top"))
    ws.row_dimensions[35].height = 22
    radar = RadarChart()
    radar.type = "marker"
    radar.style = 2
    data = Reference(wb[S.CALC], min_col=CI("H"), max_col=CI("J"), min_row=4, max_row=15)
    radar.add_data(data, titles_from_data=True)
    radar.set_categories(Reference(wb[S.CALC], min_col=2, min_row=5, max_row=15))
    radar.title = chart_title("Roda da Vida")
    radar.width, radar.height = 10.6, 8.2
    radar.y_axis.scaling.min, radar.y_axis.scaling.max = 0, 10
    radar.y_axis.majorUnit = 2
    radar.y_axis.delete = False
    radar.y_axis.txPr = axis_text(700)
    radar.x_axis.txPr = axis_text(750, INK2)
    radar.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRIDC))
    radar.legend.position = "b"
    radar.legend.txPr = axis_text(800, INK2)
    for s, c, w, dash_ in zip(radar.series, [ORANGE, BLUE, MUTED], [22000, 28000, 15000], [None, None, "dash"]):
        s.graphicalProperties.line.solidFill = c
        s.graphicalProperties.line.width = w
        if dash_:
            s.graphicalProperties.line.prstDash = dash_
        s.marker.symbol = "circle" if not dash_ else "none"
        s.marker.size = 5
        s.marker.graphicalProperties = GraphicalProperties(solidFill=c, ln=LineProperties(solidFill=WHITE))
        s.smooth = False
    place(ws, radar, "S22", "AD36")

    # ---- foco & reflexão
    sec(37, "FOCO & REFLEXÃO")

    def textbox(r1, r2, c1, c2, title, value, color, editable=False):
        merge(ws, f"{c1}{r1}:{c2}{r1}", title, F(8, True, MUTED), fill(WHITE), al(indent=1, v="bottom"),
              Border(top=side(color, "thick")))
        merge(ws, f"{c1}{r1+1}:{c2}{r2}", value, F(9.5, False, INK, italic=not editable), fill(WHITE),
              al("left", "top", True, 1), Border(bottom=side(LINE)))

    ws.row_dimensions[38].height = 16
    textbox(38, 41, "B", "N", "FOCO DO MÊS · da sua revisão mensal",
            f'=IF({X[160]}="","Ainda sem foco definido. No fim do mês, escreva em Roda da Vida → \'Foco do próximo mês\'.",{X[160]})', BLUE)
    merge(ws, "P38:AC38", "PRIORIDADES DA SEMANA · escreva aqui", F(8, True, MUTED), fill(WHITE), al(indent=1, v="bottom"),
          Border(top=side(BLUE, "thick")))
    for k, r in enumerate((39, 40, 41)):
        merge(ws, f"P{r}", f"{k + 1}", F(9, True, BLUE), fill(WHITE), al("center"))
        merge(ws, f"Q{r}:AC{r}", None, F(10, False, INK), fill(WHITE), al(indent=1),
              Border(bottom=side(LINE), left=side(LINE)))
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[42].height = 9
    textbox(43, 46, "B", "N", "VITÓRIAS DO MÊS · e o que aprendi",
            f'=IF(AND({X[161]}="",{X[162]}=""),"Preencha a revisão de "&RotuloMes&" na aba Roda da Vida.",{X[161]}&IF({X[162]}="","","  ·  Aprendi: "&{X[162]}))',
            GOOD_T)
    textbox(43, 46, "P", "AC", "ÚLTIMO REGISTRO DO DIÁRIO",
            f'=IF({X[157]}="","Use a coluna Notas da aba Saúde para registrar como foi o dia.",DAY({X[156]})&"/"&MONTH({X[156]})&" — "&{X[157]})',
            ORANGE)
    for r in range(39, 47):
        if r != 42:
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 18)
    textbox(48, 50, "B", "AC", "MINHA VISÃO · daqui a 5 anos",
            f'=IF({X[164]}="","Escreva sua visão no topo da aba Metas — é o norte de todas as outras abas.",{X[164]})', VIOLET)
    ws.row_dimensions[51].height = 9

    # ---- atenção agora
    sec(52, "ATENÇÃO AGORA")

    def list_block(r0, c1, c2, title, color, cols, src_row, n=8, bar_cols=()):
        """cols: [(c_ini, c_fim, cabeçalho, coluna_no_Cálculos, formato)]"""
        merge(ws, f"{c1}{r0}:{c2}{r0}", title, F(8.5, True, INK), fill(WHITE), al(indent=1, v="center"),
              Border(top=side(color, "thick")))
        for a, b, h, _, _ in cols:
            merge(ws, f"{a}{r0+1}:{b}{r0+1}", h, F(7.5, True, MUTED), fill(WHITE), al("left", indent=1),
                  Border(bottom=side(LINE)))
        for k in range(n):
            r = r0 + 2 + k
            for a, b, h, src, fmt in cols:
                merge(ws, f"{a}{r}:{b}{r}", f'={A(S.CALC, src, src_row + k)}', F(8, False, INK), fill(WHITE),
                      al("left", indent=1), Border(bottom=side(GRIDC)), fmt)
        for a, b, h, src, fmt in cols:
            if (a, b) in bar_cols:
                cf_bar(ws, f"{a}{r0+2}:{a}{r0+1+n}", 0, 1, color)
        first = f"{cols[0][0]}{r0+2}"
        merge(ws, f"{c1}{r0+2+n}:{c2}{r0+2+n}", f'=IF({first}="","Nada por aqui ✔","")', F(8, False, MUTED, italic=True),
              fill(WHITE), al("left", indent=1))

    ws.row_dimensions[53].height = 18
    list_block(53, "B", "O", "Tarefas atrasadas e próximas", BLUE,
               [("B", "H", "Tarefa", "C", None), ("I", "J", "Prazo", "D", DATE_S), ("K", "M", "Situação", "E", None),
                ("N", "O", "Projeto", "F", None)], CALC["tar"])
    list_block(53, "P", "V", "Contatos a fazer", MAGENTA,
               [("P", "R", "Pessoa", "C", None), ("S", "T", "Sem contato", "D", None), ("U", "V", "Status", "E", None)],
               CALC["cont"])
    list_block(53, "W", "AC", "Vencimentos & rotinas", YELLOW,
               [("W", "Z", "Item", "C", None), ("AA", "AC", "Quando", "D", None)], CALC["venc"])
    list_block(65, "B", "O", "Metas em andamento · prazo mais próximo", BLUE,
               [("B", "G", "Meta", "C", None), ("H", "I", "Progresso", "E", PCT), ("J", "K", "Esperado", "F", PCT),
                ("L", "M", "Ritmo", "G", None), ("N", "O", "Prazo", "H", DATE_S)], CALC["met"], bar_cols=[("H", "I")])
    list_block(65, "P", "V", "Próximos aniversários", MAGENTA,
               [("P", "R", "Pessoa", "C", None), ("S", "T", "Data", "D", DATE_S), ("U", "V", "Faltam", "E", None)],
               CALC["aniv"])
    list_block(65, "W", "AC", "Aprendendo agora", VIOLET,
               [("W", "AA", "Título", "C", None), ("AB", "AC", "Progresso", "E", PCT)], CALC["apr"], bar_cols=[("AB", "AC")])
    for r in range(53, 77):
        ws.row_dimensions[r].height = ws.row_dimensions[r].height or 17

    # ---- finanças
    C = wb[S.CALC]
    months = Reference(C, min_col=2, min_row=20, max_row=31)
    sec(78, '="FINANÇAS · "&Ano')
    ws["B78"] = '="FINANÇAS · ANO "&Ano'
    ch = col_chart("Receitas × despesas por mês (€)", months,
                   [Reference(C, min_col=3, max_col=4, min_row=19, max_row=31)], [BLUE, ORANGE])
    place(ws, ch, "B79", "P95")
    ch = col_chart("Despesas do mês por categoria (top 8, €)", Reference(C, min_col=7, min_row=36, max_row=43),
                   [(Reference(C, min_col=9, min_row=36, max_row=43), "Despesa")], [ORANGE], horizontal=True,
                   legend=False, labels=True, fmt="#,##0")
    place(ws, ch, "P79", "AD95")
    r0 = 95
    merge(ws, f"B{r0}:AC{r0}", "DISTRIBUIÇÃO DA RECEITA NO MÊS · referência 50/30/20", F(8, True, MUTED), None, al(indent=0, v="bottom"))
    dist = [(0, "ESSENCIAL · ref. ≤ 50%", 90), (1, "ESTILO DE VIDA · ref. ≤ 30%", 91), (2, "CRESCIMENTO · cursos, livros", 92),
            (3, "POUPANÇA · ref. ≥ 20%", 93)]
    for i, label, orow in dist:
        card(r0 + 1, i, AQUA, label, f'=IF({A(S.ORC, "D", orow)}="","–",{A(S.ORC, "D", orow)})', PCT,
             f'={A(S.ORC, "G", orow)}', 16)
    ws.row_dimensions[r0 + 1].height = 14
    ws.row_dimensions[r0 + 2].height = 24
    ws.row_dimensions[r0 + 3].height = 16
    ch = col_chart("Patrimônio líquido no fim do mês (€)", months,
                   [(Reference(C, min_col=8, min_row=20, max_row=31), "Patrimônio")], [AQUA], legend=False)
    place(ws, ch, "B100", "P116")
    ch = col_chart("Taxa de poupança por mês", months,
                   [(Reference(C, min_col=7, min_row=20, max_row=31), "Taxa de poupança")], [AQUA], pct=True, legend=False)
    place(ws, ch, "P100", "AD116")

    # ---- saúde & hábitos
    sec(116, "SAÚDE & HÁBITOS")
    s0 = CALC["strip"]
    merge(ws, "B117:AC117", "HUMOR NOS ÚLTIMOS 28 DIAS  ·  vermelho = dia ruim · cinza = neutro · azul = dia bom · vazio = sem registro",
          F(8, True, MUTED), None, al(v="bottom"))
    merge(ws, "B120:AC120", "HÁBITOS CUMPRIDOS POR DIA  ·  quanto mais escuro, mais hábitos feitos naquele dia",
          F(8, True, MUTED), None, al(v="bottom"))
    for k, c in enumerate(DCOLS):
        r = s0 + k
        for drow, vrow, src in ((118, 119, "C"), (121, 122, "D")):
            d = ws[f"{c}{drow}"]
            d.value = f"=DAY({A(S.CALC, 'B', r)})"
            d.font = F(7, False, MUTED)
            d.alignment = al("center")
            v = ws[f"{c}{vrow}"]
            v.value = f"={A(S.CALC, src, r)}"
            v.number_format = ";;;"
            v.fill = fill("FCFCFB")
            v.border = Border(left=side(PAGE, "medium"), right=side(PAGE, "medium"))
    ws.row_dimensions[119].height = ws.row_dimensions[122].height = 24
    ws.conditional_formatting.add("B119:AC119", ColorScaleRule(start_type="num", start_value=1, start_color="E34948",
                                                                mid_type="num", mid_value=3, mid_color="E6E4DE",
                                                                end_type="num", end_value=5, end_color=BLUE))
    ws.conditional_formatting.add("B122:AC122", ColorScaleRule(start_type="num", start_value=0, start_color=SEQ_LO,
                                                                end_type="num", end_value=1, end_color=SEQ_HI))
    ch = col_chart("Sono médio por mês (h)", months, [(Reference(C, min_col=9, min_row=20, max_row=31), "Sono")],
                   [ORANGE], legend=False, ymax=10)
    place(ws, ch, "B124", "P140")
    ch = col_chart("Treinos por mês", months, [(Reference(C, min_col=10, min_row=20, max_row=31), "Treinos")],
                   [ORANGE], legend=False)
    place(ws, ch, "P124", "AD140")
    ch = col_chart("Hábitos no mês de referência · % da meta", Reference(C, min_col=2, min_row=65, max_row=74),
                   [(Reference(C, min_col=3, min_row=65, max_row=74), "% da meta")], [ORANGE], pct=True, horizontal=True,
                   legend=False, ymax=1, labels=True, fmt="0%")
    place(ws, ch, "B140", "P156")
    ch = col_chart("Consistência de hábitos por mês (% de marcações possíveis)", months,
                   [(Reference(C, min_col=12, min_row=20, max_row=31), "Hábitos")], [ORANGE], pct=True, legend=False, ymax=1)
    place(ws, ch, "P140", "AD156")

    # ---- metas, estudo & lazer
    sec(156, "METAS, ESTUDO & LAZER")
    m0 = CALC["met"]
    ch = col_chart("Metas: progresso × esperado pelo tempo", Reference(C, min_col=3, min_row=m0, max_row=m0 + 7),
                   [(Reference(C, min_col=5, min_row=m0, max_row=m0 + 7), "Progresso"),
                    (Reference(C, min_col=6, min_row=m0, max_row=m0 + 7), "Esperado hoje")],
                   [BLUE, "C3C2B7"], pct=True, horizontal=True, ymax=1)
    place(ws, ch, "B157", "P173")
    ch = col_chart("Horas de estudo e de lazer por mês", months,
                   [Reference(C, min_col=13, max_col=14, min_row=19, max_row=31)], [VIOLET, MAGENTA])
    place(ws, ch, "P157", "AD173")

    cf_status(ws, f"B5:AC{MAXR}")
    ws.freeze_panes = "A4"
    ws.print_area = f"A1:AD{MAXR}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# =============================================================================
# Guia
# =============================================================================
def build_guia(ws, sample):
    sheet_title(ws, "Guia · como usar o sistema",
                "Leia uma vez. Depois, o ritmo é: 2 minutos por dia, 20 por semana, 30 por mês.", "H", legend=False)
    setw(ws, {"B": 22, "C": 30, "D": 30, "E": 30, "F": 22, "G": 14, "H": 14})
    r = 5

    def para(text, size=9, bold=False, color=INK, h=None, italic=False):
        nonlocal r
        merge(ws, f"B{r}:H{r}", text, F(size, bold, color, italic=italic), None, al(wrap=True, v="top"))
        ws.row_dimensions[r].height = h or max(16, 14 * (1 + len(text) // 150))
        r += 1

    def table(head, rows, widths=None):
        nonlocal r
        spans = widths or [("B", "B"), ("C", "D"), ("E", "H")]
        for (a, b), t in zip(spans, head):
            merge(ws, f"{a}{r}:{b}{r}", t, F(9, True, WHITE), fill(DARK), al(indent=1))
        r += 1
        for row in rows:
            for (a, b), t in zip(spans, row):
                merge(ws, f"{a}{r}:{b}{r}", t, F(9, a == "B", INK), fill(WHITE), al(wrap=True, indent=1, v="top"), BOX)
            ws.row_dimensions[r].height = max(18, 13 * (1 + max(len(str(x)) for x in row) // 70))
            r += 1
        r += 1

    if sample:
        para("ESTE É O ARQUIVO DE EXEMPLO: todos os dados são fictícios e servem para mostrar o sistema funcionando. "
             "Para usar de verdade, abra o arquivo sem '_EXEMPLO' no nome (mesma estrutura, vazio).", 10, True, CRIT_T, 30)
        r += 1
    section(ws, r, "B", "H", "COMECE AQUI · 10 minutos")
    r += 1
    for t in ["1. Config: escreva seu nome, confira o ano e ajuste suas metas (poupança, sono, treinos, estudo, lazer…).",
              "2. Metas: escreva sua visão de 5 anos e 3 a 7 metas mensuráveis, cada uma ligada a uma área da vida.",
              "3. Hábitos: dê nome a até 10 hábitos (linha azul), escolha a área e a meta de dias por semana.",
              "4. Orçamento: defina quanto pretende gastar por categoria (coluna 'Orçamento / mês').",
              "5. Relações: cadastre as pessoas importantes e de quantos em quantos dias quer falar com cada uma.",
              "6. Casa: cadastre documentos com validade, rotinas recorrentes e assinaturas.",
              "7. Roda da Vida: dê sua primeira nota (0 a 10) para cada área. Pronto: o Dashboard já tem um ponto de partida."]:
        para(t)
    r += 1
    section(ws, r, "B", "H", "RITMO DE USO")
    r += 1
    table(["Quando", "O que fazer", "Onde"], [
        ("Todo dia · 2 min", "Preencher a linha de hoje: sono, humor, energia, estresse, treino e uma nota curta. Marcar x nos hábitos cumpridos.",
         "Saúde · Hábitos (use o link 'Ir para a linha de hoje')"),
        ("Ao gastar ou receber", "Lançar a transação (ou acumule e lance 1× por semana a partir do extrato).", "Finanças"),
        ("Toda semana · 20 min", "Revisar tarefas e prazos, escrever as 3 prioridades da semana no Dashboard, registrar contatos, horas de estudo e lazer.",
         "Dashboard · Tarefas · Relações · Aprendizado · Lazer"),
        ("Fim do mês · 30 min", "Atualizar saldos do patrimônio, dar nota para cada área, escrever a revisão mensal (vitórias, o que não funcionou, aprendizado, foco do próximo mês) e atualizar o valor atual das metas.",
         "Orçamento · Roda da Vida · Metas"),
        ("Fim do ano", "Duplicar o arquivo, apagar os dados, mudar o ano em Config e começar o ano seguinte.", "Arquivo · Config"),
    ])
    section(ws, r, "B", "H", "MAPA DAS ABAS")
    r += 1
    table(["Aba", "Para que serve", "O que alimenta no Dashboard"], [
        ("Dashboard", "Tudo de uma vez: 16 indicadores, placar das 11 áreas, alertas, reflexão e gráficos. Escolha o mês no topo.", "—"),
        ("Roda da Vida", "Nota mensal (0–10) por área + revisão mensal qualitativa.", "Percepção, tendência, foco do mês, vitórias"),
        ("Metas", "Visão, propósito e metas com início, prazo e valores. Calcula progresso, ritmo esperado e horizonte (curto/médio/longo).", "Metas no ritmo, lista de metas, gráfico progresso × esperado"),
        ("Tarefas", "Projetos (inclui banco de ideias) e tarefas com prioridade, prazo e alerta automático (matriz de Eisenhower).", "Tarefas atrasadas, lista de próximas"),
        ("Finanças", "Lançamentos de receitas, despesas e aportes com categoria e forma de pagamento.", "Resultado, taxa de poupança, despesas por categoria"),
        ("Orçamento", "Planejado × realizado por categoria e por mês, patrimônio mensal e distribuição 50/30/20.", "Patrimônio, reserva de emergência, gráficos anuais"),
        ("Saúde", "Registro diário (sono, humor, energia, estresse, treino, passos, peso) + consultas e exames.", "Sono, treinos, humor, faixa de humor de 28 dias"),
        ("Hábitos", "Grade anual de hábitos com meta semanal, sequência atual e melhor sequência.", "Consistência, faixa de 28 dias, gráficos"),
        ("Carreira", "Situação atual, competências (atual × alvo), candidaturas e diário de conquistas.", "Placar de carreira"),
        ("Relações", "Pessoas importantes com frequência desejada de contato, aniversários e registro de contatos.", "Contatos em dia, contatos a fazer, aniversários"),
        ("Aprendizado", "Livros, cursos, idiomas e sessões de estudo.", "Horas de estudo, livros no ano, 'aprendendo agora'"),
        ("Lazer", "Registro de lazer (tempo, custo, satisfação) e lista de sonhos.", "Horas de lazer, satisfação"),
        ("Casa", "Documentos com validade, rotinas recorrentes e assinaturas.", "Vencimentos, alertas de casa"),
        ("Config", "Seus parâmetros, metas e todas as listas dos menus.", "Tudo"),
        ("Cálculos", "Motor transparente: mostra como cada número do Dashboard é calculado.", "Tudo"),
    ])
    section(ws, r, "B", "H", "LEGENDA")
    r += 1
    leg = [("Branco", "Célula para você preencher.", WHITE), ("Cinza", "Cálculo automático. Não digite por cima.", AUTO_BG),
           ("Azul-claro", "Cabeçalho editável (nomes dos hábitos).", EDIT_HDR), ("Amarelo-claro", "Linha de hoje / mês de referência.", TODAY_BG)]
    for name, desc, color in leg:
        merge(ws, f"B{r}", name, F(9, True, INK), fill(color), al(indent=1), BOX)
        merge(ws, f"C{r}:H{r}", desc, F(9), None, al(indent=1))
        r += 1
    for sym, desc, color in [("✔", "ok / no alvo / em dia", GOOD_T), ("▲", "atenção / vence em breve / abaixo do ritmo", WARN_T),
                             ("✖", "crítico / atrasado / vencido", CRIT_T), ("○", "sem dados ou neutro", MUTED),
                             ("↑ ↓ →", "tendência da nota da área em relação ao mês anterior", INK)]:
        merge(ws, f"B{r}", sym, F(11, True, color), None, al("center"))
        merge(ws, f"C{r}:H{r}", desc, F(9), None, al(indent=1))
        r += 1
    para("Cores das abas: azul = planejamento (Roda, Metas, Tarefas) · verde-água = dinheiro · laranja = corpo e rotina · "
         "violeta = crescimento (Carreira, Aprendizado) · rosa = pessoas e lazer · amarelo = casa · cinza = sistema.", 8, color=INK2, h=26)
    r += 1
    section(ws, r, "B", "H", "COMO O 'ÍNDICE DE VIDA' É CALCULADO")
    r += 1
    for t in ["Cada uma das 11 áreas recebe um placar 'Dados' de 0 a 10 = média dos componentes disponíveis, cada um de 0 a 1: "
              "(a) % das metas da área no ritmo, (b) % da meta dos hábitos ligados à área e (c) até 3 métricas próprias "
              "(ex.: Finanças = taxa de poupança vs. meta, reserva vs. meta, categorias dentro do orçamento).",
              "Componente sem dados é ignorado — nunca conta como zero. O Índice de vida (0–100) é a média dos placares × 10.",
              "A 'Percepção' é a sua nota na Roda da Vida. Comparar os dois é o ponto: se você se sente pior do que os dados mostram "
              "(ou melhor), há algo que vale examinar. A aba Cálculos mostra cada componente, área por área."]:
        para(t, h=30)
    r += 1
    section(ws, r, "B", "H", "ESPAÇO E PERSONALIZAÇÃO")
    r += 1
    table(["Tabela", "Capacidade pronta", "Como ampliar"], [
        ("Finanças", "1.500 lançamentos", "Insira linhas ANTES da última linha da tabela e copie as fórmulas das colunas cinza."),
        ("Tarefas · Projetos", "300 tarefas · 20 projetos", "Idem. Apague tarefas concluídas antigas se quiser enxugar."),
        ("Metas", "60 metas", "Idem."),
        ("Relações", "60 pessoas · 600 contatos", "Idem."),
        ("Aprendizado · Lazer", "80 itens · 500 sessões · 400 atividades · 50 sonhos", "Idem."),
        ("Casa", "25 documentos · 35 rotinas · 25 assinaturas", "Idem."),
        ("Saúde · Hábitos", "1 linha por dia do ano · 10 hábitos", "Datas são geradas a partir do ano em Config."),
        ("Listas dos menus", "até 35 itens por lista", "Escreva novos itens embaixo da lista na aba Config, sem deixar linhas vazias no meio."),
    ])
    para("Moeda: os valores usam o formato €. Para outra moeda, selecione as colunas de valor e troque o formato de número.", 8, color=INK2)
    para("Fórmulas compatíveis com Excel 2010 ou mais recente (Windows/Mac), Excel na web e LibreOffice. Não há macros.", 8, color=INK2)


# =============================================================================
# Dados de exemplo (fictícios)
# =============================================================================
def put(ws, row, cols_values):
    for col, v in cols_values.items():
        if v is not None and v != "":
            ws[f"{col}{row}"] = v


def fill_sample(wb, ano):
    import random
    rnd = random.Random(7)
    D = lambda m, d: dt.date(ano, m, d)
    last_day = D(10, 1)

    # ---------- Config
    wb[S.CFG]["C6"] = "Vinícius"
    wb[S.CFG]["C14"] = 78

    # ---------- Finanças
    rows = []

    def add(d, tipo, cat, desc, val, conta, fv, nota=None):
        if d <= last_day:
            rows.append((d, tipo, cat, desc, round(val, 2), conta, fv, nota))

    for m in range(1, 11):
        dim = (dt.date(ano + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).day
        dd = lambda x: D(m, min(x, dim))
        add(dd(1), "Despesa", "Moradia", "Aluguel + condomínio", 650, "PIX / transferência", "Fixo")
        if m == 10:
            continue
        add(dd(10), "Receita", "Salário", "Salário líquido", 2450, "Conta corrente", "Fixo")
        add(dd(28), "Receita", "Rendimentos", "Rendimento da reserva", 12 + m * 0.9, "Conta corrente", "Variável")
        if m in (3, 5, 8):
            add(dd(20), "Receita", "Renda extra / freelance", "Projeto BIM freelance", {3: 380, 5: 520, 8: 450}[m], "PIX / transferência", "Variável")
        if m == 6:
            add(dd(14), "Receita", "Reembolsos", "Reembolso de exame (seguro)", 85, "Conta corrente", "Variável")
        add(dd(2), "Despesa", "Transporte", "Assinatura transporte público", 38, "Cartão de débito", "Fixo")
        add(dd(3), "Despesa", "Assinaturas", "Streaming", 9.99, "Cartão de crédito", "Fixo")
        add(dd(6), "Despesa", "Assinaturas", "Música", 10.99, "Cartão de crédito", "Fixo")
        add(dd(4), "Despesa", "Saúde & bem-estar", "Academia", 34.9, "Cartão de crédito", "Fixo")
        add(dd(5), "Despesa", "Contas da casa", "Internet fibra", 27.95, "Conta corrente", "Fixo")
        add(dd(8), "Despesa", "Contas da casa", "Celular", 7.99, "Conta corrente", "Fixo")
        add(dd(12), "Despesa", "Contas da casa", "Luz", rnd.uniform(38, 72), "Conta corrente", "Variável")
        add(dd(15), "Despesa", "Contas da casa", "Gás", rnd.uniform(70, 95) if m in (1, 2, 3, 11, 12) else rnd.uniform(18, 30),
            "Conta corrente", "Variável")
        add(dd(9), "Despesa", "Dívidas & financiamentos", "Financiamento estudantil (parcela)", 90.83, "Conta corrente", "Fixo")
        add(dd(10), "Despesa", "Família & ajudas", "Ajuda à família", 75, "PIX / transferência", "Fixo")
        add(dd(7), "Despesa", "Educação & cursos", "Mensalidade Master BIM", 180, "Conta corrente", "Fixo")
        if m <= 6:
            add(dd(7), "Despesa", "Educação & cursos", "Curso de italiano", 120, "Cartão de crédito", "Fixo")
        add(dd(rnd.randint(14, 20)), "Despesa", "Cuidados pessoais", "Barbeiro", 15, "Dinheiro", "Variável")
        for _ in range(rnd.randint(5, 7)):
            add(dd(rnd.randint(1, dim)), "Despesa", "Mercado", rnd.choice(["Supermercado", "Feira", "Padaria", "Mercado do bairro"]),
                rnd.uniform(18, 70), rnd.choice(["Cartão de débito", "Cartão de crédito"]), "Variável")
        for _ in range(rnd.randint(3, 7)):
            add(dd(rnd.randint(1, dim)), "Despesa", "Restaurantes & cafés", rnd.choice(["Pizzaria", "Café com amigos", "Aperitivo", "Almoço fora", "Sushi"]),
                rnd.uniform(6, 34), "Cartão de crédito", "Variável")
        for _ in range(rnd.randint(1, 3)):
            add(dd(rnd.randint(1, dim)), "Despesa", "Lazer & passeios", rnd.choice(["Cinema", "Museu", "Show", "Passeio de bicicleta", "Escape room"]),
                rnd.uniform(8, 40), "Cartão de crédito", "Variável")
        if rnd.random() < 0.6:
            add(dd(rnd.randint(1, dim)), "Despesa", "Saúde & bem-estar", "Farmácia", rnd.uniform(6, 28), "Cartão de débito", "Variável")
        if rnd.random() < 0.5:
            add(dd(rnd.randint(1, dim)), "Despesa", "Transporte", "Trem intermunicipal", rnd.uniform(9, 32), "Cartão de crédito", "Variável")
        if rnd.random() < 0.45:
            add(dd(rnd.randint(1, dim)), "Despesa", "Roupas", rnd.choice(["Tênis de corrida", "Camisa", "Casaco", "Calça"]),
                rnd.uniform(30, 110), "Cartão de crédito", "Variável")
        if rnd.random() < 0.5:
            add(dd(rnd.randint(1, dim)), "Despesa", "Livros & materiais", "Livro", rnd.uniform(9, 24), "Cartão de crédito", "Variável")
        if rnd.random() < 0.35:
            add(dd(rnd.randint(1, dim)), "Despesa", "Casa & utensílios", rnd.choice(["Utensílios de cozinha", "Lâmpadas", "Organizadores"]),
                rnd.uniform(10, 55), "Cartão de crédito", "Variável")
        if m in (2, 5, 9):
            add(dd(rnd.randint(5, 25)), "Despesa", "Presentes & doações", "Presente de aniversário", rnd.uniform(20, 55), "Cartão de crédito", "Variável")
        add(dd(11), "Aporte", "Reserva de emergência", "Transferência para a reserva", 200, "PIX / transferência", "Fixo")
        if m >= 3:
            add(dd(11), "Aporte", "Investimentos", "Aporte mensal (ETF)", 150, "PIX / transferência", "Fixo")
    add(D(4, 3), "Despesa", "Viagens", "Páscoa — trem + hospedagem", 280, "Cartão de crédito", "Variável")
    add(D(8, 12), "Despesa", "Viagens", "Viagem para a Sicília", 640, "Cartão de crédito", "Variável", "Férias de verão")
    add(D(7, 18), "Despesa", "Impostos & taxas", "Renovação do permesso di soggiorno", 116, "Conta corrente", "Variável")
    add(D(3, 22), "Despesa", "Outros", "Correios", 12.5, "Dinheiro", "Variável")
    rows.sort(key=lambda x: x[0])
    ws = wb[S.FIN]
    for i, (d, tp, cat, desc, val, conta, fv, nota) in enumerate(rows):
        put(ws, FIN_R1 + i, {"B": d, "C": tp, "D": cat, "E": desc, "F": val, "G": conta, "H": fv, "I": nota})

    # ---------- Orçamento
    ws = wb[S.ORC]
    budget = {"Moradia": 650, "Contas da casa": 130, "Mercado": 280, "Transporte": 60, "Saúde & bem-estar": 55,
              "Dívidas & financiamentos": 91, "Impostos & taxas": 15, "Família & ajudas": 75, "Restaurantes & cafés": 100,
              "Assinaturas": 25, "Cuidados pessoais": 20, "Lazer & passeios": 60, "Viagens": 80, "Roupas": 40,
              "Casa & utensílios": 20, "Presentes & doações": 15, "Educação & cursos": 300, "Livros & materiais": 15, "Outros": 15}
    for i, (cat, _) in enumerate(CAT_DESP):
        ws[f"D{ORC_DESP1 + i}"] = budget.get(cat)
    for i, v in enumerate([2450, 150, None, 15]):
        if v:
            ws[f"D{ORC_REC1 + i}"] = v
    ws[f"D{ORC_APO1}"] = 200
    ws[f"D{ORC_APO1 + 1}"] = 150
    for i in range(9):
        c = MCOLS[i]
        ws[f"{c}79"] = round(1350 + rnd.uniform(-180, 260) + i * 25, 0)
        ws[f"{c}80"] = 4000 + 200 * (i + 1)
        ws[f"{c}81"] = round(2800 + 150 * max(0, i - 1) + 28 * i + rnd.uniform(-60, 90), 0)
        ws[f"{c}83"] = round(6200 - 90.83 * (i + 1), 0)

    # ---------- Saúde
    ws = wb[S.SAU]
    notes = ["Dia produtivo, treino bom.", "Cansado depois do trabalho.", "Ótima conversa com a família.",
             "Reunião difícil, mas resolvi bem.", "Dormi mal, cabeça pesada.", "Caminhada longa no parque, recarregou.",
             "Aula de italiano rendeu muito.", "Ansioso com prazos do Master.", "Dia leve, cozinhei em casa.",
             "Saí com amigos, me senti bem acompanhado."]
    n_days = (last_day - D(1, 1)).days + 1
    weight = 84.2
    gym_days = []
    for k in range(n_days):
        day = D(1, 1) + dt.timedelta(days=k)
        r = SAU_R1 + k
        prog = k / n_days
        if rnd.random() < 0.88:
            sono = max(4.5, min(9, rnd.gauss(6.5 + 0.8 * prog, 0.6)))
            hum = max(1, min(5, round(rnd.gauss(3.1 + 0.8 * prog, 0.8))))
            ene = max(1, min(5, round(rnd.gauss(3.0 + 0.7 * prog, 0.8))))
            est = max(1, min(5, round(rnd.gauss(3.3 - 0.8 * prog, 0.8))))
            put(ws, r, {"D": round(sono * 2) / 2, "E": max(1, min(5, round((sono - 4) / 1.1))), "F": hum, "G": ene, "H": est,
                        "K": int(rnd.gauss(7200 + 1800 * prog, 1800)), "M": round(rnd.uniform(1.2, 2.4), 1),
                        "N": max(1, min(5, round(rnd.gauss(3.2 + 0.6 * prog, 0.7))))})
            if rnd.random() < 0.18:
                ws[f"O{r}"] = rnd.choice(notes)
        wd = day.weekday()
        p_gym = {0: 0.75, 1: 0.25, 2: 0.75, 3: 0.2, 4: 0.6, 5: 0.55, 6: 0.25}[wd] * (0.75 + 0.35 * prog)
        if rnd.random() < p_gym:
            gym_days.append(k)
            tipo = rnd.choice(["Musculação", "Musculação", "Corrida", "Corrida", "Caminhada", "Yoga / alongamento", "Bicicleta"])
            put(ws, r, {"I": tipo, "J": rnd.choice([30, 40, 45, 50, 60, 75])})
        if wd == 0:
            weight = weight - rnd.uniform(-0.1, 0.36)
            ws[f"L{r}"] = round(weight, 1)
    cons = [(D(3, 10), "Check-up", "Clínico geral", "Realizado", D(3, 10) + dt.timedelta(days=365), "Tudo ok; repetir exames de sangue em 1 ano"),
            (D(3, 17), "Exame", "Exames de sangue", "Realizado", D(3, 17) + dt.timedelta(days=365), "Vitamina D um pouco baixa"),
            (D(6, 5), "Dentista", "Limpeza", "Realizado", D(12, 5), "Retorno em 6 meses"),
            (None, "Especialista", "Oftalmologista", "A agendar", D(10, 20), "Vista cansada no computador"),
            (D(10, 15), "Especialista", "Dermatologista", "Agendado", D(10, 15), "Mapeamento de pintas")]
    for i, (d, tp, prof, st, prox, nota) in enumerate(cons):
        put(ws, 27 + i, {"Q": d, "R": tp, "S": prof, "U": st, "V": prox, "X": nota})

    # ---------- Hábitos
    ws = wb[S.HAB]
    habits = [("Treinar", "Saúde física", 4, None), ("Meditar 10 min", "Saúde mental", 6, (0.45, 0.88)),
              ("Ler 20 páginas", "Aprendizado", 5, (0.5, 0.8)), ("Estudar italiano", "Aprendizado", 5, (0.5, 0.62)),
              ("Dormir até 23h30", "Saúde física", 6, (0.35, 0.6)), ("Journaling", "Saúde mental", 5, (0.4, 0.38)),
              ("Planejar o dia", "Casa & organização", 5, (0.5, 0.82)), ("Falar com alguém querido", "Amizades & social", 3, (0.3, 0.5)),
              ("Prática espiritual", "Propósito & espiritualidade", 4, (0.35, 0.32))]
    for i, (name, area, meta, p) in enumerate(habits):
        c = HAB_COLS[i]
        ws[f"{c}16"], ws[f"{c}6"], ws[f"{c}7"] = name, area, meta
    gym = set(gym_days)
    for k in range(n_days):
        prog = k / n_days
        r = HAB_R1 + k
        for i, (name, area, meta, p) in enumerate(habits):
            c = HAB_COLS[i]
            done = (k in gym) if p is None else rnd.random() < p[0] + (p[1] - p[0]) * prog
            if done:
                ws[f"{c}{r}"] = "x"

    # ---------- Roda da Vida
    ws = wb[S.RODA]
    base = [5, 4, 5, 5, 6, 7, 8, 4, 4, 5, 5]
    trend = [0.35, 0.4, 0.3, 0.15, 0.25, 0.1, 0.05, 0.3, 0.3, 0.2, 0.25]
    target = [8, 8, 8, 8, 8, 9, 9, 7, 7, 8, 8]
    for i in range(11):
        r = RODA_R1 + i
        ws[f"O{r}"] = target[i]
        for m in range(9):
            ws[f"{RCOLS[m]}{r}"] = max(1, min(10, round(base[i] + trend[i] * m + rnd.uniform(-0.8, 0.8))))
    reviews = [
        (6, "Comecei a rotina de treinos; 3 livros na fila.", "Dormi tarde quase todo dia.", "Sem plano escrito eu reajo à semana.", "Dormir até 23h30 em 5 dias da semana."),
        (6, "Primeiro mês inteiro sem estourar o orçamento.", "Pulei a meditação em dias corridos.", "Hábito pequeno vence hábito ambicioso.", "Meditar antes do café, não depois."),
        (7, "Freelance BIM entregue; elogio do cliente.", "Estudo do Master atrasou 2 semanas.", "Blocos fixos de estudo funcionam melhor.", "Recuperar o atraso do módulo 4."),
        (7, "Páscoa com amigos; voltei descansado.", "Gastei mais com restaurantes.", "Lazer planejado rende mais que lazer improvisado.", "Reconectar com 2 amigos do Brasil."),
        (7, "Concluí o curso de Revit MEP.", "Pouco tempo com a família.", "Ligar no domingo virou ritual.", "Organizar documentos do permesso."),
        (7, "Exames ok; renovei o permesso.", "Semana de muito estresse no trabalho.", "Dizer não para reunião sem pauta.", "Começar a procurar vagas melhores."),
        (8, "3 candidaturas enviadas; LinkedIn atualizado.", "Treinos caíram nas férias do escritório.", "Constância > intensidade.", "Viagem para a Sicília sem culpa."),
        (8, "Sicília! Uma semana de descanso real.", "Gasto da viagem acima do plano.", "Separar um fundo para viagens.", "Terminar o módulo 5 do Master."),
        (8, "Módulo 5 entregue; segunda entrevista marcada.", "Pouco contato com amigos daqui.", "Convidar é mais fácil do que esperar convite.", "Preparar a entrevista e fechar o diploma."),
    ]
    for m, (nota, vit, nao, apr, foco) in enumerate(reviews):
        r = REV_R1 + m
        put(ws, r, {"C": nota, "D": vit, "G": nao, "J": apr, "M": foco,
                    "P": rnd.choice(["Saúde da família.", "Amigos novos em Turim.", "Ter um trabalho estável.", "Dias de sol no fim de semana."])})

    # ---------- Metas
    ws = wb[S.METAS]
    ws["D5"] = "Engenheiro BIM reconhecido, trabalhando numa empresa que valoriza pessoas, com saúde em dia, reserva de 12 meses e uma vida social rica na Itália."
    ws["D6"] = "Construir coisas que duram — obras, relações e hábitos — e ajudar quem está começando."
    ws["D7"] = "Honestidade · constância · família · aprender sempre"
    ws["D8"] = "Constância"
    goals = [
        ("Concluir o Master BIM", "Carreira", "Ativa", dt.date(ano - 1, 10, 1), D(12, 15), "módulos", 0, 5, 8, None, "Entregar o módulo 6 até 31/10", "Abre portas para cargos de coordenação"),
        ("Atingir nível B2 em italiano", "Aprendizado", "Ativa", D(1, 8), D(12, 31), "nível", None, None, None, 0.6, "Fazer simulado de prova em outubro", "Trabalho e vida social"),
        ("Reserva de emergência de 6 meses", "Finanças", "Ativa", D(1, 1), D(12, 31), "€", 4000, 5800, 7000, None, "Manter aporte de € 200/mês", "Tranquilidade para mudar de emprego"),
        ("Chegar a 78 kg", "Saúde física", "Ativa", D(1, 1), D(12, 31), "kg", 84.2, None, 78, None, "Manter 4 treinos por semana", "Energia e saúde"),
        ("Correr 10 km em menos de 60 min", "Saúde física", "Ativa", D(3, 1), D(11, 30), "marco", None, None, None, 0.62, "Treino intervalado 1×/semana", ""),
        ("Conseguir posição em empresa melhor", "Carreira", "Ativa", D(6, 1), dt.date(ano + 1, 3, 31), "marco", None, None, None, 0.35, "Preparar a 2ª entrevista", "Ambiente mais organizado e salário maior"),
        ("Ler 12 livros em " + str(ano), "Aprendizado", "Ativa", D(1, 1), D(12, 31), "livros", 0, 9, 12, None, "Terminar o livro em italiano", ""),
        ("40 conversas longas com a família", "Família", "Ativa", D(1, 1), D(12, 31), "conversas", 0, 31, 40, None, "Ligação de domingo", "Estar presente mesmo de longe"),
        ("Rede de 8 amigos próximos em Turim", "Amizades & social", "Ativa", D(1, 1), D(12, 31), "pessoas", 2, 6, 8, None, "Convidar 1 pessoa nova por semana", ""),
        ("Validar o diploma de engenharia", "Casa & organização", "Ativa", D(2, 1), D(9, 30), "marco", None, None, None, 0.7, "Enviar a tradução juramentada", "Necessário para inscrição na ordem"),
        ("Meditar 250 sessões", "Saúde mental", "Ativa", D(1, 1), D(12, 31), "sessões", 0, 0, 250, None, "Meditar antes do café", ""),
        ("Publicar 4 artigos técnicos no LinkedIn", "Carreira", "Ativa", D(4, 1), D(12, 31), "artigos", 0, 2, 4, None, "Rascunho sobre Dynamo", "Marca pessoal"),
        ("Voluntariado 1×/mês", "Propósito & espiritualidade", "Ativa", D(1, 1), D(12, 31), "ações", 0, 6, 12, None, "Inscrever-se no mutirão de novembro", "Contribuir"),
        ("Viagem para a Sicília", "Lazer & criatividade", "Concluída", D(2, 1), D(8, 20), "marco", None, None, None, 1, "", "Primeira viagem longa na Itália"),
        ("Organizar a vida digital (backup, senhas, pastas)", "Casa & organização", "Ativa", D(9, 20), D(11, 15), "marco", None, None, None, 0.2, "Backup completo do notebook", "Paz de espírito"),
        ("Correr a meia maratona de Turim", "Saúde física", "Ativa", D(9, 1), dt.date(ano + 1, 4, 15), "km no longão", 8, 12, 21, None, "Subir o longão 1 km a cada 2 semanas", ""),
        ("Reserva de 12 meses de despesas", "Finanças", "Ativa", D(1, 1), dt.date(ano + 2, 12, 31), "€", 4000, 5800, 24000, None, "Aumentar aporte quando o salário subir", "Liberdade para escolher trabalho"),
        ("Cidadania italiana", "Casa & organização", "Ativa", D(1, 1), dt.date(ano + 3, 6, 30), "marco", None, None, None, 0.15, "Juntar certidões do bisavô", "Estabilidade de longo prazo"),
        ("Dar aulas em uma pós-graduação BIM", "Carreira", "Ativa", D(1, 1), dt.date(ano + 4, 12, 31), "marco", None, None, None, 0.1, "Concluir o Master primeiro", "Ensinar o que aprendi"),
    ]
    medit = sum(1 for k in range(n_days) if wb[S.HAB][f"E{HAB_R1 + k}"].value == "x")
    for i, g in enumerate(goals):
        r = MET_R1 + i
        name, area, st, ini, prazo, un, v0, v1, v2, man, prox, why = g
        if name.startswith("Chegar a 78"):
            v1 = round(weight, 1)
        if name.startswith("Meditar"):
            v1 = medit
        put(ws, r, {"C": name, "D": area, "E": st, "F": ini, "G": prazo, "H": un, "I": v0, "J": v1, "K": v2,
                    "L": man, "S": prox, "T": why})

    # ---------- Tarefas
    ws = wb[S.TAR]
    projects = [
        ("Master BIM — módulo 6", "Carreira", "Em andamento", D(9, 1), D(10, 31), "Concluir o Master BIM", "Modelagem MEP + relatório"),
        ("Troca de emprego", "Carreira", "Em andamento", D(6, 1), dt.date(ano + 1, 3, 31), "Conseguir posição em empresa melhor", "Funil: candidaturas, entrevistas, propostas"),
        ("Validação do diploma", "Casa & organização", "Em andamento", D(2, 1), D(10, 31), "Validar o diploma de engenharia", "Tradução, apostila, envio"),
        ("Organização digital", "Casa & organização", "Planejado", D(10, 5), D(11, 15), "", "Backup, senhas, pastas"),
        ("Natal no Brasil", "Família", "Planejado", D(10, 1), D(12, 15), "", "Passagens, presentes, agenda"),
        ("App de rotina pessoal", "Carreira", "Ideia", None, None, "", "Ideia antiga: app para organizar a rotina com lembretes inteligentes"),
        ("Canal de conteúdo BIM", "Carreira", "Ideia", None, None, "", "Vídeos curtos com dicas de Revit/Dynamo"),
    ]
    for i, (n, a, st, ini, prz, meta, nota) in enumerate(projects):
        put(ws, PRJ_R1 + i, {"B": n, "C": a, "D": st, "E": ini, "F": prz, "I": meta, "M": nota})
    tasks = [
        ("Modelar instalações hidráulicas", "Master BIM — módulo 6", "Carreira", "Alta", D(10, 10), "Em andamento", None, "Concluir o Master BIM"),
        ("Escrever relatório do módulo 6", "Master BIM — módulo 6", "Carreira", "Alta", D(10, 28), "A fazer", None, "Concluir o Master BIM"),
        ("Revisar famílias Revit do módulo", "Master BIM — módulo 6", "Carreira", "Média", D(9, 26), "A fazer", None, "Concluir o Master BIM"),
        ("Ler capítulo de coordenação 4D", "Master BIM — módulo 6", "Carreira", "Baixa", D(9, 20), "Concluída", D(9, 19), "Concluir o Master BIM"),
        ("Preparar a 2ª entrevista (case)", "Troca de emprego", "Carreira", "Alta", D(10, 6), "Em andamento", None, "Conseguir posição em empresa melhor"),
        ("Atualizar portfólio com projeto freelance", "Troca de emprego", "Carreira", "Média", D(10, 3), "A fazer", None, "Conseguir posição em empresa melhor"),
        ("Follow-up com recrutadora", "Troca de emprego", "Carreira", "Média", D(9, 29), "A fazer", None, "Conseguir posição em empresa melhor"),
        ("Enviar 3 candidaturas", "Troca de emprego", "Carreira", "Alta", D(9, 15), "Concluída", D(9, 14), "Conseguir posição em empresa melhor"),
        ("Atualizar LinkedIn", "Troca de emprego", "Carreira", "Média", D(7, 10), "Concluída", D(7, 8), "Conseguir posição em empresa melhor"),
        ("Enviar tradução juramentada", "Validação do diploma", "Casa & organização", "Alta", D(9, 25), "Aguardando", None, "Validar o diploma de engenharia"),
        ("Pagar taxa da apostila", "Validação do diploma", "Casa & organização", "Média", D(9, 10), "Concluída", D(9, 12), "Validar o diploma de engenharia"),
        ("Agendar atendimento no consulado", "Validação do diploma", "Casa & organização", "Média", D(10, 15), "A fazer", None, "Validar o diploma de engenharia"),
        ("Backup completo do notebook", "Organização digital", "Casa & organização", "Média", D(10, 12), "A fazer", None, ""),
        ("Gerenciador de senhas", "Organização digital", "Casa & organização", "Baixa", D(10, 30), "A fazer", None, ""),
        ("Pesquisar passagens para dezembro", "Natal no Brasil", "Família", "Alta", D(10, 8), "A fazer", None, ""),
        ("Lista de presentes", "Natal no Brasil", "Família", "Baixa", D(11, 20), "A fazer", None, ""),
        ("Marcar oftalmologista", "", "Saúde física", "Média", D(10, 9), "A fazer", None, ""),
        ("Renovar CNH brasileira (agendar)", "", "Casa & organização", "Alta", D(10, 4), "A fazer", None, ""),
        ("Inscrição no mutirão de novembro", "", "Propósito & espiritualidade", "Baixa", D(10, 20), "A fazer", None, "Voluntariado 1×/mês"),
        ("Rascunho de artigo sobre Dynamo", "", "Carreira", "Média", D(9, 30), "A fazer", None, "Publicar 4 artigos técnicos no LinkedIn"),
        ("Declaração de imposto (Brasil)", "", "Finanças", "Alta", D(5, 29), "Concluída", D(5, 27), ""),
        ("Revisar orçamento de setembro", "", "Finanças", "Média", D(9, 30), "Concluída", D(9, 30), "Reserva de emergência de 6 meses"),
        ("Comprar tênis de corrida", "", "Saúde física", "Baixa", D(8, 30), "Concluída", D(9, 2), ""),
        ("Ligar para a tia no aniversário", "", "Família", "Média", D(10, 28), "A fazer", None, ""),
        ("Planejar fim de semana em Gênova", "", "Lazer & criatividade", "Baixa", None, "A fazer", None, ""),
        ("Trocar operadora de internet", "", "Casa & organização", "Baixa", D(8, 15), "Cancelada", None, ""),
    ]
    for i, (n, p, a, pr, prz, st, concl, meta) in enumerate(tasks):
        put(ws, TSK_R1 + i, {"B": n, "C": p, "D": a, "E": pr, "F": prz, "G": st, "H": concl, "I": meta})

    # ---------- Carreira
    ws = wb[S.CAR]
    put(ws, 6, {"C": "Engenheiro de projetos BIM"})
    put(ws, 7, {"C": "Studio Tecnico (fictício)"})
    put(ws, 8, {"C": dt.date(ano - 2, 3, 1)})
    put(ws, 9, {"C": 2450})
    put(ws, 10, {"C": 5})
    put(ws, 11, {"C": "Coordenador BIM em empresa organizada, com salário ≥ € 2.900"})
    put(ws, 12, {"C": "Liderar a área BIM de uma construtora e dar aulas em pós-graduação"})
    comps = [("Revit / modelagem BIM", "Ferramenta", 3, 5, "Master + projetos freelance", D(12, 31), "Alta"),
             ("Coordenação e clash detection", "Técnica", 2, 4, "Módulo 6 do Master + Navisworks", D(12, 31), "Alta"),
             ("Dynamo / Python", "Ferramenta", 2, 4, "Curso online + 1 script por mês", dt.date(ano + 1, 3, 31), "Média"),
             ("Italiano técnico", "Idioma", 3, 5, "Escola + leitura técnica", D(12, 31), "Alta"),
             ("Gestão de projetos", "Gestão", 3, 4, "Liderar o próximo projeto interno", dt.date(ano + 1, 6, 30), "Média"),
             ("Comunicação e apresentação", "Comportamental", 3, 4, "Publicar artigos; apresentar no time", D(12, 31), "Média"),
             ("Inglês", "Idioma", 4, 4, "Manter com podcasts", None, "Baixa"),
             ("Orçamento de obras", "Técnica", 2, 3, "Ler norma + planilha prática", dt.date(ano + 1, 6, 30), "Baixa")]
    for i, (n, cat, a, alv, how, prz, pri) in enumerate(comps):
        put(ws, CMP_R1 + i, {"B": n, "C": cat, "D": a, "E": alv, "H": how, "I": prz, "J": pri})
    opps = [("Construtora Alfa", "Coordenador BIM", D(9, 2), "Entrevista", "Preparar o case técnico", D(10, 6), 2900, "LinkedIn"),
            ("Engenharia Beta", "BIM Specialist", D(9, 14), "Aplicado", "Follow-up por e-mail", D(9, 29), 2700, "Indicação do Marco"),
            ("Studio Gama", "Projetista MEP", D(9, 14), "Aplicado", "", None, 2600, "Site da empresa"),
            ("Infra Delta", "BIM Manager Jr.", D(7, 20), "Recusado", "", None, 3000, "Recrutadora"),
            ("Arquitetura Épsilon", "Modelador BIM", D(9, 22), "Interesse", "Pesquisar a empresa", D(10, 10), 2500, "Evento do Master"),
            ("Consultoria Zeta", "BIM Coordinator", D(8, 25), "Teste / case", "Entregar teste técnico", D(10, 3), 3100, "LinkedIn")]
    for i, (e, c, d, et, acao, da, rem, fonte) in enumerate(opps):
        put(ws, OPP_R1 + i, {"B": e, "C": c, "D": d, "E": et, "H": acao, "I": da, "J": rem, "K": fonte})
    wins = [("Entreguei modelo federado sem retrabalho", "Entrega relevante", D(2, 20), 4, "0 conflitos na revisão", "Cliente elogiou por e-mail"),
            ("Curso Revit MEP avançado", "Certificação", D(5, 30), 3, "Certificado", ""),
            ("Freelance BIM para escritório local", "Conquista", D(3, 28), 4, "€ 380 + novo cliente", ""),
            ("Feedback positivo do gestor", "Feedback recebido", D(4, 15), 3, "", "Organização da documentação"),
            ("Apresentei padrão de famílias ao time", "Palestra / conteúdo", D(6, 12), 3, "Time adotou o padrão", ""),
            ("Script Dynamo que poupa 2h/semana", "Entrega relevante", D(7, 3), 5, "≈ 8 h/mês economizadas", "Compartilhado no repositório do time"),
            ("Primeiro artigo no LinkedIn", "Palestra / conteúdo", D(8, 28), 2, "1.200 visualizações", ""),
            ("Módulo 5 do Master aprovado", "Conquista", D(9, 10), 4, "Nota 28/30", "")]
    for i, (n, t, d, imp, res, ev) in enumerate(wins):
        put(ws, WIN_R1 + i, {"B": n, "C": t, "D": d, "E": imp, "F": res, "H": ev})

    # ---------- Relações
    ws = wb[S.REL]
    people = [("Mãe", "Família", "Íntimo", 7, dt.date(1965, 11, 3), "Gosta de orquídeas", 6),
              ("Pai", "Família", "Íntimo", 7, dt.date(1962, 10, 19), "Futebol, churrasco", 7),
              ("Rafael (irmão)", "Família", "Íntimo", 14, dt.date(1995, 4, 22), "Começou mestrado", 12),
              ("Tia Regina", "Família", "Próximo", 60, dt.date(1958, 10, 28), "", None),
              ("Chiara", "Parceria", "Íntimo", 3, dt.date(1996, 12, 8), "Ama cinema e trilhas", 3),
              ("Marco", "Amizade", "Próximo", 14, dt.date(1991, 10, 11), "Colega do Master", 12),
              ("Pedro (Brasil)", "Amizade", "Próximo", 30, dt.date(1990, 2, 14), "Amigo de infância", None),
              ("Lucas", "Amizade", "Próximo", 21, dt.date(1992, 6, 30), "", 18),
              ("Francesca", "Amizade", "Social", 30, dt.date(1994, 3, 9), "Grupo de corrida", 25),
              ("Davide", "Amizade", "Social", 30, dt.date(1989, 8, 25), "Vizinho, joga basquete", 26),
              ("Paolo (gestor)", "Trabalho", "Social", 30, dt.date(1978, 5, 15), "", 20),
              ("Prof.ª Rossi", "Mentoria", "Próximo", 45, dt.date(1970, 1, 20), "Orientadora do Master", None),
              ("Grupo de voluntariado", "Comunidade", "Social", 30, None, "Mutirões mensais", 28)]
    inter = []
    for i, (n, rel_, circ, freq, aniv, nota, step) in enumerate(people):
        put(ws, PPL_R1 + i, {"B": n, "C": rel_, "D": circ, "E": freq, "I": aniv, "M": nota})
        stop = {"Tia Regina": D(6, 8), "Pedro (Brasil)": D(7, 20), "Prof.ª Rossi": D(7, 2)}.get(n, last_day)
        stepd = step or freq
        d = D(1, 1) + dt.timedelta(days=rnd.randint(1, stepd))
        while d <= stop:
            tipo = {"Família": rnd.choice(["Ligação", "Videochamada", "Videochamada", "Mensagem"]),
                    "Parceria": rnd.choice(["Encontro", "Encontro", "Encontro", "Videochamada"]),
                    "Comunidade": "Evento / grupo"}.get(rel_, rnd.choice(["Encontro", "Mensagem", "Ligação", "Encontro"]))
            dur = {"Mensagem": 10, "Ligação": rnd.choice([20, 30, 45]), "Videochamada": rnd.choice([30, 45, 60]),
                   "Encontro": rnd.choice([60, 90, 120, 180]), "Evento / grupo": 180}[tipo]
            inter.append((d, n, tipo, max(1, min(5, round(rnd.gauss(3.8 if tipo != "Mensagem" else 2.8, 0.7)))), dur,
                          rnd.choice(["", "", "Conversa boa", "Planejamos algo juntos", "Falamos de trabalho", "Rimos muito"])))
            d = d + dt.timedelta(days=max(1, int(stepd * rnd.uniform(0.7, 1.25))))
        if n == "Tia Regina":
            inter.append((D(1, 15), n, "Ligação", 4, 30, ""))
    inter.sort(key=lambda x: x[0])
    for i, (d, n, tp, ql, du, nt) in enumerate(inter[: INT_R2 - INT_R1 + 1]):
        put(ws, INT_R1 + i, {"Q": d, "R": n, "S": tp, "T": ql, "U": du, "V": nt})

    # ---------- Aprendizado
    ws = wb[S.APR]
    items = [("Master BIM — módulos", "Curso", "Carreira", "Em andamento", dt.date(ano - 1, 10, 1), None, 8, 5, None, "Coordenação 4D muda o jeito de planejar obra", "Módulo 6 até 31/10"),
             ("Italiano B1 → B2", "Idioma", "Aprendizado", "Em andamento", D(1, 8), None, 120, 78, None, "Ler em voz alta acelera a fluência", "Simulado em outubro"),
             ("Python para engenheiros (Dynamo)", "Curso", "Carreira", "Em andamento", D(5, 2), None, 40, 14, None, "", ""),
             ("Il nome della rosa (em italiano)", "Livro", "Aprendizado", "Em andamento", D(9, 3), None, 500, 120, None, "", ""),
             ("O Evangelho Segundo o Espiritismo", "Livro", "Propósito & espiritualidade", "Em andamento", D(2, 1), None, 400, 230, None, "", ""),
             ("Hábitos Atômicos", "Livro", "Aprendizado", "Concluído", D(1, 2), D(1, 28), 320, 320, 5, "Sistemas > metas; identidade antes de resultado", ""),
             ("Deep Work", "Livro", "Carreira", "Concluído", D(2, 1), D(2, 26), 300, 300, 4, "Blocos de foco de 90 min", ""),
             ("Mindset", "Livro", "Saúde mental", "Concluído", D(3, 1), D(3, 24), 310, 310, 4, "", ""),
             ("Essencialismo", "Livro", "Aprendizado", "Concluído", D(4, 2), D(4, 25), 270, 270, 5, "Menos, porém melhor", ""),
             ("Revit MEP avançado", "Curso", "Carreira", "Concluído", D(3, 15), D(5, 30), 24, 24, 4, "", ""),
             ("A Psicologia Financeira", "Livro", "Finanças", "Concluído", D(5, 3), D(5, 29), 300, 300, 5, "Riqueza é o que você não vê", ""),
             ("Comunicação Não-Violenta", "Livro", "Amor & parceria", "Concluído", D(6, 2), D(6, 30), 280, 280, 4, "Observar sem julgar", ""),
             ("Sapiens", "Livro", "Aprendizado", "Concluído", D(7, 1), D(8, 5), 460, 460, 4, "", ""),
             ("O Poder do Agora", "Livro", "Saúde mental", "Concluído", D(8, 6), D(8, 30), 240, 240, 3, "", ""),
             ("Pai Rico, Pai Pobre", "Livro", "Finanças", "Concluído", D(9, 1), D(9, 21), 200, 200, 3, "", ""),
             ("Certificação Autodesk Revit Professional", "Certificação", "Carreira", "Quero fazer", None, None, None, None, None, "", "Depois do Master")]
    for i, it in enumerate(items):
        n, tp, ar, st, ini, fim, tot, at, nota, apr, prox = it
        put(ws, APR_R1 + i, {"B": n, "C": tp, "D": ar, "E": st, "F": ini, "G": fim, "H": tot, "I": at, "L": nota, "M": apr, "N": prox})
    logs = []
    d = D(1, 2)
    while d <= last_day:
        wd = d.weekday()
        if wd in (1, 3) and rnd.random() < 0.85:
            logs.append((d, "Italiano B1 → B2", rnd.choice([1, 1.5]), rnd.choice(["Aula", "Aula", "Conversação"]), ""))
        if wd in (5, 6) and rnd.random() < 0.8:
            logs.append((d, "Master BIM — módulos", rnd.choice([1.5, 2, 2.5, 3]), rnd.choice(["Aula", "Prática", "Exercícios"]), ""))
        if wd == 2 and d >= D(5, 2) and rnd.random() < 0.6:
            logs.append((d, "Python para engenheiros (Dynamo)", 1, "Prática", ""))
        if wd == 0 and rnd.random() < 0.35:
            logs.append((d, "Revisão de vocabulário técnico", 0.5, "Revisão", ""))
        d += dt.timedelta(days=1)
    for i, (d, it, h, at, nt) in enumerate(logs[: LOG_R2 - LOG_R1 + 1]):
        put(ws, LOG_R1 + i, {"Q": d, "R": it, "S": h, "T": at, "U": nt})

    # ---------- Lazer
    ws = wb[S.LAZ]
    acts = [("Trilha nas colinas", "Esporte / ar livre", "Chiara", 4, 0), ("Cinema", "Cultura", "Chiara", 2.5, 18),
            ("Jogo de basquete", "Esporte / ar livre", "Davide", 2, 0), ("Museu Egípcio", "Cultura", "Marco", 3, 15),
            ("Jantar com amigos", "Social", "Marco, Francesca", 3, 30), ("Tocar violão", "Música", "", 1, 0),
            ("Série no sofá", "Descanso", "Chiara", 2, 0), ("Pedal no parque Valentino", "Esporte / ar livre", "", 2, 0),
            ("Cozinhar receita nova", "Hobby criativo", "Chiara", 2, 12), ("Jogos de tabuleiro", "Jogos", "Lucas, Davide", 3, 0),
            ("Aperitivo no centro", "Social", "Francesca", 2, 14), ("Show de jazz", "Música", "Chiara", 3, 25)]
    lz = []
    d = D(1, 3)
    while d <= last_day:
        if rnd.random() < 0.3 + (0.15 if d.weekday() >= 5 else 0):
            a = rnd.choice(acts)
            lz.append((d, a[0], a[1], a[2], a[3], a[4], max(1, min(5, round(rnd.gauss(4, 0.7)))), "Sim" if rnd.random() < 0.85 else "Não", ""))
        d += dt.timedelta(days=1)
    lz.append((D(4, 4), "Páscoa em Florença", "Viagem", "Amigos", 16, 0, 5, "Sim", "Viagem curta, muito bem aproveitada"))
    lz.append((D(8, 13), "Sicília — semana de férias", "Viagem", "Chiara", 40, 0, 5, "Sim", "Palermo, Cefalù, Taormina"))
    lz.sort(key=lambda x: x[0])
    for i, (d, a, cat, quem, h, c, sat, rep, nt) in enumerate(lz[: LAZ_R2 - LAZ_R1 + 1]):
        put(ws, LAZ_R1 + i, {"B": d, "C": a, "D": cat, "E": quem, "F": h, "G": c, "H": sat, "I": rep, "J": nt})
    dreams = [("Ver a aurora boreal na Noruega", "Viagem", 2500, dt.date(ano + 2, 2, 1), "Sonho", "Média", ""),
              ("Correr a meia maratona de Turim", "Esporte / ar livre", 40, dt.date(ano + 1, 4, 15), "Planejando", "Alta", ""),
              ("Aprender a cozinhar risotto de verdade", "Hobby criativo", 80, D(11, 30), "Agendado", "Baixa", "Aula marcada"),
              ("Viagem para a Sicília", "Viagem", 700, D(8, 12), "Realizado", "Alta", ""),
              ("Fim de semana em Gênova", "Viagem", 250, D(11, 8), "Planejando", "Média", ""),
              ("Show na Arena de Verona", "Música", 150, dt.date(ano + 1, 7, 1), "Sonho", "Média", ""),
              ("Esquiar nos Alpes", "Esporte / ar livre", 400, dt.date(ano + 1, 1, 20), "Planejando", "Alta", ""),
              ("Levar os pais para conhecer a Itália", "Viagem", 3000, dt.date(ano + 1, 9, 1), "Sonho", "Alta", "")]
    for i, (n, cat, c, q_, st, pr, nt) in enumerate(dreams):
        put(ws, BKT_R1 + i, {"L": n, "M": cat, "N": c, "O": q_, "P": st, "Q": pr, "R": nt})

    # ---------- Casa
    ws = wb[S.CASA]
    docs = [("Permesso di soggiorno", "Eu", "—", dt.date(ano - 1, 11, 20), D(11, 20), "Juntar documentos e agendar na Questura", ""),
            ("Passaporte brasileiro", "Eu", "—", dt.date(ano - 4, 5, 10), dt.date(ano + 5, 5, 9), "", ""),
            ("CNH brasileira", "Eu", "—", dt.date(ano - 9, 10, 25), D(10, 25), "Agendar renovação no consulado", "Usada para a conversão da patente"),
            ("Carta d'identità", "Eu", "—", dt.date(ano - 1, 2, 1), dt.date(ano + 9, 2, 1), "", ""),
            ("Tessera sanitaria", "Eu", "—", dt.date(ano - 1, 11, 25), D(11, 25), "Renova junto com o permesso", ""),
            ("Codice fiscale", "Eu", "—", None, None, "", "Não vence"),
            ("Contrato de aluguel", "Eu", "—", dt.date(ano - 2, 7, 1), dt.date(ano + 2, 6, 30), "", "Aviso prévio de 6 meses"),
            ("Seguro do apartamento", "Eu", "—", D(1, 15), dt.date(ano + 1, 1, 15), "", "")]
    for i, (n, tit, num, em, val, acao, nt) in enumerate(docs):
        put(ws, DOC_R1 + i, {"B": n, "C": tit, "D": num, "E": em, "F": val, "I": acao, "J": nt})
    rots = [("Trocar roupa de cama", "Limpeza", 7, D(9, 28), "Eu", 15), ("Limpeza pesada do banheiro", "Limpeza", 14, D(9, 14), "Eu", 40),
            ("Limpar a geladeira", "Limpeza", 30, D(9, 6), "Eu", 30), ("Backup do computador", "Digital / backup", 30, D(8, 20), "Eu", 20),
            ("Revisar senhas importantes", "Digital / backup", 90, D(7, 1), "Eu", 30), ("Conferir extratos e lançamentos", "Finanças / admin", 7, D(9, 27), "Eu", 20),
            ("Fechar o mês no Orçamento", "Finanças / admin", 30, D(9, 30), "Eu", 30), ("Limpar filtro do ar-condicionado", "Manutenção", 90, D(6, 15), "Eu", 20),
            ("Regar e adubar as plantas", "Plantas", 7, D(9, 29), "Eu", 10), ("Revisão da bicicleta", "Veículo", 180, D(4, 10), "Eu", 60),
            ("Descongelar o freezer", "Limpeza", 90, D(7, 20), "Eu", 45), ("Organizar papéis e documentos", "Finanças / admin", 30, D(9, 5), "Eu", 30)]
    for i, (n, cat, fq, ult, resp, dur) in enumerate(rots):
        put(ws, ROT_R1 + i, {"B": n, "C": cat, "D": fq, "E": ult, "I": resp, "J": dur})
    subs = [("Streaming de vídeo", "Assinaturas", 9.99, "Mensal", D(10, 3), "Médio"), ("Música", "Assinaturas", 10.99, "Mensal", D(10, 6), "Alto"),
            ("Compras online (frete grátis)", "Assinaturas", 49.9, "Anual", dt.date(ano + 1, 2, 10), "Médio"),
            ("Leitura digital ilimitada", "Assinaturas", 9.99, "Mensal", D(10, 14), "Baixo"),
            ("Armazenamento em nuvem", "Assinaturas", 2.99, "Mensal", D(10, 20), "Alto"), ("Academia", "Saúde & bem-estar", 34.9, "Mensal", D(10, 4), "Alto"),
            ("Celular", "Contas da casa", 7.99, "Mensal", D(10, 8), "Alto"), ("Internet fibra", "Contas da casa", 27.95, "Mensal", D(10, 5), "Alto"),
            ("App de idiomas (premium)", "Educação & cursos", 84, "Anual", D(12, 1), "Médio"),
            ("Rede profissional (premium)", "Assinaturas", 39.99, "Mensal", D(10, 18), "Baixo")]
    for i, (n, cat, v, per, ren, uso) in enumerate(subs):
        put(ws, ASS_R1 + i, {"B": n, "C": cat, "D": v, "E": per, "H": ren, "J": uso})


# =============================================================================
# Montagem
# =============================================================================
def build(path, ano, sample):
    K.clear()
    CALC.clear()
    LIST_COL.clear()
    wb = Workbook()
    wb._fonts = IndexedList([Font(name=FONT, size=10)])   # fonte padrão do arquivo
    wb.remove(wb.active)
    for name in ORDER:
        wb.create_sheet(name)
    build_config(wb, wb[S.CFG], ano, "")
    build_financas(wb[S.FIN])
    build_orcamento(wb[S.ORC])
    build_saude(wb[S.SAU])
    build_habitos(wb[S.HAB])
    build_metas(wb[S.METAS])
    build_tarefas(wb[S.TAR])
    build_carreira(wb[S.CAR])
    build_relacoes(wb[S.REL])
    build_aprendizado(wb[S.APR])
    build_lazer(wb[S.LAZ])
    build_casa(wb[S.CASA])
    build_roda(wb[S.RODA])
    build_calculos(wb[S.CALC])
    build_dashboard(wb, wb[S.DASH], sample)
    build_guia(wb[S.GUIA], sample)
    # listas dinâmicas que vêm de outras abas
    wb.defined_names["lst_Projetos"] = DefinedName("lst_Projetos", attr_text=(
        f"OFFSET({A(S.TAR, 'B', PRJ_R1)},0,0,MAX(1,COUNTA({A(S.TAR, 'B', PRJ_R1, 'B', PRJ_R2)})),1)"))
    wb.defined_names["lst_Metas"] = DefinedName("lst_Metas", attr_text=(
        f"OFFSET({A(S.METAS, 'C', MET_R1)},0,0,MAX(1,COUNTA({A(S.METAS, 'C', MET_R1, 'C', MET_R2)})),1)"))
    wb.defined_names["lst_Pessoas"] = DefinedName("lst_Pessoas", attr_text=(
        f"OFFSET({A(S.REL, 'B', PPL_R1)},0,0,MAX(1,COUNTA({A(S.REL, 'B', PPL_R1, 'B', PPL_R2)})),1)"))
    wb.defined_names["lst_ItensApr"] = DefinedName("lst_ItensApr", attr_text=(
        f"OFFSET({A(S.APR, 'B', APR_R1)},0,0,MAX(1,COUNTA({A(S.APR, 'B', APR_R1, 'B', APR_R2)})),1)"))
    if sample:
        fill_sample(wb, ano)
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = ws.sheet_view.zoomScale or 100
    wb[S.CALC].sheet_properties.tabColor = GRAY
    wb.active = 0
    wb[S.DASH].sheet_view.tabSelected = True
    for ws in wb.worksheets[1:]:
        ws.sheet_view.tabSelected = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


# =============================================================================
# Valores pré-calculados (opcional, requer LibreOffice)
# =============================================================================
# O openpyxl grava fórmulas sem resultado. O Excel recalcula ao abrir, mas pré-visualizações
# (celular, navegador) mostrariam células vazias. Aqui o LibreOffice calcula uma cópia e os
# resultados são gravados como cache no arquivo original — estilos, gráficos e validações
# continuam exatamente como o openpyxl gerou. O Excel recalcula tudo de novo ao abrir.
_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalcularESalvar()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


def recalcular_libreoffice(src, dst, timeout=300):
    if not shutil.which("soffice"):
        return False
    shutil.copyfile(src, dst)
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="lo_perfil_") as prof:
        url = Path(prof).as_uri()
        subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       env=env, capture_output=True, timeout=120)
        macro_dir = Path(prof) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            return False
        (macro_dir / "Module1.xba").write_text(_MACRO, encoding="utf-8")
        r = subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
                            "vnd.sun.star.script:Standard.Module1.RecalcularESalvar?language=Basic&location=application",
                            str(dst)], env=env, capture_output=True, timeout=timeout)
    return r.returncode == 0


def _serial(v):
    if isinstance(v, dt.datetime):
        d = v - dt.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, dt.date):
        return (v - dt.date(1899, 12, 30)).days
    if isinstance(v, dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


def injetar_valores(path, calc_path):
    from openpyxl import load_workbook
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        vals = load_workbook(calc_path, data_only=True)
    pat = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?)><f>(.*?)</f><v ?/></c>')
    with zipfile.ZipFile(path) as zin:
        infos = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in infos}
    total = feitos = erros = 0
    for i, name in enumerate(ORDER):
        key = f"xl/worksheets/sheet{i + 1}.xml"
        xml = data[key].decode("utf-8")
        ws = vals[name]
        total += xml.count("<f>")

        def rep(m):
            nonlocal feitos, erros
            ref, style, f = m.groups()
            v = _serial(ws[ref].value)
            feitos += 1
            if v is None:
                return f'<c r="{ref}"{style} t="str"><f>{f}</f><v></v></c>'
            if isinstance(v, bool):
                return f'<c r="{ref}"{style} t="b"><f>{f}</f><v>{int(v)}</v></c>'
            if isinstance(v, (int, float)):
                return f'<c r="{ref}"{style}><f>{f}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
            v = str(v)
            if v.startswith("#"):
                erros += 1
            return f'<c r="{ref}"{style} t="str"><f>{f}</f><v>{escape(v)}</v></c>'

        data[key] = pat.sub(rep, xml).encode("utf-8")
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, data[info.filename])
    os.replace(tmp, path)
    return total, feitos, erros


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--ano", type=int, default=2026)
    ap.add_argument("--saida", default=os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument("--sem-valores", action="store_true", help="não pré-calcular com LibreOffice")
    a = ap.parse_args()
    for sample in (False, True):
        name = f"Sistema_de_Vida_{a.ano}{'_EXEMPLO' if sample else ''}.xlsx"
        path = build(os.path.join(a.saida, name), a.ano, sample)
        msg = "fórmulas sem valores em cache (o Excel calcula ao abrir)"
        if not a.sem_valores:
            with tempfile.TemporaryDirectory() as tmpd:
                calc = os.path.join(tmpd, "calc.xlsx")
                if recalcular_libreoffice(path, calc):
                    total, feitos, erros = injetar_valores(path, calc)
                    msg = f"{feitos}/{total} fórmulas com valor em cache, {erros} erro(s)"
                    if feitos != total or erros:
                        raise SystemExit(f"Falha ao gravar valores: {msg}")
        print(f"{path}: {msg}")


if __name__ == "__main__":
    main()
